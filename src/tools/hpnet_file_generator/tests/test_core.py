import pytest
from openpyxl import Workbook
from pathlib import Path
from tools.hpnet_file_generator.models.data_models import PersonRecord, Parcel, SourceFile, ProfileConfig, ActionStatus
from tools.hpnet_file_generator.utils.text_normalizer import normalize_person_name, extract_stt_and_name
from tools.hpnet_file_generator.core.person_matcher import PersonMatcher
from tools.hpnet_file_generator.core.naming_engine import NamingEngine
from tools.hpnet_file_generator.core.action_planner import ActionPlanner
from tools.hpnet_file_generator.core.excel_reader import ExcelReader
from tools.hpnet_file_generator.core.source_scanner import SourceScanner


@pytest.mark.parametrize('filename,name', [
    ('1.Nguyễn Thị Hảo_0001.pdf', 'Nguyễn Thị Hảo'),
    ('10. Nguyễn Thị Ngọt_0001.pdf', 'Nguyễn Thị Ngọt'),
    ('102 Lê Văn Thuận_0001.pdf', 'Lê Văn Thuận'),
    ('110 Lê Văn Thời 0001.pdf', 'Lê Văn Thời'),
    ('Nguyễn Văn A.pdf', 'Nguyễn Văn A'),
    ('Nguyễn Văn A_0001_002.pdf', 'Nguyễn Văn A'),
])
def test_scan_matches_only_person_name(tmp_path, filename, name):
    (tmp_path / filename).write_bytes(b'pdf')
    sources = SourceScanner(str(tmp_path), ['.pdf']).scan()
    record = PersonRecord(name, normalize_person_name(name))
    PersonMatcher([record], sources).match()
    assert sources[0].normalized_name == normalize_person_name(name)
    assert sources[0].matched_person is record
    assert sources[0].filename == filename

def test_normalize_person_name():
    assert normalize_person_name(" NGUYỄN   VĂN A ") == "nguyễn văn a"
    assert normalize_person_name("nguyễn văn a") == "nguyễn văn a"
    assert normalize_person_name("ĐỖ ĐỨC ĐẠT") == "đỗ đức đạt"
    assert normalize_person_name("Nguyễn Văn A (Người đại diện)") == "nguyễn văn a"

def test_extract_stt():
    stt, name = extract_stt_and_name("1. Nguyễn Văn A")
    assert stt == "1"
    assert name == "Nguyễn Văn A"
    
    stt, name = extract_stt_and_name("001 - Trần B")
    assert stt == "001"
    assert name == "Trần B"
    
    stt, name = extract_stt_and_name("Lê C")
    assert stt is None
    assert name == "Lê C"

def test_person_matcher():
    records = [
        PersonRecord(ho_ten="Nguyễn Văn A", normalized_name="nguyễn văn a"),
        PersonRecord(ho_ten="Trần B", normalized_name="trần b"),
    ]
    records[0].parcels.add(Parcel("70", "300"))
    
    sources = [
        SourceFile(original_path=Path("1. Nguyễn Văn A.pdf"), filename="1. Nguyễn Văn A.pdf", normalized_name="nguyễn văn a", extension=".pdf"),
        SourceFile(original_path=Path("Lê C.pdf"), filename="Lê C.pdf", normalized_name="lê c", extension=".pdf")
    ]
    
    matcher = PersonMatcher(records, sources)
    matcher.match()
    
    assert sources[0].matched_person == records[0]
    assert sources[1].matched_person is None

def test_person_matcher_ambiguous():
    records = [
        PersonRecord(ho_ten="A", normalized_name="a", secondary_key="1"),
        PersonRecord(ho_ten="A", normalized_name="a", secondary_key="2"),
    ]
    
    sources = [
        SourceFile(original_path=Path("A.pdf"), filename="A.pdf", normalized_name="a", extension=".pdf")
    ]
    
    matcher = PersonMatcher(records, sources)
    matcher.match()
    assert sources[0].matched_person is None
    assert sources[0].is_ambiguous == True


def test_person_matcher_allows_stt_drift_for_unique_name():
    record = PersonRecord(ho_ten="A", normalized_name="a", secondary_key="2")
    source = SourceFile(Path("1. A.pdf"), "1. A.pdf", "a", ".pdf", stt="1")

    PersonMatcher([record], [source]).match()

    assert source.matched_person is record


def test_person_matcher_ignores_stt_when_excel_name_is_unique():
    record = PersonRecord(ho_ten="A", normalized_name="a", secondary_key="2")
    sources = [
        SourceFile(Path("1. A.pdf"), "1. A.pdf", "a", ".pdf", stt="1"),
        SourceFile(Path("2. A.pdf"), "2. A.pdf", "a", ".pdf", stt="2"),
    ]

    PersonMatcher([record], sources).match()

    assert sources[0].matched_person is record
    assert sources[1].matched_person is record


def test_person_matcher_normalizes_stt_before_matching():
    record = PersonRecord(ho_ten="A", normalized_name="a", secondary_key="1")
    other = PersonRecord(ho_ten="A", normalized_name="a", secondary_key="2")
    source = SourceFile(Path("001. A.pdf"), "001. A.pdf", "a", ".pdf", stt="001")

    PersonMatcher([record, other], [source]).match()

    assert source.matched_person is record

def test_naming_engine():
    config = ProfileConfig(ma_dvhc="10930", prefix="CHUACOGIAY")
    engine = NamingEngine(config)
    
    class DummyAction:
        def __init__(self):
            self.parcel = Parcel("70", "300")
            self.suffix = "DDK"
            self.source_file = SourceFile(Path("A.pdf"), "A.pdf", "a", ".pdf")
            self.source_file.matched_person = PersonRecord("A", "a")
            self.source_file.stt = "1"
            
    name = engine.generate_filename(DummyAction())
    assert name == "CHUACOGIAY_10930_70_300-DDK.pdf"

def test_action_planner_reserves_multiple_conflict_paths(tmp_path):
    config = ProfileConfig(suffixes=["DDK"])
    
    sf1 = SourceFile(Path("A.pdf"), "A.pdf", "a", ".pdf")
    sf1.matched_person = PersonRecord("A", "a")
    sf1.matched_person.parcels.add(Parcel("70", "300"))
    
    sf2 = SourceFile(Path("B.pdf"), "B.pdf", "b", ".pdf")
    sf2.matched_person = PersonRecord("B", "b")
    # Cùng thửa (giả lập 2 người có thửa trùng nhau, dù đã bị loại ở bước đọc Excel nhưng ở Planner vẫn phải bắt conflict)
    sf2.matched_person.parcels.add(Parcel("70", "300"))

    sf3 = SourceFile(Path("C.pdf"), "C.pdf", "c", ".pdf")
    sf3.matched_person = PersonRecord("C", "c")
    sf3.matched_person.parcels.add(Parcel("70", "300"))
    
    planner = ActionPlanner([sf1, sf2, sf3], str(tmp_path), config)
    actions = planner.build_plan()
    
    assert len(actions) == 3
    statuses = [a.status for a in actions]
    assert statuses.count(ActionStatus.CONFLICT) == 3
    assert {a.conflict_path.name for a in actions if a.conflict_path} == {
        "CHUACOGIAY__70_300-DDK__CONFLICT_001.pdf",
        "CHUACOGIAY__70_300-DDK__CONFLICT_002.pdf",
        "CHUACOGIAY__70_300-DDK__CONFLICT_003.pdf",
    }


def test_excel_reader_keeps_single_header_behavior(tmp_path):
    path = tmp_path / "single-header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["STT", "Tên hộ", "Số tờ", "Số thửa"])
    sheet.append([1, "Nguyễn Văn A", 10, 20])
    workbook.save(path)

    reader = ExcelReader(str(path))
    assert reader.header_depth(sheet.title, 1) == 1
    assert reader.get_headers(sheet.title, 1)[:4] == ["STT", "Tên hộ", "Số tờ", "Số thửa"]
    records = reader.read_data(sheet.title, 1, {
        "ho_ten": "Tên hộ", "so_to": "Số tờ", "so_thua": "Số thửa"
    })
    assert records[0].raw_rows == [2]
    assert records[0].parcels == {Parcel("10", "20")}


def test_excel_reader_flattens_two_level_merged_headers_and_skips_both_rows(tmp_path):
    path = tmp_path / "two-level-header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells("A4:A5")
    sheet.merge_cells("B4:B5")
    sheet.merge_cells("C4:E4")
    sheet["A4"] = "STT"
    sheet["B4"] = "Tên hộ"
    sheet["C4"] = "BĐ 2004"
    sheet["C5"] = "Số thửa"
    sheet["D5"] = "Tờ BĐ"
    sheet["E5"] = "Diện tích"
    sheet["A6"] = 1
    sheet["B6"] = "Nguyễn Văn A"
    sheet["C6"] = 55
    sheet["D6"] = 12
    sheet["E6"] = 100
    workbook.save(path)

    reader = ExcelReader(str(path))
    assert reader.header_depth(sheet.title, 4) == 2
    assert reader.get_headers(sheet.title, 4)[:5] == [
        "STT", "Tên hộ", "BĐ 2004 / Số thửa", "BĐ 2004 / Tờ BĐ", "BĐ 2004 / Diện tích"
    ]
    records = reader.read_data(sheet.title, 4, {
        "ho_ten": "Tên hộ",
        "so_to": "BĐ 2004 / Tờ BĐ",
        "so_thua": "BĐ 2004 / Số thửa",
        "secondary_key": "STT",
    })
    assert records[0].raw_rows == [6]
    assert records[0].secondary_key == "1"
    assert records[0].parcels == {Parcel("12", "55")}


def test_excel_reader_resolves_exact_running_max_secondary_formula(tmp_path):
    path = tmp_path / "formula-stt.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["STT", "Tên hộ", "Số tờ", "Số thửa"])
    sheet.append([1, "Nguyễn Văn A", 10, 20])
    sheet.append(["=MAX($A$2:A2)+1", "Nguyễn Văn B", 11, 21])
    workbook.save(path)

    reader = ExcelReader(str(path))
    records = reader.read_data(sheet.title, 1, {
        "ho_ten": "Tên hộ", "so_to": "Số tờ", "so_thua": "Số thửa", "secondary_key": "STT"
    })

    assert [(record.ho_ten, record.secondary_key) for record in records] == [
        ("Nguyễn Văn A", "1"), ("Nguyễn Văn B", "2")
    ]


def test_one_household_with_three_merged_rows_keeps_all_parcels(tmp_path):
    path = tmp_path / "merged-household.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["STT", "Tên hộ", "Số tờ", "Số thửa"])
    sheet.merge_cells("A2:A4")
    sheet.merge_cells("B2:B4")
    sheet["A2"] = 1
    sheet["B2"] = "Nguyễn Văn A"
    sheet["C2"] = 1
    sheet["D2"] = 2
    sheet["C3"] = 3
    sheet["D3"] = 4
    sheet["C4"] = 5
    sheet["D4"] = 6
    workbook.save(path)

    reader = ExcelReader(str(path))
    records = reader.read_data(sheet.title, 1, {
        "ho_ten": "Tên hộ", "so_to": "Số tờ", "so_thua": "Số thửa", "secondary_key": "STT"
    })

    assert len(records) == 1
    assert records[0].parcels == {Parcel("1", "2"), Parcel("3", "4"), Parcel("5", "6")}


def test_one_household_with_blank_continuation_rows_keeps_all_parcels(tmp_path):
    path = tmp_path / "blank-continuation-household.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["STT", "Tên hộ", "Số tờ", "Số thửa"])
    sheet.append([1, "Nguyễn Văn A", 1, 2])
    sheet.append([None, None, 3, 4])
    sheet.append([None, None, 5, 6])
    sheet.append([None, "Tổng DT", None, None])
    workbook.save(path)

    records = ExcelReader(str(path)).read_data(sheet.title, 1, {
        "ho_ten": "Tên hộ", "so_to": "Số tờ", "so_thua": "Số thửa", "secondary_key": "STT"
    })

    assert len(records) == 1
    assert records[0].parcels == {Parcel("1", "2"), Parcel("3", "4"), Parcel("5", "6")}


def test_excel_reader_rejects_letters_and_non_positive_parcels_with_reasons(tmp_path):
    path = tmp_path / "invalid-parcels.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["STT", "Tên hộ", "Số tờ", "Số thửa"])
    sheet.append([1, "Nguyễn Văn A", 92, 182])
    sheet.append([None, None, 92, "CN"])
    sheet.append([None, None, 0, 183])
    sheet.append([None, None, 92, -1])
    workbook.save(path)

    records = ExcelReader(str(path)).read_data(sheet.title, 1, {
        "ho_ten": "Tên hộ", "so_to": "Số tờ", "so_thua": "Số thửa", "secondary_key": "STT"
    })

    assert records[0].parcels == {Parcel("92", "182")}
    assert len(records[0].data_issues) == 3
    assert all("hãy sửa dữ liệu nguồn" in issue for issue in records[0].data_issues)


def test_source_scanner_reads_nested_folders(tmp_path):
    nested = tmp_path / "hộ 01" / "giấy tờ"
    nested.mkdir(parents=True)
    (nested / "Nguyễn Văn A.pdf").write_bytes(b"pdf")

    sources = SourceScanner(str(tmp_path), [".pdf"]).scan()

    assert [source.filename for source in sources] == ["Nguyễn Văn A.pdf"]


def test_shared_parcel_is_an_explicit_warning_not_a_silent_drop(tmp_path):
    path = tmp_path / "shared.xlsx"
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["STT", "Tên hộ", "Số tờ", "Số thửa"])
    sheet.append([1, "Nguyễn Văn A", 92, 330])
    sheet.append([2, "Nguyễn Văn B", 92, 330])
    workbook.save(path)
    records = ExcelReader(str(path)).read_data(sheet.title, 1, {
        "ho_ten": "Tên hộ", "so_to": "Số tờ", "so_thua": "Số thửa", "secondary_key": "STT"
    })
    source_path = tmp_path / "Nguyễn Văn A.pdf"; source_path.write_bytes(b"pdf")
    source = SourceFile(source_path, source_path.name, "nguyễn văn a", ".pdf", matched_person=records[0])

    actions = ActionPlanner([source], str(tmp_path / "out"), ProfileConfig()).build_plan()

    assert records[0].parcels == set()
    assert len(actions) == 1 and actions[0].status == ActionStatus.WARNING
    assert "thuộc nhiều tên hộ" in actions[0].reason


def test_matcher_suggests_close_excel_name():
    record = PersonRecord("Nguyễn Bá Tài", "nguyễn bá tài", secondary_key="52")
    source = SourceFile(Path("52. Nguyễn Bá Tải.pdf"), "52. Nguyễn Bá Tải.pdf", "nguyễn bá tải", ".pdf", stt="52")

    PersonMatcher([record], [source]).match()

    assert source.matched_person is None
    assert "Nguyễn Bá Tài" in source.match_issue
    assert "STT: 52" in source.match_issue


def test_ambiguous_name_without_stt_lists_excel_candidates():
    records = [
        PersonRecord("Lê Thị Lan", "lê thị lan", secondary_key="37"),
        PersonRecord("Lê Thị Lan", "lê thị lan", secondary_key="93"),
    ]
    source = SourceFile(Path("Lê Thị Lan.pdf"), "Lê Thị Lan.pdf", "lê thị lan", ".pdf")

    PersonMatcher(records, [source]).match()

    assert source.is_ambiguous
    assert "37" in source.match_issue and "93" in source.match_issue


def test_excel_reader_skips_summary_rows(tmp_path):
    path = tmp_path / "summary-row.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["STT", "Tên hộ", "Số tờ", "Số thửa"])
    sheet.append([1, "Nguyễn Văn A", 1, 2])
    sheet.append([None, "Tổng DT", None, None])
    workbook.save(path)

    records = ExcelReader(str(path)).read_data(sheet.title, 1, {
        "ho_ten": "Tên hộ", "so_to": "Số tờ", "so_thua": "Số thửa"
    })

    assert [record.ho_ten for record in records] == ["Nguyễn Văn A"]


@pytest.mark.parametrize("label", [
    "Thiếu HS",
    "Thiếu hồ sơ",
    "Sai diện tích",
    "Không khớp diện tích",
    "Diện tích đất không khớp với tờ khai",
])
def test_excel_reader_skips_non_person_note_rows(tmp_path, label):
    path = tmp_path / "non-person-label.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["STT", "Tên hộ", "Số tờ", "Số thửa"])
    sheet.append([1, "Nguyễn Văn A", 1, 2])
    sheet.append([None, label, 3, 4])
    workbook.save(path)

    records = ExcelReader(str(path)).read_data(sheet.title, 1, {
        "ho_ten": "Tên hộ", "so_to": "Số tờ", "so_thua": "Số thửa", "secondary_key": "STT"
    })

    assert [record.ho_ten for record in records] == ["Nguyễn Văn A"]
    assert records[0].parcels == {Parcel("1", "2")}


def test_one_pdf_creates_one_output_per_parcel(tmp_path):
    source_path = tmp_path / "Nguyễn Văn A.pdf"
    source_path.write_bytes(b"source-pdf")
    person = PersonRecord("Nguyễn Văn A", "nguyễn văn a")
    person.parcels.update({Parcel("1", "2"), Parcel("3", "4"), Parcel("5", "6")})
    source = SourceFile(source_path, source_path.name, "nguyễn văn a", ".pdf", matched_person=person)

    actions = ActionPlanner([source], str(tmp_path / "out"), ProfileConfig(ma_dvhc="mdvhc", suffixes=["DDK"])).build_plan()

    assert len(actions) == 3
    assert {a.target_filename for a in actions} == {
        "CHUACOGIAY_mdvhc_1_2-DDK.pdf",
        "CHUACOGIAY_mdvhc_3_4-DDK.pdf",
        "CHUACOGIAY_mdvhc_5_6-DDK.pdf",
    }
