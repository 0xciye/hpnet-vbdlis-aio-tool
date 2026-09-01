from collections import defaultdict
from contextlib import ExitStack
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata
from openpyxl import load_workbook
from openpyxl.cell.cell import ERROR_CODES
from openpyxl.utils import get_column_letter, column_index_from_string
from .models import ColumnMapping, Inspection, NoticeRecord, UserError, file_hash
from .fields import has_content


def clean(value):
    return "" if value is None else " ".join(str(value).split())


def folded(value):
    return "".join(c for c in unicodedata.normalize("NFD", clean(value).casefold().replace("đ", "d"))
                   if not unicodedata.combining(c))


def is_summary(value):
    # Only complete labels, never a prefix that could match a person's name.
    return folded(value).rstrip(" :.;…：") in {"tong dt", "tong cong", "cong", "tong dien tich"}


def cell_problem(source, cached, address, label):
    errors = {*ERROR_CODES, "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!", "#UNKNOWN!", "#CONNECT!"}
    if source.data_type == "e" or cached.data_type == "e" or clean(cached.value).upper() in errors:
        return (f"Dữ liệu {label} không hợp lệ tại {address}: ô Excel đang báo lỗi «{clean(cached.value) or clean(source.value)}». "
                "Mở Excel, sửa lỗi tại ô này rồi lưu file và kiểm tra lại.")
    if source.data_type == "f" and not clean(cached.value):
        return (f"Dữ liệu {label} không hợp lệ tại {address}: ô chứa công thức nhưng chưa có kết quả để đọc. "
                "Mở Excel, tính lại công thức và lưu file; hoặc nhập giá trị đúng vào ô này rồi kiểm tra lại.")
    return ""


def identifier(value):
    text = clean(value)
    if not text:
        return ""
    if not re.fullmatch(r"[0-9]+(?:\.0+)?", text):
        raise ValueError("Số tờ/thửa phải là số nguyên dương.")
    number = Decimal(text)
    if number <= 0 or number > 2147483647:
        raise ValueError("Số tờ/thửa phải là số nguyên dương.")
    return str(int(number))


def area_text(value):
    text = clean(value)
    if not text:
        return ""
    if not re.fullmatch(r"[0-9]+(?:[.,][0-9]+)?", text):
        raise ValueError("Diện tích phải là số dương, không chứa lỗi công thức hay đơn vị.")
    try:
        number = Decimal(text.replace(",", "."))
    except InvalidOperation:
        raise ValueError("Diện tích không phải số.") from None
    if number <= 0:
        raise ValueError("Diện tích phải lớn hơn 0.")
    return format(number, "f").rstrip("0").rstrip(".") if "." in format(number, "f") else format(number, "f")


def optional_excel_row(value, label):
    if value is None or str(value).strip() == "":
        return None
    text = str(value).strip()
    if not re.fullmatch(r"[0-9]+", text) or int(text) < 1:
        raise UserError(f"{label} phải là số dòng Excel nguyên dương.")
    return int(text)


def workbook_info(path, sheet_name=None, header_row=None, depth=2):
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        names = wb.sheetnames
        if sheet_name and sheet_name not in names:
            raise UserError("Trang tính đã chọn không còn trong file Excel.")
        ws = wb[sheet_name or names[0]]
        if header_row is None:
            candidates = []
            for number in range(1, min(ws.max_row, 30) + 1):
                texts = [folded(c.value) for c in ws[number]]
                score = sum(any(key in t for key in ("ten ho", "ho ten", "to bd", "thua", "dien tich", "dt bd")) for t in texts)
                candidates.append((score, -number))
            header_row = -max(candidates)[1]
            depth = 2 if any(m.min_row == header_row and m.max_row == header_row + 1 for m in ws.merged_cells.ranges) else 1
        if not 1 <= header_row <= ws.max_row or depth not in (1, 2) or header_row + depth > ws.max_row + 1:
            raise UserError("Dòng tiêu đề/số tầng không nằm trong dữ liệu của trang tính.")
        def header_value(row, col):
            cell = ws.cell(row, col)
            if cell.value is not None:
                return clean(cell.value)
            for merged in ws.merged_cells.ranges:
                if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                    return clean(ws.cell(merged.min_row, merged.min_col).value)
            return ""
        headers = {}
        for col in range(1, ws.max_column + 1):
            labels = list(dict.fromkeys(header_value(r, col) for r in range(header_row, header_row + depth)))
            headers[get_column_letter(col)] = " / ".join(v for v in labels if v) or "[Không có tiêu đề]"
        suggestions = {}
        for field_name, aliases in {"household_index": ("stt", "so tt", "stt ho", "so thu tu"),
                                   "owner": ("ten ho", "ho va ten", "ho ten"), "sheet": ("to bd moi", "to ban do moi"),
                                   "parcel": ("thua bd moi", "thua ban do moi"), "area": ("dt bd", "dien tich ban do"),
                                   "location": ("xu dong", "vi tri",), "identity": ("cccd", "cmnd", "giay to nhan than", "so dinh danh")}.items():
            matches = [col for col, label in headers.items() if any(a in folded(label) for a in aliases)]
            suggestions[field_name] = matches[0] if len(matches) == 1 else ""
        return {"source": str(Path(path).resolve()), "sheets": names, "sheet": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                "header_row": header_row, "depth": depth, "headers": headers, "suggestions": suggestions}
    except (OSError, ValueError) as exc:
        if isinstance(exc, UserError):
            raise
        raise UserError("Không đọc được cấu trúc Excel. Kiểm tra file .xlsx và trang tính đã chọn.") from exc
    finally:
        wb.close()


def inspect_workbook(path, sheet_name, header_row, depth, mapping, *, require_identity=False,
                     start_row=None, end_row=None):
    path = Path(path).resolve()
    digest = file_hash(path)
    info = workbook_info(path, sheet_name, header_row, depth)
    mapping.validate(info["columns"])
    records, blanks, summaries, name_only = [], [], [], []
    groups = defaultdict(list)
    current_household = ""; household_problem = ""
    current_owner = ""; owner_row = None; owner_problem = ""
    current_identity = ""; identity_row = None; identity_problem = ""
    with ExitStack() as stack:
        wb = load_workbook(path, read_only=True, data_only=True)
        stack.callback(wb.close)
        formulas = load_workbook(path, read_only=True, data_only=False)
        stack.callback(formulas.close)
        ws = wb[sheet_name]
        source_ws = formulas[sheet_name]
        indices = {k: column_index_from_string(v)-1 for k,v in vars_mapping(mapping).items() if v}
        first = header_row + depth
        requested_start = optional_excel_row(start_row, "Dòng bắt đầu")
        requested_end = optional_excel_row(end_row, "Dòng kết thúc")
        if (requested_start is None) != (requested_end is None):
            raise UserError("Hãy nhập đủ cả dòng bắt đầu và dòng kết thúc, hoặc để trống cả hai để xử lý toàn bộ file.")
        row_options = {"min_row": first}
        if requested_start is not None:
            if requested_start > requested_end:
                raise UserError("Dòng bắt đầu phải nhỏ hơn hoặc bằng dòng kết thúc.")
            if requested_start < first:
                raise UserError(f"Dòng bắt đầu phải từ dòng {first} trở đi, sau phần tiêu đề.")
            last_data_row = first - 1
            for number, source_cells in enumerate(source_ws.iter_rows(min_row=first), first):
                if any(clean(cell.value) for cell in source_cells):
                    last_data_row = number
            if last_data_row < first:
                raise UserError("Trang tính không có dòng dữ liệu nào sau phần tiêu đề.")
            if requested_start > last_data_row:
                raise UserError(f"Dòng bắt đầu {requested_start} vượt quá dòng cuối có dữ liệu ({last_data_row}).")
            if requested_end > last_data_row:
                raise UserError(f"Dòng kết thúc {requested_end} vượt quá dòng cuối có dữ liệu ({last_data_row}).")
            # Scan from the first data row to preserve the owning household when a selected
            # range begins on a parcel/member row. Only selected rows become output records.
            row_options = {"min_row": first, "max_row": requested_end}
        for number, (cells, source_cells) in enumerate(zip(ws.iter_rows(**row_options), source_ws.iter_rows(**row_options), strict=True), row_options["min_row"]):
            selected = requested_start is None or number >= requested_start
            if not any(clean(c.value) for c in source_cells):
                if selected: blanks.append(number)
                continue
            def value(key):
                return cells[indices[key]].value if key in indices else None
            def problem(key, label):
                if key not in indices:
                    return ""
                return cell_problem(source_cells[indices[key]], cells[indices[key]], f"{getattr(mapping,key)}{number}", label)
            owner = clean(value("owner"))
            if is_summary(owner):
                if selected: summaries.append(number)
                continue
            marker_issue = problem("household_index", "STT hộ")
            marker = clean(value("household_index"))
            starts_household = False
            if marker_issue:
                current_household, household_problem = "", marker_issue
                current_owner, owner_row, owner_problem = "", number, ""
                current_identity, identity_row, identity_problem = "", None, ""
            elif marker:
                try:
                    current_household = identifier(marker)
                except ValueError:
                    current_household = ""
                    household_problem = f"Dữ liệu STT hộ không hợp lệ tại {mapping.household_index}{number}: «{marker}». STT phải là số nguyên dương."
                    current_owner, owner_row, owner_problem = "", number, ""
                    current_identity, identity_row, identity_problem = "", None, ""
                else:
                    starts_household = True; household_problem = ""
                    issue = problem("owner", "tên hộ")
                    if owner and not has_content(owner) and not issue:
                        issue = f"Dữ liệu tên hộ không hợp lệ tại {mapping.owner}{number}: chỉ có dấu/khoảng trống, chưa có họ tên. Hãy nhập họ tên đúng vào ô nguồn."
                    if issue:
                        current_owner, owner_row, owner_problem = "", number, issue
                    elif not owner:
                        current_owner, owner_row = "", number
                        owner_problem = f"Thiếu tên chủ hộ tại {mapping.owner}{number}, cùng dòng với STT hộ {current_household}."
                    else:
                        current_owner, owner_row, owner_problem = owner, number, ""
                    current_identity, identity_row, identity_problem = "", None, ""
                    identity_issue = problem("identity", "giấy tờ nhân thân")
                    raw_identity = value("identity")
                    if identity_issue:
                        current_identity, identity_row, identity_problem = "", number, identity_issue
                    elif clean(raw_identity) and not owner_problem:
                        # Only the identity on the STT row belongs to the household head.
                        current_identity = clean(raw_identity)
                        if isinstance(raw_identity, (int, float)) and not isinstance(raw_identity, bool) and raw_identity == int(raw_identity):
                            current_identity = str(int(raw_identity))
                            fmt = source_cells[indices["identity"]].number_format
                            if re.fullmatch(r"0+", fmt): current_identity = current_identity.zfill(len(fmt))
                        identity_row, identity_problem = number, ""
            if not selected:
                continue
            if not any(clean(source_cells[indices[k]].value) for k in ("sheet", "parcel", "area")):
                name_only.append(number); continue
            errors = []
            if household_problem:
                errors.append(household_problem + " Không xác định được chủ hộ; không kế thừa tên của hộ phía trên.")
            elif not current_household:
                errors.append(f"Thiếu STT hộ tại {mapping.household_index}{number}; chưa xác định được dòng chủ hộ gần nhất phía trên.")
            if identity_problem:
                errors.append(identity_problem)
            elif require_identity and not has_content(current_identity):
                address = f"{mapping.identity}{owner_row or number}" if mapping.identity else "bước 3 (Đối chiếu cột)"
                errors.append(f"Thiếu giấy tờ nhân thân tại {address}. Chọn đúng cột và bổ sung giấy tờ của hộ này; không lấy giấy tờ của hộ trước.")
            if owner_problem:
                errors.append(owner_problem + (f" Dòng {number} không được kế thừa tên hộ trước ô lỗi này." if number != owner_row else " Không dùng tên hộ phía trên để thay thế."))
            elif not current_owner:
                errors.append(f"Thiếu tên hộ tại {mapping.owner}{number}; chưa có tên phía trên để kế thừa.")
            normalized = {}
            for key, label, convert in (("sheet", "tờ BĐ mới", identifier), ("parcel", "thửa BĐ mới", identifier), ("area", "diện tích", area_text)):
                raw = value(key)
                issue = problem(key, label)
                if issue:
                    normalized[key] = ""; errors.append(issue); continue
                try:
                    normalized[key] = convert(raw)
                    if not normalized[key]:
                        errors.append(f"Thiếu {label} tại {getattr(mapping,key)}{number}.")
                except ValueError:
                    normalized[key] = ""
                    errors.append(f"Dữ liệu {label} không hợp lệ tại {getattr(mapping,key)}{number}: «{clean(raw)}». Hãy kiểm tra ô nguồn.")
            location_issue = problem("location", "xứ đồng")
            if location_issue: errors.append(location_issue)
            record = NoticeRecord(number, current_owner, owner_row, normalized["sheet"], normalized["parcel"],
                                  normalized["area"], "" if location_issue else clean(value("location")), errors,
                                  identity=current_identity, identity_row=identity_row, household_number=current_household)
            records.append(record)
            if record.sheet and record.parcel:
                groups[(record.sheet, record.parcel)].append(record)
        for group in groups.values():
            if len(group) > 1:
                for record in group:
                    record.duplicate_rows = [r.source_row for r in group]
    if file_hash(path) != digest:
        raise UserError("File Excel đã thay đổi trong khi đọc. Hãy kiểm tra lại dữ liệu.")
    return Inspection(path, sheet_name, digest, records, info["rows"], blanks, summaries, name_only,
                      header_row, depth, mapping, requested_start, requested_end)


def vars_mapping(mapping):
    from dataclasses import asdict
    return asdict(mapping)
