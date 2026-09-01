from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import json
import uuid


@dataclass
class LogEntry:
    row: int
    owner: str
    sheet: str
    parcel: str
    number: int | None
    status: str
    error_type: str
    detail: str
    filename: str = ""


class RunLog:
    def __init__(self, output, inspection):
        directory = Path(output) / "Nhat_ky"
        directory.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S") + "_" + uuid.uuid4().hex[:8]
        self.txt = directory / f"Log_Tao_Thong_Bao_{stamp}.txt"
        self.xlsx = self.txt.with_suffix(".xlsx")
        self.entries = []
        self.inspection = inspection
        self.handle = self.txt.open("x", encoding="utf-8-sig")
        selected = (f"Phạm vi xử lý: dòng {inspection.selected_start_row}–{inspection.selected_end_row} (bao gồm hai đầu)\n"
                    if inspection.selected_start_row is not None else "")
        self.handle.write(f"NHẬT KÝ TẠO THÔNG BÁO ĐẤT ĐAI\nNguồn: {inspection.source}\nTrang tính: {inspection.sheet_name}\n"
                          f"Dòng tiêu đề: {inspection.header_row}; số tầng: {inspection.header_depth}\n"
                          f"{selected}"
                          f"Dòng thửa: {len(inspection.records)}; hợp lệ: {len(inspection.valid_records)}; "
                          f"dòng tổng bỏ qua: {len(inspection.summary_rows)}; dòng trống: {len(inspection.blank_rows)}\n"
                          "Tìm ô lỗi: mở đúng trang tính Excel, nhấn Ctrl+G rồi nhập địa chỉ ô ghi trong chi tiết.\n"
                          "Số chỉ được dùng khi file được tạo thành công. Khi lỗi, số chưa dùng được giữ cho thửa tiếp theo.\n\n")
        self.handle.flush()

    def append(self, entry):
        self.entries.append(entry)
        self.handle.write(f"[{entry.status}] Dòng Excel {entry.row} | Hộ: {entry.owner or '[CHƯA CÓ TÊN]'} | "
                          f"Tờ: {entry.sheet} | Thửa: {entry.parcel} | Số thông báo: {entry.number if entry.number is not None else '[CHƯA CẤP]'}\n"
                          f"{entry.detail}\nFile: {entry.filename or '[KHÔNG TẠO]'}\n\n")
        self.handle.flush()

    def finish(self):
        self.handle.close()
        wb = Workbook()
        ws = wb.active; ws.title = "Kết quả từng dòng"
        ws.append(["STT", "Tên hộ", "Tờ", "Thửa", "Dòng Excel", "Số thông báo", "Trạng thái", "Loại lỗi", "Chi tiết", "Tên file"])
        for index, entry in enumerate(self.entries, 1):
            ws.append([index, entry.owner, entry.sheet, entry.parcel, entry.row, entry.number,
                       entry.status, entry.error_type, entry.detail, entry.filename])
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # User data must never become an Excel formula.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163C51")
        for col, width in zip("ABCDEFGHIJ", [8,28,10,10,13,16,24,24,85,60]):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        meta = wb.create_sheet("Đối chiếu nguồn")
        meta.append(["Mục", "Giá trị"])
        range_text = (f"{self.inspection.selected_start_row}–{self.inspection.selected_end_row}"
                      if self.inspection.selected_start_row is not None else "Toàn bộ dữ liệu")
        for label, value in [("File",str(self.inspection.source)),("Trang tính",self.inspection.sheet_name),
                             ("Phạm vi dòng xử lý",range_text),
                             ("SHA-256",self.inspection.source_hash),
                             ("Dòng tổng bỏ qua",", ".join(map(str,self.inspection.summary_rows))),
                             ("Dòng trống bỏ qua",", ".join(map(str,self.inspection.blank_rows))),
                             ("Dòng không có thửa",", ".join(map(str,self.inspection.name_only_rows)))]:
            meta.append([label,value])
        meta.column_dimensions["A"].width = 26; meta.column_dimensions["B"].width = 100
        for row in meta:
            for cell in row:
                if isinstance(cell.value,str): cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        wb.save(self.xlsx); wb.close()


def export_inspection(inspection,output):
    """Available even when every record is invalid; never allocates a number."""
    from .models import file_hash,UserError
    if file_hash(inspection.source)!=inspection.source_hash:
        raise UserError("Excel đã thay đổi. Hãy kiểm tra lại trước khi xuất báo cáo.")
    log=RunLog(output,inspection)
    try:
        for record in inspection.records:
            detail=" ".join(record.errors)
            if record.duplicate_rows:
                detail+=" Trùng tờ/thửa tại dòng "+", ".join(map(str,record.duplicate_rows))+"; toàn bộ nhóm không được tạo."
            if record.valid:
                detail=f"Đủ dữ liệu nguồn; tên lấy từ dòng {record.owner_row}. Chưa tạo Word hoặc cấp số."
            else:
                detail+=" Không tạo file, không cấp số."
            log.append(LogEntry(record.source_row,record.owner,record.sheet,record.parcel,None,record.status,
                                "" if record.valid else record.status,detail))
        log.finish()
    finally:
        if not log.handle.closed: log.handle.close()
    return {"txt":log.txt,"xlsx":log.xlsx}
