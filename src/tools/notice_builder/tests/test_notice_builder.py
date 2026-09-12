from dataclasses import asdict, replace
from datetime import date
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from xml.dom import minidom
import json
import os
import hashlib
import pytest
from openpyxl import Workbook, load_workbook
from tools.notice_builder.paths import resource, default_template_path, default_template_fields
from tools.notice_builder.core import (BatchConfig, ColumnMapping, NoticeService, NumberPool,
    UserError, WordTemplate, inspect_workbook, workbook_info)
from tools.notice_builder.core.models import NoticeRecord, file_hash
from tools.notice_builder.core.numbering import parse_numbers
from tools.notice_builder.core.reader import area_text, identifier, is_summary
from tools.notice_builder.core.renderer import WORD_NS, paragraph_nodes, node_text
from tools.notice_builder.core import service as service_module
from tools.notice_builder.template_config import default_template_config, template_configs, template_config_for_path


@pytest.fixture
def config():
    return BatchConfig("12345","Thôn thử, xã thử","Thôn thử","Xã thử","Xã thử, thành phố thử","Xã thử",31,8,2026,start_number=100,template_fields=default_template_fields())


@pytest.fixture
def service():
    return NoticeService(default_template_path(),
                         json.loads(resource("config/legal_defaults.json").read_text(encoding="utf-8")))


def workbook(path, rows, header=1, other=False):
    wb=Workbook(); ws=wb.active; ws.title="Dữ liệu"
    if other: wb.create_sheet("Không chọn").append(["DỮ LIỆU KHÔNG ĐƯỢC ĐỌC"])
    for col,title in {1:"STT",2:"Tên hộ",7:"Tờ BĐ mới",8:"Thửa BĐ mới",11:"DT BĐ (m2)",12:"Vị trí/Xứ đồng",13:"Diện tích giao"}.items():
        ws.cell(header,col,title)
    household_number=0
    for index,row in enumerate(rows,header+1):
        if row[0] is not None and not is_summary(row[0]):
            household_number+=1; ws.cell(index,1,household_number)
        for col,value in zip((2,7,8,11,12),row):
            cell=ws.cell(index,col,value)
            if isinstance(value,str) and value.startswith("="): cell.data_type="s"
        ws.cell(index,4,"000000000001")  # Synthetic identity; leading zero must be preserved.
    wb.save(path); wb.close()
    return inspect_workbook(path,"Dữ liệu",header,1,ColumnMapping(identity="D"),require_identity=True)


def ready(tmp_path,service,config,rows=None):
    data=workbook(tmp_path/"source.xlsx", rows or [("HỘ THỬ A",72,175,357,None),(None,72,176,120.25,"Đồng thử"),("HỘ THỬ B",72,177,500,None)])
    output=tmp_path/"output"
    preview=service.preview(data,config,output,0,tmp_path/"preview")
    return data,output,service.approve(preview)


def all_text(payload):
    with ZipFile(BytesIO(payload)) as z:
        doc=minidom.parseString(z.read("word/document.xml"))
    return "\n".join("".join(node_text(n) for n in paragraph_nodes(p)) for p in doc.getElementsByTagNameNS(WORD_NS,"p"))


@pytest.mark.parametrize("text,expected",[("1",[1]),("1,3,5",[1,3,5]),("5-7",[5,6,7]),("1,3,5-13",[1,3,*range(5,14)]),(" 1 , 3 - 5 ",[1,3,4,5])])
def test_number_list(text,expected): assert parse_numbers(text)==expected


@pytest.mark.parametrize("text",[""," ","13-5","abc","1,,3","1,3,3,5","0","-1","1-3,3","1,","2.5","1-100001","2147483648"])
def test_bad_number_list(text):
    with pytest.raises(UserError): parse_numbers(text)


def test_peek_commit_and_continuation():
    pool=NumberPool("list",text="1,3,5-13",continuation=15)
    assert pool.preview_number(11)==15 and pool.peek()==1
    actual=[]
    for _ in range(14):
        n=pool.peek(); assert pool.peek()==n; actual.append(n); pool.commit(n)
    assert actual==[1,3,*range(5,14),15,16,17]
    with pytest.raises(UserError): pool.commit(99)
    with pytest.raises(UserError): NumberPool("list",text="1,3,5",continuation=3)
    finite=NumberPool("list",text="1"); finite.commit(1); assert finite.peek() is None


def test_number_date_groups_override_default_date(config):
    cfg=replace(config,number_mode="list",number_list="300-400",number_date_rules=[
        {"numbers":"300-350","date":"2026-08-25"},
        {"numbers":"351-400","date":"2026-08-28"},
    ])
    pool=NumberPool.from_config(cfg); default=date(2026,8,31)
    assert pool.date_for(300,default)==pool.date_for(350,default)==date(2026,8,25)
    assert pool.date_for(351,default)==pool.date_for(400,default)==date(2026,8,28)
    assert pool.date_for(401,default)==default


def test_blank_number_list_uses_sparse_date_groups_as_ordered_number_queue(config):
    cfg=replace(config,number_mode="list",number_list="",number_date_rules=[
        {"numbers":"397-398,401-410","date":"2026-08-22"},
        {"numbers":"543","date":"2026-08-25"},
        {"numbers":"735,777-778,799","date":"2026-08-28"},
    ])
    pool=NumberPool.from_config(cfg); default=date(2026,8,31)
    expected=[397,398,*range(401,411),543,735,777,778,799]
    actual=[]
    while pool.peek() is not None:
        number=pool.peek(); actual.append(number)
        assert pool.date_for(number,default) in {date(2026,8,22),date(2026,8,25),date(2026,8,28)}
        pool.commit(number)
    assert actual==expected
    assert pool.date_for(397,default)==date(2026,8,22)
    assert pool.date_for(543,default)==date(2026,8,25)
    assert pool.date_for(799,default)==date(2026,8,28)


def test_blank_number_list_group_queue_can_continue_after_highest_number(config):
    cfg=replace(config,number_mode="list",number_list="",continue_number=800,number_date_rules=[
        {"numbers":"397,543,799","date":"2026-08-25"},
    ])
    pool=NumberPool.from_config(cfg); default=date(2026,8,31)
    assert [pool.preview_number(index) for index in range(5)]==[397,543,799,800,801]
    assert pool.date_for(799,default)==date(2026,8,25)
    assert pool.date_for(800,default)==default


def test_blank_number_list_requires_at_least_one_number_date_group(config):
    cfg=replace(config,number_mode="list",number_list="",number_date_rules=[])
    with pytest.raises(UserError,match="Danh sách số hoặc thêm ít nhất một nhóm"):
        NumberPool.from_config(cfg)


@pytest.mark.parametrize("rules,message",[
    ([{"numbers":"300-350","date":"2026-08-25"},{"numbers":"350-360","date":"2026-08-28"}],"nhiều nhóm ngày"),
    ([{"numbers":"299-300","date":"2026-08-25"}],"không nằm trong Danh sách"),
    ([{"numbers":"300-350","date":"25/08/2026"}],"không hợp lệ"),
    ([{"numbers":"","date":"2026-08-25"}],"nhập đủ"),
])
def test_invalid_number_date_groups(config,rules,message):
    cfg=replace(config,number_mode="list",number_list="300-400",number_date_rules=rules)
    with pytest.raises(UserError,match=message): NumberPool.from_config(cfg)


def test_render_uses_date_for_each_notice_number(tmp_path,service,config):
    record=workbook(tmp_path/"date-rules.xlsx",[("HỘ A",1,1,100,None)]).records[0]
    cfg=replace(config,number_mode="list",number_list="300,351",number_date_rules=[
        {"numbers":"300","date":"2026-08-25"},{"numbers":"351","date":"2026-08-28"}])
    pool=NumberPool.from_config(cfg)
    first=all_text(service.template.render(service.values(record,cfg,300,pool.date_for(300,date(2026,8,31)))))
    second=all_text(service.template.render(service.values(record,cfg,351,pool.date_for(351,date(2026,8,31)))))
    assert "ngày 25 tháng 08 năm 2026" in first
    assert "ngày 28 tháng 08 năm 2026" in second


@pytest.mark.parametrize("value,expected",[(72,"72"),(72.0,"72"),(" 072.00 ","72")])
def test_identifier(value,expected): assert identifier(value)==expected


@pytest.mark.parametrize("value",[0,-1,"12a","CN","#REF!","72.5",True,"../CN"])
def test_bad_identifier(value):
    with pytest.raises(ValueError): identifier(value)


@pytest.mark.parametrize("value,expected",[(357,"357"),(1000,"1000"),("120,2500","120.25"),("0.0001","0.0001"),(None,"")])
def test_area(value,expected): assert area_text(value)==expected


@pytest.mark.parametrize("value",[0,-1,"#REF!","5m2","NaN","1.000,50"])
def test_bad_area(value):
    with pytest.raises(ValueError): area_text(value)


def test_headers_grouping_summary_and_other_sheet(tmp_path):
    path=tmp_path/"grouped.xlsx"
    data=workbook(path,[(None,1,1,20,None),(" HỘ A ",None,None,None,None),(None,72,175,357,None),
        ("Tổng DT",72,175,"#REF!",None),(None,72,176,120.25,None),("TỔNG CỘNG",1,1,1,None),
        ("Cộng",1,1,1,None),("HỘ B",None,2,100,None),(None,1,None,100,None),(None,1,2,None,None)],header=6,other=True)
    info=workbook_info(path,"Dữ liệu")
    assert info["header_row"]==6 and len(info["sheets"])==2
    assert info["suggestions"]=={"household_index":"A","owner":"B","sheet":"G","parcel":"H","area":"K","location":"L","identity":"","birth_date":""}
    assert len(data.records)==6 and data.records[0].status=="THIẾU DỮ LIỆU"
    assert data.records[1].owner==data.records[2].owner=="HỘ A"
    assert data.records[1].owner_row==8 and len(data.summary_rows)==3
    assert any("G14" in e for e in data.records[3].errors)
    assert any("H15" in e for e in data.records[4].errors)
    assert any("K16" in e for e in data.records[5].errors)


def test_changed_mapping_and_two_tier(tmp_path):
    path=tmp_path/"two.xlsx"; wb=Workbook(); ws=wb.active; ws.title="Nguồn"
    ws.append(["Bảng thống kê"]); ws.append(["Tên hộ","Tờ BĐ mới","Thửa BĐ mới","Diện tích bản đồ"])
    for col in "ABC": ws.merge_cells(f"{col}2:{col}3")
    ws["D3"]="m2"; ws["E2"]="STT hộ"; ws.merge_cells("E2:E3")
    ws.append(["HỘ KHÁC",43,120,200,1]); ws.append([None,None,None,None,None]); ws.append([None,43,121,201,None])
    wb.save(path); wb.close()
    info=workbook_info(path)
    assert info["header_row"]==2 and info["depth"]==2
    data=inspect_workbook(path,"Nguồn",2,2,ColumnMapping("A","B","C","D","",household_index="E"))
    assert len(data.valid_records)==2 and data.blank_rows==[5]
    assert data.records[1].owner=="HỘ KHÁC" and data.records[1].location==""


def test_optional_excel_row_range_and_blank_backward_compatibility(tmp_path):
    path=tmp_path/"range.xlsx"; wb=Workbook(); ws=wb.active; ws.title="Nguồn"
    ws.append(["Tên hộ","Tờ BĐ mới","Thửa BĐ mới","Diện tích bản đồ","Giấy tờ nhân thân","STT hộ"])
    ws.append(["HỘ CŨ",1,1,100,"GIẤY CŨ",1]); ws.append([None,1,2,110,None,None])
    ws.append(["HỘ MỚI",2,3,120,"GIẤY MỚI",2]); ws.append([None,2,4,130,None,None])
    wb.save(path); wb.close()
    mapping=ColumnMapping("A","B","C","D","","E",household_index="F")
    full=inspect_workbook(path,"Nguồn",1,1,mapping,require_identity=True)
    blank=inspect_workbook(path,"Nguồn",1,1,mapping,require_identity=True,start_row="",end_row="")
    selected=inspect_workbook(path,"Nguồn",1,1,mapping,require_identity=True,start_row=4,end_row=5)
    assert [asdict(record) for record in blank.records]==[asdict(record) for record in full.records]
    assert blank.signature()==full.signature() and blank.selected_start_row is None
    assert [record.source_row for record in selected.records]==[4,5]
    assert [record.owner for record in selected.records]==["HỘ MỚI","HỘ MỚI"]
    assert all(record.identity=="GIẤY MỚI" for record in selected.records)
    assert (selected.selected_start_row,selected.selected_end_row)==(4,5)


@pytest.mark.parametrize("start,end,message",[
    (5,4,"nhỏ hơn hoặc bằng"),
    (1,2,"sau phần tiêu đề"),
    (2,None,"nhập đủ cả dòng bắt đầu"),
    (2,99,"vượt quá dòng cuối có dữ liệu"),
])
def test_invalid_excel_row_range(tmp_path,start,end,message):
    path=tmp_path/"bad-range.xlsx"; wb=Workbook(); ws=wb.active; ws.title="Nguồn"
    ws.append(["Tên hộ","Tờ","Thửa","Diện tích","Giấy tờ","STT hộ"]); ws.append(["HỘ A",1,1,100,"GIẤY A",1])
    wb.save(path); wb.close()
    with pytest.raises(UserError,match=message):
        inspect_workbook(path,"Nguồn",1,1,ColumnMapping("A","B","C","D","","E",household_index="F"),
                         require_identity=True,start_row=start,end_row=end)


@pytest.mark.parametrize("same_owner,count",[(False,2),(False,3),(True,2)])
def test_all_duplicate_members_blocked(tmp_path,same_owner,count):
    rows=[("HỘ A",72,175,357,None),("HỘ X",72,176,20,None)]
    rows.extend([("HỘ A" if same_owner else f"HỘ {i}","072", "175.0",10,None) for i in range(count-1)])
    data=workbook(tmp_path/"dup.xlsx",rows)
    assert len(data.valid_records)==1
    blocked=[r for r in data.records if r.duplicate_rows]
    assert len(blocked)==count and all(len(r.duplicate_rows)==count for r in blocked)


def test_repeated_parcel_on_member_rows_creates_one_notice(tmp_path):
    path=tmp_path/"members-repeat.xlsx"; wb=Workbook(); ws=wb.active; ws.title="Nguồn"
    ws.append(["STT hộ","Tên hộ","Tờ","Thửa","Diện tích","Giấy tờ"])
    ws.append([1,"NGUYỄN VĂN CHỦ",92,330,100,"030064013684"])
    ws.append([None,"NGUYỄN THỊ THÀNH VIÊN",92,330,100,"030200012345"])
    ws.append([None,"NGUYỄN VĂN THÀNH VIÊN",92,330,100,"030202067890"])
    wb.save(path); wb.close()

    data=inspect_workbook(path,"Nguồn",1,1,ColumnMapping(owner="B",sheet="C",parcel="D",area="E",location="",identity="F",household_index="A"),require_identity=True)

    assert len(data.records)==1 and len(data.valid_records)==1
    assert [person.name for person in data.records[0].household_people]==[
        "NGUYỄN VĂN CHỦ","NGUYỄN THỊ THÀNH VIÊN","NGUYỄN VĂN THÀNH VIÊN"
    ]


def test_three_members_repeated_for_two_parcels_create_two_notices(tmp_path):
    path=tmp_path/"members-two-parcels.xlsx"; wb=Workbook(); ws=wb.active; ws.title="Nguồn"
    ws.append(["STT hộ","Tên hộ","Tờ","Thửa","Diện tích","Giấy tờ"])
    people=[("NGUYỄN VĂN CHỦ","030064013684"),("NGUYỄN THỊ B","030200012345"),("NGUYỄN VĂN C","030202067890")]
    for parcel in (100,101):
        for index,(name,cccd) in enumerate(people):
            ws.append([1 if parcel==100 and index==0 else None,name,10,parcel,100,cccd])
    wb.save(path); wb.close()

    data=inspect_workbook(path,"Nguồn",1,1,ColumnMapping(owner="B",sheet="C",parcel="D",area="E",location="",identity="F",household_index="A"),require_identity=True)

    assert [(record.sheet,record.parcel) for record in data.valid_records]==[("10","100"),("10","101")]
    assert all(len(record.household_people)==3 for record in data.valid_records)


def test_bad_mapping(tmp_path):
    with pytest.raises(UserError): ColumnMapping("B","G","H","M","L").validate(12)
    with pytest.raises(UserError): ColumnMapping("B","B","H","K","L").validate(12)
    with pytest.raises(UserError,match="diện tích"): ColumnMapping(area="").validate(12)
    with pytest.raises(UserError,match="STT hộ"): ColumnMapping(household_index="").validate(12)


def test_render_values_layout_parts_and_empty_location(tmp_path,service,config):
    data=workbook(tmp_path/"source.xlsx",[("HỘ A & <B>",72,175,"357,25",None)])
    payload=service.template.render(service.values(data.records[0],config,100)); text=all_text(payload)
    assert "{{" not in text and "HỘ A & <B>" in text and "000000000001" in text
    assert "xứ đồng ," not in text and "xứ đồng" not in text.casefold()
    assert text.count("357.25")==3 and "LUC" in text and "chưa xác định" in text and "2074" not in text
    with ZipFile(BytesIO(service.template.raw)) as original,ZipFile(BytesIO(payload)) as result:
        assert original.namelist()==result.namelist()
        for name in original.namelist():
            if name!="word/document.xml": assert original.read(name)==result.read(name),name
        old=minidom.parseString(original.read("word/document.xml")); new=minidom.parseString(result.read("word/document.xml"))
        for tag in ("pPr","rPr","sectPr"):
            assert [n.toxml() for n in old.getElementsByTagNameNS(WORD_NS,tag)]==[n.toxml() for n in new.getElementsByTagNameNS(WORD_NS,tag)]


def test_split_run_header_table_and_textbox(tmp_path,service,config):
    template=tmp_path/"split.docx"
    with ZipFile(BytesIO(service.template.raw)) as source,ZipFile(template,"w") as target:
        for item in source.infolist():
            if item.filename=="word/header1.xml": continue
            data=source.read(item)
            if item.filename=="word/document.xml":
                text=data.decode("utf-8")
                text=text.replace("{{HO_TEN}}","{{HO</w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>_TE</w:t></w:r><w:r><w:t>N}}")
                text=text.replace("</w:body>",'<w:tbl><w:tr><w:tc><w:p><w:r><w:t>{{SO_TO}}</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:p><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>{{SO_THUA}}</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p></w:body>')
                data=text.encode("utf-8")
            target.writestr(item,data)
        target.writestr("word/header1.xml",f'<w:hdr xmlns:w="{WORD_NS}"><w:p><w:r><w:t>{{{{HO_</w:t></w:r><w:r><w:t>TEN}}}}</w:t></w:r></w:p></w:hdr>')
    renderer=WordTemplate(template)
    row=workbook(tmp_path/"src.xlsx",[("HỘ THỬ",43,999,20,"Đồng xanh")]).records[0]
    payload=renderer.render(service.values(row,config,3))
    assert "HỘ THỬ" in all_text(payload) and "xứ đồng Đồng xanh" in all_text(payload)
    with ZipFile(BytesIO(payload)) as z:
        assert b"{{" not in z.read("word/header1.xml")
        assert "HỘ THỬ" in z.read("word/header1.xml").decode()


def test_preview_requires_confirmation_and_does_not_consume(tmp_path,service,config):
    data=workbook(tmp_path/"src.xlsx",[("HỘ A",1,1,20,None)])
    output=tmp_path/"out"
    with pytest.raises(UserError,match="xác nhận"): service.generate(data,config,output,"invalid")
    assert not output.exists()
    preview=service.preview(data,config,output,0,tmp_path/"preview")
    assert not output.exists() and "100" in all_text(preview.path.read_bytes())
    token=service.approve(preview); result=service.generate(data,config,output,token)
    assert result["success"]==1 and result["entries"][0].number==100
    with pytest.raises(UserError): service.generate(data,config,output,token)


@pytest.mark.parametrize("failure",["write","render","exists","race"])
def test_failed_number_reused(tmp_path,service,config,monkeypatch,failure):
    data,output,token=ready(tmp_path,service,config)
    original_publish=service_module.publish; original_render=service.template.render
    calls=[]
    if failure=="exists":
        output.mkdir(); (output/"CHUACOGIAY_12345_72_176-TBXN.docx").write_bytes(b"DO NOT CHANGE")
    def publish(payload,target):
        calls.append(target.name)
        if len(calls)==2 and failure in ("write","race"):
            raise PermissionError(13,"Locked") if failure=="write" else FileExistsError()
        return original_publish(payload,target)
    def render(values):
        if values["SO_THUA"]=="176" and failure=="render": raise ValueError("Mẫu thử lỗi")
        return original_render(values)
    monkeypatch.setattr(service_module,"publish",publish); monkeypatch.setattr(service.template,"render",render)
    result=service.generate(data,config,output,token)
    assert result["success"]==2 and [e.number for e in result["entries"]]==[100,101,101]
    assert result["entries"][2].status=="THÀNH CÔNG" and not result["log_error"]
    if failure=="exists": assert (output/"CHUACOGIAY_12345_72_176-TBXN.docx").read_bytes()==b"DO NOT CHANGE"
    assert not list(output.glob(".thongbao-*"))
    wb=load_workbook(result["xlsx"]); assert wb.active.max_row==4; wb.close()


def test_no_number_for_invalid_and_duplicates(tmp_path,service,config):
    rows=[("HỘ A",1,1,20,None),("HỘ B",1,1,20,None),("HỘ C",1,2,None,None),("HỘ D",1,3,20,None)]
    data,output,token=ready(tmp_path,service,config,rows)
    result=service.generate(data,config,output,token)
    assert result["success"]==1 and [e.number for e in result["entries"]]==[None,None,None,100]


def test_cancel_records_unprocessed(tmp_path,service,config):
    data,output,token=ready(tmp_path,service,config); state=[]
    result=service.generate(data,config,output,token,lambda *args:state.append(1),lambda:bool(state))
    assert result["success"]==result["processed"]==1 and result["cancelled"]
    assert [e.status for e in result["entries"]]==["THÀNH CÔNG","CHƯA XỬ LÝ","CHƯA XỬ LÝ"]
    assert [e.number for e in result["entries"]]==[100,None,None]


@pytest.mark.parametrize("change",["config","source","preview"])
def test_stale_preview_blocked(tmp_path,service,config,change):
    data=workbook(tmp_path/"src.xlsx",[("HỘ A",1,1,20,None)]); output=tmp_path/"out"
    preview=service.preview(data,config,output,0,tmp_path/"preview")
    if change=="preview":
        preview.path.write_bytes(b"edited")
        with pytest.raises(UserError): service.approve(preview)
    else:
        token=service.approve(preview)
        if change=="source": data.source.write_bytes(data.source.read_bytes()+b"changed")
        else: config=replace(config,day=30)
        with pytest.raises(UserError): service.generate(data,config,output,token)
    assert not output.exists()


def test_finite_number_pool_and_log_formula_escape(tmp_path,service,config):
    config=replace(config,number_mode="list",number_list="7")
    data,output,token=ready(tmp_path,service,config,[("=HỘ THỬ",1,1,20,None),("HỘ B",1,2,20,None)])
    result=service.generate(data,config,output,token)
    assert result["success"]==1 and result["entries"][1].number is None
    assert result["entries"][1].status=="HẾT SỐ THÔNG BÁO"
    wb=load_workbook(result["xlsx"]); assert wb.active["B2"].data_type=="s"; wb.close()


def test_config_and_filename(config):
    from tools.notice_builder.core.models import NoticeRecord
    row=NoticeRecord(1,"Không được vào tên file",1,"72","175","20","")
    assert service_module.safe_filename(config,row)==("CHUACOGIAY_12345_72_175-TBXN.docx",False)
    assert service_module.safe_filename(replace(config,prefix="CHUACAPGIAY"),row)==("CHUACAPGIAY_12345_72_175-TBXN.docx",False)
    name,cleaned=service_module.safe_filename(replace(config,suffix='A:/B?'),row)
    assert cleaned and name=="CHUACOGIAY_12345_72_175-A__B_.docx"
    with pytest.raises(UserError): replace(config,day=31,month=2).validate()
    with pytest.raises(UserError): replace(config,commune_code="").validate()


def test_validation_log_without_valid_rows(tmp_path):
    from tools.notice_builder.core.logging import export_inspection
    data=workbook(tmp_path/"src.xlsx",[("HỘ A",1,1,None,None),("HỘ B",1,None,20,None)])
    result=export_inspection(data,tmp_path/"reports")
    assert result["txt"].is_file() and result["xlsx"].is_file()
    assert not list(tmp_path.rglob("*.docx"))
    text=result["txt"].read_text(encoding="utf-8-sig"); assert "K2" in text and "H3" in text


def test_log_write_error_stops_before_next_file(tmp_path,service,config,monkeypatch):
    from tools.notice_builder.core.logging import RunLog
    data,output,token=ready(tmp_path,service,config)
    original=RunLog.append
    def fail(self,entry):
        original(self,entry)
        raise PermissionError("log locked")
    monkeypatch.setattr(RunLog,"append",fail)
    result=service.generate(data,config,output,token)
    assert result["success"]==1 and result["cancelled"] and result["log_error"]
    assert len(list(output.glob("*.docx")))==1 and result["entries"][0].number==100


def test_xlsx_log_error_keeps_txt_and_created_files(tmp_path,service,config,monkeypatch):
    from tools.notice_builder.core.logging import RunLog
    data,output,token=ready(tmp_path,service,config)
    def fail(self): raise PermissionError("no space")
    monkeypatch.setattr(RunLog,"finish",fail)
    result=service.generate(data,config,output,token)
    assert result["success"]==3 and result["log_error"] and result["txt"].is_file() and result["xlsx"] is None


def test_invalid_template_and_unresolved_tokens(tmp_path,service):
    broken=tmp_path/"bad.docx"; broken.write_bytes(b"not a docx")
    with pytest.raises(UserError): WordTemplate(broken)
    with pytest.raises(UserError,match="placeholder"): service.template.render({})


def test_default_template_spacing_is_normalized_and_research_copy_matches():
    packaged=default_template_path()
    root=Path(__file__).resolve().parents[4]
    research=root/"research"/"notice_template"/packaged.name
    assert research.read_bytes()==packaged.read_bytes()
    evidence=json.loads((research.parent/"template-evidence.json").read_text(encoding="utf-8"))
    with ZipFile(packaged) as archive:
        document=minidom.parseString(archive.read("word/document.xml"))
    checked=set()
    for paragraph in document.getElementsByTagNameNS(WORD_NS,"p"):
        para_id=paragraph.getAttribute("w14:paraId")
        if para_id not in evidence["changed_paragraphs"]: continue
        checked.add(para_id)
        text="".join(node_text(node) for node in paragraph.getElementsByTagNameNS(WORD_NS,"t"))
        assert "  " not in text
        assert paragraph.getElementsByTagNameNS(WORD_NS,"jc")[0].getAttribute("w:val")=="left"
    assert checked==set(evidence["changed_paragraphs"])


@pytest.fixture(scope="module")
def app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication(["notice-test","-platform","offscreen"])


def test_ui_standalone_and_config_restore(app,tmp_path,monkeypatch,config):
    from dataclasses import asdict
    from PySide6.QtCore import QEventLoop, QTimer
    from tools.notice_builder.ui import MainWindow
    monkeypatch.setenv("APPDATA",str(tmp_path/"profile")); monkeypatch.chdir(tmp_path)
    data=workbook(tmp_path/"src.xlsx",[("HỘ A",1,1,20,None)],header=6)
    window=MainWindow(); window.show(); app.processEvents()
    assert window.stack.count()==8 and not window.confirm_button.isEnabled()
    assert not window.windowIcon().isNull()
    assert window.template_preview.pixmap() and not window.template_preview.pixmap().isNull()
    assert window.template_preview.accessibleName()=="Ảnh xem trước mẫu Word có placeholder"
    assert window.inputs["commune_code"].placeholderText()=="Ví dụ: 10930"
    assert window.inputs["owner_address"].placeholderText().startswith("Ví dụ: thôn")
    assert window.inputs["prefix"].placeholderText()=="Ví dụ: CHUACOGIAY hoặc CHUACAPGIAY"
    assert window.output.placeholderText().startswith("Ví dụ: D:/Ho so")
    assert "Số thiếu khác ngày" in window.numbering_cases_note.text()
    assert "Hết số" in window.numbering_cases_note.text()
    for key,value in asdict(config).items():
        widget=window.inputs.get(key)
        if widget:
            if hasattr(widget,"setValue"): widget.setValue(value)
            else: widget.setText(str(value))
    window.source.setText(str(data.source)); window.output.setText(str(tmp_path/"out"))
    info=workbook_info(data.source); window.source_ready(info)
    window.mapping["area"].setCurrentIndex(window.mapping["area"].findData("M"))
    window.row_start.setText("7"); window.row_end.setText("7")
    window.number_mode.setCurrentIndex(window.number_mode.findData("list")); window.number_list.setText("300-400")
    window.add_number_date_rule("300-350","2026-08-25"); window.add_number_date_rule("351-400","2026-08-28")
    assert "SU_DUNG_CHUNG" not in window.template_inputs
    window.template_inputs["NGUOI_KY"].setText("NGƯỜI KÝ THỬ")
    window.optional_empty.setCurrentIndex(window.optional_empty.findData("dots"))
    window.save_settings(); window.close()
    restored=MainWindow(); errors=[]; restored.error=errors.append
    restored.read_source()
    loop=QEventLoop(); restored.job.finished.connect(loop.quit); QTimer.singleShot(15000,loop.quit); loop.exec(); app.processEvents()
    assert not errors and restored.job is None
    assert restored.header.value()==6 and restored.mapping["area"].currentData()=="M"
    assert restored.row_start.text()=="7" and restored.row_end.text()=="7"
    assert restored.number_mode.currentData()=="list" and restored.number_list.text()=="300-400"
    assert restored.number_date_rules()==[{"numbers":"300-350","date":"2026-08-25"},{"numbers":"351-400","date":"2026-08-28"}]
    restored.navigate(1)
    loop=QEventLoop(); restored.job.finished.connect(loop.quit); QTimer.singleShot(15000,loop.quit); loop.exec(); app.processEvents()
    assert not errors and restored.job is None and restored.steps.currentRow()==2
    assert restored.mapping["area"].currentData()=="M"
    assert "SU_DUNG_CHUNG" not in restored.template_inputs
    assert restored.template_inputs["NGUOI_KY"].isHidden()
    assert not restored.empty_location.isHidden()
    assert restored.optional_empty.isHidden()
    assert restored.optional_empty.currentData()=="dots"
    assert restored.inputs["commune_code"].text()=="12345" and not restored.confirm_button.isEnabled()
    restored.close(); app.processEvents()


def test_ui_confirmation_cancel_then_accept(app,tmp_path,monkeypatch,config,service):
    from dataclasses import asdict
    from PySide6.QtCore import QEventLoop,QTimer
    from PySide6.QtWidgets import QMessageBox
    from tools.notice_builder.ui import MainWindow
    monkeypatch.setenv("APPDATA",str(tmp_path/"profile"))
    data=workbook(tmp_path/"source.xlsx",[("HỘ THỬ",1,1,20,None)])
    output=tmp_path/"out"; window=MainWindow(); errors=[]; window.error=errors.append
    for key,value in asdict(config).items():
        field=window.inputs.get(key)
        if field:
            if hasattr(field,"setValue"): field.setValue(value)
            else: field.setText(str(value))
    window.start_number.setValue(config.start_number); window.output.setText(str(output))
    window.inspection_ready(data)
    preview=service.preview(data,config,output,0,tmp_path/"preview")
    window.preview_ready((service,preview))
    choice=["Quay lại kiểm tra"]
    def select(dialog):
        assert dialog.defaultButton().text()=="Quay lại kiểm tra"
        next(b for b in dialog.buttons() if b.text()==choice[0]).click()
        return 0
    monkeypatch.setattr(QMessageBox,"exec",select)
    window.confirm_batch(); assert not output.exists() and window.job is None and not errors
    choice[0]="Đồng ý tạo"; window.confirm_batch()
    assert window.job is not None,errors
    loop=QEventLoop(); window.job.finished.connect(loop.quit); QTimer.singleShot(15000,loop.quit); loop.exec(); app.processEvents()
    assert not errors and window.result["success"]==1 and window.job is None
    assert not window.confirm_button.isEnabled()
    window.close()


def test_real_workbook_read_only():
    path=os.environ.get("NOTICE_REAL_WORKBOOK")
    if not path: pytest.skip("Set NOTICE_REAL_WORKBOOK for the user's read-only regression input")
    before=file_hash(path); info=workbook_info(path)
    assert info["header_row"]==5 and info["depth"]==2 and info["rows"]==1001
    data=inspect_workbook(path,info["sheet"],5,2,ColumnMapping())
    assert len(data.records)==559 and len(data.valid_records)==556
    assert len(data.summary_rows)==217 and len(data.blank_rows)==156 and len(data.name_only_rows)==63
    assert not any(r.duplicate_rows for r in data.records)
    assert [r.source_row for r in data.records if not r.valid]==[943,944,945]
    assert sum(r.owner_row!=r.source_row for r in data.records)==246
    assert sum(not r.location for r in data.records)==313
    assert before==file_hash(path)


def test_real_workbook_new_identity_requirement():
    path=os.environ.get("NOTICE_REAL_WORKBOOK")
    if not path: pytest.skip("Set NOTICE_REAL_WORKBOOK for private data")
    before=file_hash(path); info=workbook_info(path)
    data=inspect_workbook(path,info["sheet"],5,2,ColumnMapping(identity="D"),require_identity=True)
    assert len(data.records)==559 and len(data.valid_records)==555
    assert [r.source_row for r in data.records if not r.valid]==[279,943,944,945]
    assert "D279" in " ".join(next(r for r in data.records if r.source_row==279).errors)
    assert before==file_hash(path)


def edge_workbook(path, rows):
    wb=Workbook(); ws=wb.active; ws.title="Nguồn"
    ws.append(["Tên hộ","Tờ BĐ mới","Thửa BĐ mới","Diện tích bản đồ","Giấy tờ nhân thân","Xứ đồng","STT hộ"])
    serial=0
    for row in rows:
        values=list(row); explicit_marker=values[6] if len(values)>6 else "AUTO"
        owner=values[0] if values else None
        if explicit_marker=="AUTO":
            if owner is not None and not is_summary(owner): serial+=1; marker=serial
            else: marker=None
        else: marker=explicit_marker
        padded=values[:6]+[None]*(6-len(values[:6]))
        ws.append(padded+[marker])
    wb.save(path); wb.close()
    return path


def inspect_edge(path, strict=True):
    return inspect_workbook(path,"Nguồn",1,1,ColumnMapping("A","B","C","D","F","E",household_index="G"),require_identity=strict)


def test_uncached_running_household_numbers_preserve_owners_and_range(tmp_path):
    path=edge_workbook(tmp_path/"sequence.xlsx",[
        ["HỘ A",1,1,100,"GIẤY A",None,5],
        ["Tổng DT",None,None,None,None,None,None],
        ["HỘ B",1,2,100,"GIẤY B",None,"=MAX($G$2:G3)+1"],
        ["THÀNH VIÊN B",1,3,100,"GIẤY THÀNH VIÊN",None,None],
        ["HỘ C",1,4,100,None,None,"=MAX($G$2:G5)+1"],
        [None,1,5,100,None,None,None]])
    before=file_hash(path); data=inspect_edge(path)
    assert [r.household_number for r in data.records]==["5","6","6","7","7"]
    assert [r.owner for r in data.records]==["HỘ A","HỘ B","HỘ B","HỘ C","HỘ C"]
    assert [r.identity for r in data.records]==["GIẤY A","GIẤY B","GIẤY B","",""]
    assert len(data.valid_records)==3 and all("E6" in " ".join(r.errors) for r in data.records[-2:])
    selected=inspect_workbook(path,"Nguồn",1,1,data.mapping,require_identity=True,start_row=5,end_row=5)
    assert selected.records==[data.records[2]]
    assert file_hash(path)==before


@pytest.mark.parametrize("formula",["=1+1","=MAX($G$2:G4)+1","=MAX($G$3:G2)+1","=MAX($F$2:F2)+1","=MAX($G$2:G2)+2","#REF!"])
def test_unknown_household_formula_never_borrows_previous_identity(tmp_path,formula):
    path=edge_workbook(tmp_path/"unknown.xlsx",[
        ["HỘ A",1,1,100,"GIẤY A",None,1],
        ["HỘ B",1,2,100,"GIẤY B",None,formula],
        ["HỘ C",1,3,100,"GIẤY C",None,"=MAX($G$2:G3)+1"]])
    data=inspect_edge(path)
    assert len(data.valid_records)==1
    assert all(not r.owner and not r.identity and "STT hộ" in " ".join(r.errors) for r in data.records[1:])


def test_household_formula_uses_exact_range_even_when_numbers_repeat(tmp_path):
    path=edge_workbook(tmp_path/"range.xlsx",[
        ["HỘ A",1,1,100,"GIẤY A",None,1],
        ["HỘ B",1,2,100,"GIẤY B",None,"=MAX($G$2:G2)+1"],
        ["HỘ C",1,3,100,"GIẤY C",None,"=MAX($G$2:G2)+1"],
        ["HỘ D",1,4,100,"GIẤY D",None,"=MAX($G$2:G4)+1"]])
    data=inspect_edge(path)
    assert len(data.valid_records)==4
    assert [r.household_number for r in data.records]==["1","2","2","3"]
    assert [r.identity for r in data.records]==["GIẤY A","GIẤY B","GIẤY C","GIẤY D"]


@pytest.mark.parametrize("bad_name",['="HỘ B"',"#REF!","#VALUE!","#N/A","...."])
@pytest.mark.parametrize("name_only",[False,True])
def test_bad_owner_blocks_inheritance_until_next_valid_owner(tmp_path,bad_name,name_only):
    path=edge_workbook(tmp_path/"edge.xlsx",[
        ["HỘ A",1,1,100,"GIẤY A"],
        [bad_name,None if name_only else 1,None if name_only else 2,None if name_only else 100,"GIẤY B"],
        [None,1,3,100], ["HỘ C",1,4,100,"GIẤY C"], [None,1,5,100]])
    before=file_hash(path); data=inspect_edge(path)
    bad=[r for r in data.records if r.source_row in (3,4)]
    assert bad and all(not r.valid and r.owner=="" and any("A3" in e for e in r.errors) for r in bad)
    assert [r.owner for r in data.valid_records]==["HỘ A","HỘ C","HỘ C"]
    assert before==file_hash(path)


def test_cached_formula_owner_is_read_and_blank_literal_inherits(tmp_path):
    from xml.etree import ElementTree as ET
    path=edge_workbook(tmp_path/"cached.xlsx",[["HỘ A",1,1,100,"GIẤY A"],['="HỘ B"',1,2,100,"GIẤY B"],[None,1,3,100]])
    ns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(path) as z: entries=[(i,z.read(i)) for i in z.infolist()]
    with ZipFile(path,"w") as z:
        for item,payload in entries:
            if item.filename=="xl/worksheets/sheet1.xml":
                root=ET.fromstring(payload); cell=root.find('.//s:c[@r="A3"]',ns); cell.set("t","str")
                cell.find("s:v",ns).text="HỘ B"; payload=ET.tostring(root,encoding="utf-8")
            z.writestr(item,payload)
    data=inspect_edge(path)
    assert [r.owner for r in data.valid_records]==["HỘ A","HỘ B","HỘ B"]
    assert data.records[-1].identity=="GIẤY B" and data.records[-1].identity_row==3


@pytest.mark.parametrize("label",["Tổng DT:"," TỔNG DIỆN TÍCH... ","tổng cộng：","Cộng.","TONG DT…"])
def test_punctuated_summary_not_owner(tmp_path,label):
    data=inspect_edge(edge_workbook(tmp_path/"summary.xlsx",[["TỐNG VĂN CÔNG",1,1,100,"GIẤY A"],[label,1,2,"#REF!"],[None,1,3,100],["CÔNG THỊ TỐNG",1,4,100,"GIẤY B"]]))
    assert data.summary_rows==[3] and len(data.valid_records)==3
    assert data.records[1].owner=="TỐNG VĂN CÔNG" and data.records[2].owner=="CÔNG THỊ TỐNG"


@pytest.mark.parametrize("column",["B","C","D","E","F"])
def test_mapped_formula_without_cache_logs_cell(tmp_path,column):
    path=edge_workbook(tmp_path/"formula.xlsx",[["HỘ A",1,1,100,"GIẤY A","Đồng A"]])
    wb=load_workbook(path); wb.active[f"{column}2"]="=1+1"; wb.save(path); wb.close()
    row=inspect_edge(path).records[0]
    assert not row.valid and any(f"{column}2" in e and "công thức" in e for e in row.errors)


@pytest.mark.parametrize("formula,expected",[("=755-678","77"),("= 100.25 - 0.15 ","100.1"),("=10-0","10")])
def test_area_literal_subtraction_without_cache(tmp_path,formula,expected):
    path=edge_workbook(tmp_path/"area.xlsx",[["HỘ A",41,441,formula,"GIẤY A"]])
    before=file_hash(path); row=inspect_edge(path).records[0]
    assert row.valid and row.area==expected and file_hash(path)==before


@pytest.mark.parametrize("formula",["=1-2","=2-2","=B2-C2","=SUM(755,-678)","=755-678+1","=1,5-0,5"])
def test_invalid_or_unsupported_area_formula_still_blocks(tmp_path,formula):
    path=edge_workbook(tmp_path/"bad-area.xlsx",[["HỘ A",41,441,formula,"GIẤY A"]])
    row=inspect_edge(path).records[0]
    assert not row.valid and any("D2" in e for e in row.errors)


def test_identity_required_no_cross_household_and_no_number(tmp_path,service,config):
    data=inspect_edge(edge_workbook(tmp_path/"identity.xlsx",[["HỘ A",1,1,100,"001234567890"],[None,1,2,100],
        ["HỘ B",1,3,100],[None,1,4,100],["HỘ C",1,5,100,"GIẤY C"]]))
    assert [r.identity for r in data.records]==["001234567890","001234567890","","","GIẤY C"]
    assert all("E4" in " ".join(r.errors) for r in data.records[2:4])
    preview=service.preview(data,config,tmp_path/"out",0,tmp_path/"preview")
    result=service.generate(data,config,tmp_path/"out",service.approve(preview))
    assert [e.number for e in result["entries"]]==[100,101,None,None,102]
    assert "E4" in result["txt"].read_text(encoding="utf-8-sig")


def test_service_cannot_export_missing_identity_even_without_strict_inspection(tmp_path,service,config):
    data=inspect_edge(edge_workbook(tmp_path/"identity.xlsx",[["HỘ A",1,1,100,"GIẤY A"],["HỘ B",1,2,100]]),False)
    with pytest.raises(UserError,match="Giấy tờ nhân thân"): service.preview(data,config,tmp_path/"out",1,tmp_path/"pre")
    preview=service.preview(data,config,tmp_path/"out",0,tmp_path/"pre")
    result=service.generate(data,config,tmp_path/"out",service.approve(preview))
    assert result["success"]==1 and result["entries"][1].number is None


from tools.notice_builder.core.fields import REQUIRED_TOKENS, REQUIRED_COMMON, OPTIONAL_COMMON


@pytest.mark.parametrize("token",list(REQUIRED_TOKENS))
@pytest.mark.parametrize("empty",["", "   ", "...."])
def test_all_mandatory_placeholders_reject_empty(tmp_path,service,config,token,empty):
    row=workbook(tmp_path/"src.xlsx",[("HỘ A",1,1,100,None)]).records[0]
    values=service.values(row,config,100); values[token]=empty
    with pytest.raises(UserError,match="bắt buộc"): service.template.render(values)


@pytest.mark.parametrize("token",list(REQUIRED_COMMON))
def test_config_requires_explicit_common_inputs(config,token):
    values=dict(config.template_fields); values[token]="...."
    with pytest.raises(UserError,match="Cần nhập"): replace(config,template_fields=values).validate()


@pytest.mark.parametrize("mode,expected",[("blank","..."),("dots","...")])
def test_optional_slots_and_editable_common_fields(tmp_path,service,config,mode,expected):
    row=workbook(tmp_path/"src.xlsx",[("HỘ A",1,1,100,None)]).records[0]
    fields={**config.template_fields,"NGUOI_KY":"NGƯỜI KÝ THỬ"}
    cfg=replace(config,template_fields=fields,optional_empty=mode)
    values=service.values(row,cfg,100)
    assert all(values[k]==expected for k in OPTIONAL_COMMON)
    assert values["DIEN_TICH"]==values["SU_DUNG_CHUNG"]=="100"
    text=all_text(service.template.render(values))
    assert "Nguyễn Khắc Nghĩa" in text and "NGƯỜI KÝ THỬ" not in text and "{{" not in text


def test_shared_area_is_always_derived_from_parcel_area(tmp_path,service,config):
    row=workbook(tmp_path/"shared-area.xlsx",[("HỘ A",1,1,245,None)]).records[0]
    legacy_fields={**config.template_fields,"SU_DUNG_CHUNG":"999"}
    values=service.values(row,replace(config,template_fields=legacy_fields),100)
    assert values["DIEN_TICH"]==values["SU_DUNG_CHUNG"]=="245"
    text=all_text(service.template.render(values))
    assert "999" not in text and text.count("245")>=3


def test_ui_mapping_preserved_on_next_and_reset_for_other_workbook(app,tmp_path,monkeypatch):
    from tools.notice_builder.ui import MainWindow
    monkeypatch.setenv("APPDATA",str(tmp_path/"profile"))
    path=edge_workbook(tmp_path/"mapping.xlsx",[["HỘ A",1,1,100,"GIẤY A",250]])
    window=MainWindow(); window.run_job=lambda fn,cb:cb(fn(None))
    window.source.setText(str(path)); window.read_source()
    window.mapping["location"].setCurrentIndex(0); area=window.mapping["area"]; area.setCurrentIndex(area.findData("F"))
    window.navigate(1)
    assert area.currentData()=="F" and window.mapping["location"].currentData()==""
    data=inspect_workbook(path,"Nguồn",1,1,ColumnMapping(**{k:c.currentData() for k,c in window.mapping.items()}),require_identity=True)
    assert data.records[0].area=="250"
    other=edge_workbook(tmp_path/"other.xlsx",[["HỘ B",2,2,300,"GIẤY B"]])
    window.source.setText(str(other)); window.read_source()
    assert area.currentData()=="D"
    window.close()


def test_saved_template_path_fallbacks(tmp_path):
    from tools.notice_builder.paths import saved_template_path
    # template_kind=default always returns the bundled MAU_22
    assert saved_template_path({"template_kind": "default", "template": str(tmp_path / "old.docx")}) == default_template_path()
    # missing file falls back to default
    assert saved_template_path({"template": str(tmp_path / "nonexistent.docx")}) == default_template_path()
    # existing custom .docx is preserved
    custom = tmp_path / "custom.docx"; custom.write_bytes(b"PK custom")
    assert saved_template_path({"template": str(custom)}) == custom


def test_canonical_template_preserves_geometry_and_only_authorized_spacing():
    root=Path(__file__).resolve().parents[4]
    reference=root/"research/notice_template/MAU_22_THONG_BAO_XAC_NHAN_KET_QUA_DANG_KY_DAT_DAI.docx"
    evidence=json.loads((reference.parent/"template-evidence.json").read_text(encoding="utf-8"))
    assert file_hash(reference)==evidence["canonical_sha256"]==file_hash(default_template_path())
    assert reference.read_bytes()==default_template_path().read_bytes()
    tokens={slot["token"] for slot in evidence["slots"]}
    assert len(tokens)==28 and WordTemplate(reference).tokens==tokens
    assert set(REQUIRED_TOKENS)<=tokens and set(OPTIONAL_COMMON)<=tokens
    assert "SU_DUNG_CHUNG" in tokens
    assert {"SU_DUNG_CHUNG_VO_CHONG","THOI_DIEM_SU_DUNG","THOI_HAN_SU_DUNG"}.isdisjoint(tokens)
    with ZipFile(reference) as source:
        assert source.namelist()==[item["part"] for item in evidence["inventory"]]
        for item in evidence["inventory"]:
            assert hashlib.sha256(source.read(item["part"])).hexdigest()==item["sha256"],item["part"]


def test_template_configuration_keeps_mau_22_default_and_values_in_one_source():
    configs={item.id:item for item in template_configs()}
    default=default_template_config()
    assert default.id=="MAO_DIEN" and default.is_default
    assert default.path==default_template_path()
    assert default.static_values["TEN_XA"]=="MAO ĐIỀN"
    assert default.static_values["NGUOI_KY"]=="Nguyễn Khắc Nghĩa"
    assert configs["CAM_GIANG"].static_values["TEN_XA"]=="CẨM GIANG"
    assert configs["CAM_GIANG"].static_values["DIA_DIEM"]=="Cẩm Giang"
    assert configs["CAM_GIANG"].static_values["DON_VI_LUU"]=="KT. (1)"
    with ZipFile(default.path) as archive:
        xml=archive.read("word/document.xml").decode("utf-8")
    assert xml.count("{{TEN_XA}}") == 1
    assert xml.count("{{DIA_DIEM}}") == 1
    source_root=Path(__file__).resolve().parents[2]
    python_source="\n".join(path.read_text(encoding="utf-8") for path in source_root.rglob("*.py")
                            if "tests" not in path.parts)
    assert "MAO ĐIỀN" not in python_source and "Mao Điền" not in python_source


def cam_giang_household(path):
    wb=Workbook(); ws=wb.active; ws.title="Nguồn"
    ws.append(["Họ và tên","Tờ BĐ mới","Thửa BĐ mới","Diện tích bản đồ","CCCD","Xứ đồng","STT hộ","Ngày sinh"])
    ws.append(["NGUYỄN VĂN CHỦ",77,234,537,"030064013684","Đồng thử",1,"28/08/1964"])
    ws.append(["NGUYỄN THỊ THÀNH VIÊN",None,None,None,"030200012345",None,None,"02/03/2000"])
    ws.append(["NGUYỄN VĂN THÀNH VIÊN",None,None,None,"030202067890",None,None,2002])
    wb.save(path); wb.close()
    mapping=ColumnMapping(owner="A",sheet="B",parcel="C",area="D",identity="E",location="F",household_index="G",birth_date="H")
    return inspect_workbook(path,"Nguồn",1,1,mapping,require_identity=True)


def table_text_rows(payload, table_index):
    with ZipFile(BytesIO(payload)) as archive:
        document=minidom.parseString(archive.read("word/document.xml"))
    table=document.getElementsByTagNameNS(WORD_NS,"tbl")[table_index]
    return [" ".join(node_text(node) for node in row.getElementsByTagNameNS(WORD_NS,"t"))
            for row in table.getElementsByTagNameNS(WORD_NS,"tr")]


def test_cam_giang_uses_vbdlis_head_model_and_dynamic_member_rows(tmp_path,config):
    data=cam_giang_household(tmp_path/"household.xlsx")
    record=data.records[0]
    assert [person.is_head for person in record.household_people]==[True,False,False]
    assert [person.name for person in record.members]==["NGUYỄN THỊ THÀNH VIÊN","NGUYỄN VĂN THÀNH VIÊN"]
    assert [person.birth_date for person in record.members]==["02/03/2000","2002"]

    cam=next(item for item in template_configs() if item.id=="CAM_GIANG")
    service=NoticeService(cam.path,json.loads(resource("config/legal_defaults.json").read_text(encoding="utf-8")))
    configured=replace(config,day=9,month=9,year=2026)
    payload=service.template.render(service.values(record,configured,23))
    mao_service=NoticeService(default_template_path(),json.loads(resource("config/legal_defaults.json").read_text(encoding="utf-8")))
    mao_values=mao_service.values(record,configured,23)
    assert "MEMBERS" not in mao_values and "NGAY_SINH" not in mao_values
    text=all_text(payload); member_rows=table_text_rows(payload,2)
    assert "{{" not in text and "Sinh năm 28/08/1964" in text
    assert "XÃ CẨM GIANG" in text and "Cẩm Giang, ngày 09 tháng 09 năm 2026" in text
    assert "Lưu: VT, KT. (1)." in text
    assert text.count("23/TB-UBND")==2
    assert text.count("ngày 09 tháng 09 năm 2026")==2
    assert len(member_rows)==6
    assert not any("NGUYỄN VĂN CHỦ" in row for row in member_rows)
    assert any("NGUYỄN THỊ THÀNH VIÊN" in row and "02/03/2000" in row and "030200012345" in row for row in member_rows)
    assert any("NGUYỄN VĂN THÀNH VIÊN" in row and "2002" in row and "030202067890" in row for row in member_rows)


def test_required_fields_are_isolated_per_template(config):
    mao=default_template_config()
    cam=next(item for item in template_configs() if item.id=="CAM_GIANG")
    assert "NGAY_SINH" not in mao.required_fields and "NGAY_SINH" in cam.required_fields
    no_tax={**config.template_fields,"CO_QUAN_THUE":""}
    replace(config,template_fields=no_tax).validate(
        key for key in mao.user_fields if key in mao.required_fields)
    replace(config,template_fields=no_tax).validate(
        key for key in cam.user_fields if key in cam.required_fields)

    record=NoticeRecord(2,"CHỦ HỘ",2,"1","1","100","",identity="012345678901")
    mao_service=NoticeService(mao.path,json.loads(resource("config/legal_defaults.json").read_text(encoding="utf-8")))
    cam_service=NoticeService(cam.path,json.loads(resource("config/legal_defaults.json").read_text(encoding="utf-8")))
    mao_service.template.render(mao_service.values(record,config,1))
    with pytest.raises(UserError,match="Ngày sinh chủ hộ"):
        cam_service.template.render(cam_service.values(record,config,1))


def test_cam_giang_template_contract_is_separate_from_mau_22():
    cam=next(item for item in template_configs() if item.id=="CAM_GIANG")
    assert template_config_for_path(cam.path)==cam
    renderer=WordTemplate(cam.path,cam.required_fields,cam.document_rules)
    assert "NGAY_SINH" in renderer.tokens
    assert {"MEMBER_NAME","MEMBER_BIRTH_DATE","MEMBER_IDENTITY","MEMBER_ADDRESS"} <= renderer.tokens
    assert "NGAY_SINH" not in WordTemplate(default_template_path()).tokens


@pytest.mark.parametrize("fallback,expected", [("blank", ""), ("village", "Thôn thử")])
def test_mao_defaults_ignore_old_inputs_and_location_fallback(tmp_path, service, config, fallback, expected):
    row = workbook(tmp_path / "defaults.xlsx", [("HỘ A", 1, 1, 100, None)]).records[0]
    cfg = replace(config, village="Thôn thử", empty_location=fallback, optional_empty="dots",
                  template_fields={**config.template_fields, **dict.fromkeys((*OPTIONAL_COMMON, "NGUOI_KY", "CO_QUAN_THUE", "DON_VI_LUU", "CHI_NHANH_VP_DKDD"), "giá trị cũ")})
    values = service.values(row, cfg, 1)
    assert values["XU_DONG"] == expected
    from tools.notice_builder.core.renderer import paragraphs, paragraph_nodes, node_text
    rendered = NoticeService(default_template_path(), service.legal).template.render(values)
    with ZipFile(BytesIO(rendered)) as archive:
        document = minidom.parseString(archive.read("word/document.xml"))
        lines = ["".join(node_text(n) for n in paragraph_nodes(p)) for p in paragraphs(document)]
        address = next(line for line in lines if line.startswith("c) Địa chỉ:") and cfg.administrative_address.strip() in line)
        location = f"xứ đồng {expected}, " if expected else ""
        assert address == f"c) Địa chỉ: {location}thôn Thôn thử, {cfg.administrative_address.strip()}"
        assert "..." not in address
        document.unlink()
    assert values["CO_QUAN_THUE"] == "Thuế Cơ sở 11 TP Hải Phòng"
    assert values["DON_VI_LUU"] == "XDNN&MT"
    assert values["CHI_NHANH_VP_DKDD"] == "Cẩm Giàng"
    assert all(values[key] == "..." for key in OPTIONAL_COMMON)
    cam = next(item for item in template_configs() if item.id == "CAM_GIANG")
    cam_service = NoticeService(cam.path, service.legal)
    assert cam_service.values(row, replace(cfg, empty_location="village"), 1) == cam_service.values(row, replace(cfg, empty_location="blank"), 1)
