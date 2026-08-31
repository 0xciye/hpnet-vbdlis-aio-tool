from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ITEM49 = re.compile(r"^(?P<stem>CHUACOGIAY_.+_.+_.+)-TBXN\.pdf, (?P=stem)-DDK\.pdf$")


class GoldenReferenceAnalyzer:
    def analyze(self, path: str | Path) -> dict[str, Any]:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            ws = workbook[workbook.sheetnames[0]]
            data_rows = range(5, ws.max_row + 1)
            populated = {}
            for col in range(1, 62):
                count = sum(
                    ws.cell(row, col).value is not None and str(ws.cell(row, col).value).strip() != ""
                    for row in data_rows
                )
                if count:
                    populated[get_column_letter(col)] = count
            roles = Counter(str(ws[f"N{row}"].value or "") for row in data_rows)
            return {
                "sheet": ws.title,
                "rows": ws.max_row - 4,
                "populated_columns": populated,
                "active_field_count": len(populated),
                "item49_pass": all(ITEM49.fullmatch(str(ws[f"AX{row}"].value or "")) for row in data_rows),
                "roles": dict(roles),
                "roles_pass": set(roles) <= {"Chủ hộ", "Thành viên hộ gia đình"},
                "gcn_blank_pass": all(ws[f"D{row}"].value in (None, "") and ws[f"E{row}"].value in (None, "") for row in data_rows),
                "cccd_text_pass": all(ws[f"I{row}"].number_format == "@" for row in data_rows),
                "stt_pass": [int(str(ws[f"A{row}"].value).strip()) for row in data_rows] == list(range(1, ws.max_row - 3)),
            }
        finally:
            workbook.close()
