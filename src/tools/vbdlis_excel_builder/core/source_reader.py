from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


@dataclass(slots=True)
class SourceColumn:
    index: int
    letter: str
    header: str

    @property
    def display(self) -> str:
        return f"{self.letter} - {self.header or '[Không có tiêu đề]'}"


class SourceReader:
    def sheet_names(self, path: str | Path) -> list[str]:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    def detect_header_row(self, path: str | Path, sheet_name: str, scan_rows: int = 20) -> int:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = workbook[sheet_name]
            best_row, best_score = 1, -1
            keywords = {
                "stt", "họ tên", "họ và tên", "cccd", "ngày sinh", "số tờ",
                "tờ bản đồ", "số thửa", "diện tích", "xứ đồng", "vị trí",
            }
            for row_idx in range(1, min(scan_rows, ws.max_row) + 1):
                values = [cell.value for cell in ws[row_idx]]
                nonblank = [str(value).strip() for value in values if value not in (None, "")]
                matches = sum(
                    1 for text in nonblank
                    if any(keyword in text.casefold() for keyword in keywords)
                )
                score = len(nonblank) + matches * 5
                if score > best_score:
                    best_row, best_score = row_idx, score
            return best_row
        finally:
            workbook.close()

    @staticmethod
    def _merge_header_values(top_values: list[Any], bottom_values: list[Any]) -> list[str]:
        """Build stable column names from one or two header rows.

        Excel merged cells expose the value only in the first cell. Carrying
        that value to the right keeps a two-level header useful for mapping.
        """
        size = max(len(top_values), len(bottom_values))
        merged: list[str] = []
        current_top = ""
        for index in range(size):
            top = top_values[index] if index < len(top_values) else None
            bottom = bottom_values[index] if index < len(bottom_values) else None
            top_text = "" if top is None else str(top).strip()
            bottom_text = "" if bottom is None else str(bottom).strip()
            if top_text:
                current_top = top_text
            if top_text == bottom_text:
                merged.append(top_text or bottom_text)
            elif current_top and bottom_text:
                merged.append(f"{current_top} - {bottom_text}")
            else:
                merged.append(bottom_text or current_top)
        return merged

    def _header_values(
        self,
        ws,
        header_row: int,
        header_row_2: int | None = None,
    ) -> list[str]:
        top_values = list(next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
        if header_row_2 is None:
            return ["" if value is None else str(value).strip() for value in top_values]
        if header_row_2 <= header_row:
            raise ValueError("Dòng tiêu đề thứ hai phải lớn hơn dòng tiêu đề thứ nhất.")
        bottom_values = list(next(ws.iter_rows(min_row=header_row_2, max_row=header_row_2, values_only=True)))
        return self._merge_header_values(top_values, bottom_values)

    def headers(
        self,
        path: str | Path,
        sheet_name: str,
        header_row: int,
        header_row_2: int | None = None,
    ) -> list[SourceColumn]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = workbook[sheet_name]
            values = self._header_values(ws, header_row, header_row_2)
            return [
                SourceColumn(index, get_column_letter(index), value)
                for index, value in enumerate(values, 1)
            ]
        finally:
            workbook.close()

    def preview(
        self,
        path: str | Path,
        sheet_name: str,
        header_row: int,
        limit: int = 50,
        header_row_2: int | None = None,
    ) -> tuple[list[SourceColumn], list[dict[str, Any]]]:
        headers = self.headers(path, sheet_name, header_row, header_row_2)
        rows = self.read_rows(path, sheet_name, header_row, max_rows=limit, header_row_2=header_row_2)
        return headers, rows

    def read_rows(
        self,
        path: str | Path,
        sheet_name: str,
        header_row: int,
        max_rows: int | None = None,
        header_row_2: int | None = None,
    ) -> list[dict[str, Any]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = workbook[sheet_name]
            header_values = self._header_values(ws, header_row, header_row_2)
            data_start = max(header_row, header_row_2 or header_row) + 1
            result: list[dict[str, Any]] = []
            for offset, values in enumerate(
                ws.iter_rows(min_row=data_start, values_only=True), start=data_start
            ):
                if max_rows is not None and len(result) >= max_rows:
                    break
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                row: dict[str, Any] = {"_row": offset}
                for index, value in enumerate(values, 1):
                    letter = get_column_letter(index)
                    row[letter] = value
                    if index <= len(header_values):
                        header = header_values[index - 1]
                        if header:
                            row[header] = value
                result.append(row)
            return result
        finally:
            workbook.close()
