from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from tools.vbdlis_excel_builder.models import FieldSchema, Household, ValidationIssue
from .diagnostic_logger import DiagnosticContext, DiagnosticLogger


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D9E2F3")
STAT_LABELS = {
    "households": "Số hộ",
    "source_households": "Số hộ trong nguồn",
    "source_people": "Số người trong nguồn",
    "source_parcels": "Số thửa trong nguồn",
    "households_skipped": "Số hộ đã bỏ qua",
    "people_skipped": "Số người không xuất (kể cả hộ bị bỏ)",
    "people_skipped_individually": "Số người bị bỏ riêng do CCCD sai",
    "parcels_skipped": "Số thửa không xuất do hộ bị bỏ",
    "people": "Số người",
    "summary_rows_skipped": "Số dòng tổng hợp đã bỏ qua",
    "parcels": "Số thửa",
    "duplicates": "Số dòng thửa trùng đã loại",
    "output_rows": "Số dòng kết quả",
    "with_gcn": "Số thửa có GCN",
    "without_gcn": "Số thửa chưa có GCN",
    "valid_cccd": "Số dòng CCCD hợp lệ",
    "missing_cccd": "Số dòng CCCD thiếu/không hợp lệ",
    "location_fallbacks": "Số thửa dùng địa chỉ thay Xứ đồng",
    "errors": "Số lỗi cần sửa",
    "warnings": "Số cảnh báo cần đối chiếu",
    "info": "Số thông tin xử lý",
}


def _style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    for column_cells in ws.columns:
        width = min(45, max(11, max(len(str(cell.value or "")) for cell in column_cells) + 2))
        ws.column_dimensions[column_cells[0].column_letter].width = width


class ReportWriter:
    def write(
        self,
        path: str | Path,
        stats: dict[str, Any],
        households: list[Household],
        issues: list[ValidationIssue],
        schemas: list[FieldSchema],
        output_rows: list[dict[str, Any]],
        diagnostic_context: DiagnosticContext | None = None,
    ) -> Path:
        output = Path(path)
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Tong_quan"
        summary.append(["Chỉ tiêu", "Giá trị"])
        for key, value in stats.items():
            if isinstance(value, (list, dict, tuple, set)):
                value = json.dumps(value, ensure_ascii=False, default=str)
            label = STAT_LABELS.get(key, key.replace("verify_", "Kiểm tra: "))
            summary.append([label, value])

        per_household = workbook.create_sheet("Theo_ho")
        per_household.append(["Hộ", "Dòng nguồn", "Số người được xuất", "Số thửa được xuất", "Số dòng kết quả"])
        for household in households:
            per_household.append(
                [
                    household.household_id,
                    household.source_row,
                    len(household.people),
                    len(household.parcels),
                    len(household.people) * len(household.parcels),
                ]
            )

        warnings = workbook.create_sheet("Canh_bao")
        warnings.append(["Mức độ", "Mã tra cứu", "Dòng nguồn", "Hộ", "Người", "Vấn đề", "Giá trị", "Ảnh hưởng", "Cách xử lý", "Vị trí và dữ liệu cần kiểm tra"])
        context = diagnostic_context or DiagnosticContext(households=households, schemas=schemas, output_rows=output_rows)
        for record in DiagnosticLogger.entries(context, issues):
            issue = record.source_issue
            warnings.append(
                [{"ERROR": "LỖI", "WARNING": "CẢNH BÁO", "INFO": "THÔNG TIN"}[issue.severity.value],
                 issue.code, issue.source_row, issue.household, issue.person, record.title + ": " + record.reason,
                 issue.value, record.effect, record.action, record.location + "\n" + "\n".join(record.evidence)]
            )

        gcn = workbook.create_sheet("GCN")
        gcn.append(
            [
                "Hộ", "Tờ/Thửa", "Có GCN", "Số phát hành", "Ngày cấp",
                "Số vào sổ", "Loại GCN", "Trường GCN còn thiếu", "Dòng nguồn", "Mâu thuẫn",
            ]
        )
        conflict_households = {issue.household for issue in issues if issue.code == "GCN_CONFLICT"}
        for household in households:
            for parcel in household.parcels:
                values = parcel.certificate.values
                expected = ("gcn_issue_number", "gcn_issue_date", "gcn_registry_number", "gcn_type")
                missing = [key for key in expected if parcel.certificate.has_data and not values.get(key)]
                gcn.append(
                    [
                        household.household_id,
                        f"{parcel.sheet_number}/{parcel.parcel_number}",
                        "Có" if parcel.certificate.has_data else "Chưa có",
                        values.get("gcn_issue_number", ""),
                        values.get("gcn_issue_date", ""),
                        values.get("gcn_registry_number", ""),
                        values.get("gcn_type", ""),
                        ", ".join(missing),
                        parcel.source_row,
                        "Có" if household.household_id in conflict_households else "Không",
                    ]
                )

        mapping = workbook.create_sheet("Mapping")
        mapping.append(["Mục", "Cột", "Tên", "Phân loại", "Mode", "Source", "Default", "Required", "Fallback"])
        for schema in schemas:
            mapping.append(
                [
                    schema.item,
                    schema.column,
                    schema.name,
                    schema.classification.value,
                    schema.mode,
                    schema.source,
                    schema.default,
                    "Có" if schema.required else "Không",
                    schema.fallback,
                ]
            )

        for ws in workbook.worksheets:
            _style_sheet(ws)
        summary.column_dimensions["A"].width = 58
        warnings.column_dimensions["F"].width = 65
        warnings.column_dimensions["G"].width = 45
        warnings.column_dimensions["H"].width = 60
        warnings.column_dimensions["I"].width = 75
        warnings.column_dimensions["J"].width = 75
        warnings.column_dimensions["B"].hidden = True
        gcn.column_dimensions["H"].width = 32
        mapping.column_dimensions["C"].width = 75
        mapping.column_dimensions["D"].width = 24
        mapping.column_dimensions["G"].width = 55
        for row in range(2, mapping.max_row + 1):
            mapping[f"C{row}"].alignment = Alignment(vertical="top", wrap_text=True)
            mapping[f"G{row}"].alignment = Alignment(vertical="top", wrap_text=True)
            mapping.row_dimensions[row].height = 32
        # These are audit sheets, not the official output template. Wrap long
        # explanations and grow rows so nontechnical readers can read them in Excel.
        for ws in workbook.worksheets:
            for row in ws.iter_rows():
                lines = 1
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    width = max(8, ws.column_dimensions[cell.column_letter].width - 2)
                    count = sum(max(1, math.ceil(len(line) / width)) for line in str(cell.value or "").split("\n"))
                    lines = max(lines, count)
                ws.row_dimensions[row[0].row].height = min(409, max(26, 16 * lines + 12))
        # Store source strings as literal text, never executable Excel formulas.
        for ws in workbook.worksheets:
            for row in ws:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
        workbook.save(output)
        workbook.close()
        return output
