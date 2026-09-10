from __future__ import annotations

import hashlib
from copy import copy
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from tools.vbdlis_excel_builder.models import FieldSchema, MappingProfile
from tools.vbdlis_excel_builder.utils.paths import unique_output_path

from .template_schema_reader import TemplateSchemaReader


TEXT_COLUMNS = {"A", "B", "C", "H", "I", "J", "K", "L", "M", "N", "T", "U", "X", "AX", "AY", "AZ", "BA", "BB"}
MIN_COLUMN_WIDTH = 12
MAX_COLUMN_WIDTH = 80


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class WorkbookWriter:
    def __init__(self, template_path: str | Path, schema_reader: TemplateSchemaReader):
        self.template_path = Path(template_path)
        self.schema_reader = schema_reader

    @staticmethod
    def _copy_row_style(source_ws, source_row: int, target_ws, target_row: int, max_col: int) -> None:
        if source_ws.row_dimensions[source_row].height is not None:
            target_ws.row_dimensions[target_row].height = source_ws.row_dimensions[source_row].height
        for col in range(1, max_col + 1):
            source = source_ws.cell(source_row, col)
            target = target_ws.cell(target_row, col)
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            target.protection = copy(source.protection)
            target.alignment = copy(source.alignment)

    @staticmethod
    def _fit_column_widths(ws, schemas: list[FieldSchema], first_data_row: int, last_data_row: int) -> None:
        """Make exported columns readable without changing the template workbook.

        The official template uses narrow widths intended for data entry.  That
        leaves values such as addresses, file names and land-use descriptions
        clipped in the generated workbook.  Measure both headers and output
        values, retain a sensible upper bound, and let Excel show the complete
        value in the formula bar when a value is exceptionally long.
        """
        for schema in schemas:
            column = schema.column
            max_length = 0
            for row in range(1, last_data_row + 1):
                value = ws[f"{column}{row}"].value
                if value in (None, ""):
                    continue
                lines = str(value).splitlines() or [""]
                max_length = max(max_length, max(len(line) for line in lines))
            # A small amount of padding prevents the final character touching
            # the cell border.  Never shrink a deliberately wider template
            # column.
            width = min(MAX_COLUMN_WIDTH, max(MIN_COLUMN_WIDTH, max_length + 2))
            current = ws.column_dimensions[column].width or 0
            ws.column_dimensions[column].width = max(current, width)

    def write(
        self,
        rows: list[dict[str, Any]],
        schemas: list[FieldSchema],
        profile: MappingProfile,
        output_folder: str | Path,
        output_filename: str,
        overwrite: bool = False,
        progress: Callable[[int, str], None] | None = None,
    ) -> tuple[Path, dict[str, Any]]:
        notify = progress or (lambda _value, _message: None)
        if not self.template_path.exists():
            raise FileNotFoundError(f"Không tìm thấy template: {self.template_path}")
        notify(68, "Đang chuẩn bị file mẫu VBDLIS")
        template_hash_before = _sha256(self.template_path)
        output_path = unique_output_path(output_folder, output_filename, overwrite)
        workbook = load_workbook(self.template_path, data_only=False)
        main_ws, item_row = self.schema_reader.find_schema_sheet(workbook)
        style_ws = workbook["CachNhap"] if "CachNhap" in workbook.sheetnames else main_ws
        style_row = 5 if style_ws.max_row >= 5 else item_row
        max_col = max(column_index_from_string(schema.column) for schema in schemas)

        if main_ws.max_row > item_row:
            main_ws.delete_rows(item_row + 1, main_ws.max_row - item_row)

        total_rows = len(rows)
        last_percent = -1
        for completed, (index, row) in enumerate(
            enumerate(rows, start=item_row + 1), start=1
        ):
            self._copy_row_style(style_ws, style_row, main_ws, index, max_col)
            for schema in schemas:
                cell = main_ws[f"{schema.column}{index}"]
                value = row.get(schema.column, "")
                cell.value = None if value == "" else value
                if schema.column in TEXT_COLUMNS:
                    cell.number_format = "@"
            percent = 70 + (18 * completed // max(total_rows, 1))
            if percent != last_percent:
                notify(percent, f"Đang ghi dòng {completed:,}/{total_rows:,}")
                last_percent = percent

        if not profile.keep_reference_sheets:
            for ws in list(workbook.worksheets):
                if ws.title != main_ws.title:
                    workbook.remove(ws)

        notify(90, "Đang căn chỉnh độ rộng cột")
        self._fit_column_widths(main_ws, schemas, item_row + 1, item_row + len(rows))
        notify(94, "Đang lưu file Excel")
        workbook.save(output_path)
        workbook.close()
        if _sha256(self.template_path) != template_hash_before:
            raise RuntimeError("Template gốc đã thay đổi ngoài ý muốn.")
        notify(96, "Đang kiểm tra file sau khi lưu")
        verification = self.verify(output_path, len(rows), main_ws.title, item_row)
        return output_path, verification

    @staticmethod
    def verify(path: Path, expected_rows: int, sheet_name: str, item_row: int) -> dict[str, Any]:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            ws = workbook[sheet_name]
            actual_rows = max(0, ws.max_row - item_row)
            stt_values = [ws.cell(row, 1).value for row in range(item_row + 1, item_row + 1 + expected_rows)]
            expected_stt = list(range(1, expected_rows + 1))
            normalized_stt = []
            for value in stt_values:
                try:
                    normalized_stt.append(int(str(value).strip()))
                except (TypeError, ValueError):
                    normalized_stt.append(value)
            cccd_text_ok = all(
                ws[f"I{row}"].number_format == "@"
                for row in range(item_row + 1, item_row + 1 + expected_rows)
            )
            cccd_required_ok = all(
                str(ws[f"I{row}"].value or "").strip() != ""
                for row in range(item_row + 1, item_row + 1 + expected_rows)
            )
            item49_ok = all(
                isinstance(ws[f"AX{row}"].value, str)
                and "-TBXN.pdf, " in ws[f"AX{row}"].value
                and ws[f"AX{row}"].value.endswith("-DDK.pdf")
                for row in range(item_row + 1, item_row + 1 + expected_rows)
            )
            area_copy_ok = all(
                ws[f"Y{row}"].value == ws[f"AA{row}"].value
                for row in range(item_row + 1, item_row + 1 + expected_rows)
            )
            formula_errors = []
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")):
                        formula_errors.append(cell.coordinate)
            checks = {
                "exists": path.exists(),
                "actual_rows": actual_rows,
                "expected_rows": expected_rows,
                "row_count_ok": actual_rows == expected_rows,
                "stt_ok": normalized_stt == expected_stt,
                "cccd_text_ok": cccd_text_ok,
                "cccd_required_ok": cccd_required_ok,
                "item49_ok": item49_ok,
                "area_copy_ok": area_copy_ok,
                "header_merges": len(ws.merged_cells.ranges),
                "formula_errors": formula_errors,
            }
            checks["pass"] = all(
                checks[key]
                for key in (
                    "exists", "row_count_ok", "stt_ok", "cccd_text_ok",
                    "cccd_required_ok", "item49_ok", "area_copy_ok",
                )
            ) and not formula_errors
            return checks
        finally:
            workbook.close()
