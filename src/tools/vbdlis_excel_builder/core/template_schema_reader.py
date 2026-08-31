from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from tools.vbdlis_excel_builder.models import FieldClassification, FieldSchema


class TemplateSchemaReader:
    def __init__(self, schema_config: str | Path | None = None):
        self.schema_config = Path(schema_config) if schema_config else None

    @staticmethod
    def _merged_value(ws, row: int, col: int) -> Any:
        value = ws.cell(row, col).value
        if value not in (None, ""):
            return value
        for merged in ws.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                return ws.cell(merged.min_row, merged.min_col).value
        return None

    def find_schema_sheet(self, workbook) -> tuple[Any, int]:
        best = None
        for ws in workbook.worksheets:
            for row in range(1, min(ws.max_row, 12) + 1):
                values = [ws.cell(row, col).value for col in range(1, min(ws.max_column, 100) + 1)]
                numeric_items = sum(1 for value in values if str(value).strip() in {str(i) for i in range(1, 51)})
                if best is None or numeric_items > best[0]:
                    best = (numeric_items, ws, row)
        if best is None or best[0] < 10:
            raise ValueError("Không tìm thấy dòng Mục trong template VBDLIS.")
        return best[1], best[2]

    def read(self, path: str | Path) -> tuple[str, int, list[FieldSchema]]:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            ws, item_row = self.find_schema_sheet(workbook)
            configured: dict[str, dict[str, Any]] = {}
            if self.schema_config and self.schema_config.exists():
                payload = json.loads(self.schema_config.read_text(encoding="utf-8"))
                configured = {item["field_id"]: item for item in payload.get("fields", [])}
            fields: list[FieldSchema] = []
            for col in range(1, ws.max_column + 1):
                item = ws.cell(item_row, col).value
                header_parts = []
                for row in range(1, item_row):
                    value = self._merged_value(ws, row, col)
                    if value not in (None, ""):
                        text = " ".join(str(value).split())
                        if text not in header_parts:
                            header_parts.append(text)
                if item in (None, "") and not header_parts:
                    continue
                column = get_column_letter(col)
                item_text = "" if item is None else str(item).strip()
                field_id = f"MUC_{item_text}" if item_text.isdigit() else f"COL_{column}"
                base = FieldSchema(
                    field_id=field_id,
                    item=item_text,
                    column=column,
                    name=" / ".join(header_parts),
                    classification=FieldClassification.TEMPLATE_OPTIONAL,
                )
                if field_id in configured:
                    config = dict(configured[field_id])
                    config["column"] = column
                    config["item"] = item_text
                    config["name"] = base.name
                    base = FieldSchema.from_dict(config)
                fields.append(base)
            return ws.title, item_row, fields
        finally:
            workbook.close()
