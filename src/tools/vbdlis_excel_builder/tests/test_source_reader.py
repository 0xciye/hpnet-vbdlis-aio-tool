from openpyxl import Workbook

from tools.vbdlis_excel_builder.core.source_reader import SourceColumn, SourceReader
from tools.vbdlis_excel_builder.core.auto_mapper import AutoMapper
from tools.vbdlis_excel_builder.core.household_parser import HouseholdParser
from tools.vbdlis_excel_builder.models import MappingProfile


def test_source_reader_supports_two_level_headers_and_merged_groups(tmp_path):
    source = tmp_path / "two-level.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Thông tin thửa"
    sheet["A2"] = "Số tờ"
    sheet["B2"] = "Số thửa"
    sheet["C1"] = "Họ tên"
    sheet["A3"] = 12
    sheet["B3"] = 34
    sheet["C3"] = "Nguyễn Văn A"
    workbook.save(source)

    reader = SourceReader()
    headers, rows = reader.preview(source, "Data", 1, 10, header_row_2=2)

    assert [column.header for column in headers[:3]] == [
        "Thông tin thửa - Số tờ",
        "Thông tin thửa - Số thửa",
        "Họ tên",
    ]
    assert rows == [{
        "_row": 3,
        "A": 12,
        "Thông tin thửa - Số tờ": 12,
        "B": 34,
        "Thông tin thửa - Số thửa": 34,
        "C": "Nguyễn Văn A",
        "Họ tên": "Nguyễn Văn A",
    }]


def test_auto_mapper_recognizes_ten_ho_as_person_name():
    headers = SourceReader()._merge_header_values(["", "", ""], ["STT", "Tên hộ", "CCCD"])
    columns = [SourceColumn(index, letter, header)
               for index, (letter, header) in enumerate(zip(("A", "B", "C"), headers), 1)]
    assert AutoMapper().suggest(columns)["person_name"] == "B"


def test_formula_stt_without_cached_value_starts_a_new_household(tmp_path):
    source = tmp_path / "formula-stt.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["STT", "Họ tên", "CCCD", "Số tờ", "Số thửa", "Diện tích"])
    sheet.append([1, "Nguyễn Văn A", "031080001234", 10, 100, 200])
    sheet.append(["=MAX($A$2:A2)+1", "Trần Thị B", "031185001235", 11, 101, 300])
    sheet.append([None, "Lê Văn C", "031090001236", None, None, None])
    workbook.save(source)

    rows = SourceReader().read_rows(
        source,
        "Data",
        1,
        formula_presence_columns={"A"},
    )
    profile = MappingProfile(source_mapping={
        "household_stt": "A",
        "person_name": "B",
        "cccd": "C",
        "sheet_number": "D",
        "parcel_number": "E",
        "area": "F",
    })
    households, _, _ = HouseholdParser().parse(rows, profile)

    assert rows[1]["A"] is None
    assert rows[1]["_formula_cells"] == ("A",)
    assert [[person.is_head for person in household.people] for household in households] == [
        [True],
        [True, False],
    ]
