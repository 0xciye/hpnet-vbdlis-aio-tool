from __future__ import annotations

import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from tools.vbdlis_excel_builder.core.golden_reference_analyzer import GoldenReferenceAnalyzer
from tools.vbdlis_excel_builder.core.template_schema_reader import TemplateSchemaReader
from tools.vbdlis_excel_builder.core.workbook_writer import WorkbookWriter
from tools.vbdlis_excel_builder.models import MappingProfile


PROJECT = Path(__file__).resolve().parents[1]
WORKSPACE = PROJECT.parent
OFFICIAL = PROJECT / "resources" / "BieuMauThuThapThongTinGiayChungNhan.xlsx"
GOLDEN = WORKSPACE / "Copy of tk1, tk2 (chạy lượt 2).xlsx"
SCHEMA = PROJECT / "config" / "template_schema.json"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class GoldenRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not GOLDEN.exists():
            raise unittest.SkipTest("Golden reference chỉ dùng trong môi trường phát triển.")

    def test_golden_rules(self):
        result = GoldenReferenceAnalyzer().analyze(GOLDEN)
        self.assertEqual(result["rows"], 174)
        self.assertEqual(result["active_field_count"], 24)
        self.assertTrue(result["item49_pass"])
        self.assertTrue(result["roles_pass"])
        self.assertTrue(result["gcn_blank_pass"])
        self.assertTrue(result["stt_pass"])

    def test_round_trip_critical_values_and_template_immutability(self):
        reader = TemplateSchemaReader(SCHEMA)
        main_sheet, _, schemas = reader.read(OFFICIAL)
        workbook = load_workbook(GOLDEN, data_only=False)
        ws = workbook[workbook.sheetnames[0]]
        rows = []
        for row_index in range(5, ws.max_row + 1):
            rows.append(
                {
                    get_column_letter(col): ws.cell(row_index, col).value
                    for col in range(1, 62)
                }
            )
        workbook.close()
        official_hash = sha256(OFFICIAL)
        with tempfile.TemporaryDirectory() as directory:
            output, verification = WorkbookWriter(OFFICIAL, reader).write(
                rows,
                schemas,
                MappingProfile(keep_reference_sheets=False),
                directory,
                "golden_round_trip.xlsx",
            )
            self.assertTrue(verification["pass"], verification)
            self.assertEqual(sha256(OFFICIAL), official_hash)
            produced = load_workbook(output, data_only=False)
            out_ws = produced[main_sheet]
            golden_wb = load_workbook(GOLDEN, data_only=False)
            golden_ws = golden_wb[golden_wb.sheetnames[0]]
            for col in ("A", "B", "C", "H", "I", "J", "N", "T", "U", "X", "Y", "AA", "AX", "AY", "AZ", "BA", "BB"):
                self.assertEqual(
                    [out_ws[f"{col}{row}"].value for row in range(5, 179)],
                    [golden_ws[f"{col}{row}"].value for row in range(5, 179)],
                    col,
                )
            self.assertEqual(len(produced.sheetnames), 1)
            self.assertGreater(len(out_ws.merged_cells.ranges), 0)
            produced.close()
            golden_wb.close()


if __name__ == "__main__":
    unittest.main()
