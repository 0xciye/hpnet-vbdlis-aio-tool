from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tools.vbdlis_excel_builder.core import BuilderService
from tools.vbdlis_excel_builder.models import MappingProfile


PROJECT = Path(__file__).resolve().parents[1]


class IntegrationExportTests(unittest.TestCase):
    def test_source_to_output_and_audit_report(self):
        with tempfile.TemporaryDirectory() as directory:
            temp = Path(directory)
            source = temp / "nguon.xlsx"
            workbook = Workbook()
            ws = workbook.active
            ws.title = "Du lieu"
            ws.append(["STT", "Họ và tên", "CCCD", "Ngày sinh", "Số tờ", "Số thửa", "Diện tích", "Xứ đồng", "Số GCN"])
            ws.append([1, "Nguyễn Văn A", "031080001234", "1980", 10, 300, 100, "", ""])
            ws.append([None, "Trần Thị B", "031185001235", "1985", 20, 400, 200, "Xứ Hai", "CS001"])
            ws.append([None, "TỔNG DT", None, None, None, None, 300])
            workbook.save(source)
            workbook.close()

            service = BuilderService(
                PROJECT / "resources" / "BieuMauThuThapThongTinGiayChungNhan.xlsx",
                PROJECT / "config" / "template_schema.json",
                diagnostics_dir=temp / "logs",
            )
            profile = MappingProfile(
                commune_code="10930",
                address="Thôn Test, Xã Test, Thành phố Hải Phòng",
                output_folder=str(temp),
                output_filename="ket_qua.xlsx",
                export_audit_report=True,
                source_mapping={
                    "household_stt": "A",
                    "person_name": "B",
                    "cccd": "C",
                    "birth_date": "D",
                    "sheet_number": "E",
                    "parcel_number": "F",
                    "area": "G",
                    "land_location": "H",
                    "gcn_issue_number": "I",
                },
            )
            result = service.process(source, "Du lieu", 1, profile)
            self.assertTrue(result.can_export, [issue.to_dict() for issue in result.issues])
            self.assertEqual(result.stats["output_rows"], 4)
            self.assertEqual(result.stats["people"], 2)
            self.assertEqual(result.stats["summary_rows_skipped"], 1)
            output, report, verification = service.export(result, profile, temp, "ket_qua.xlsx")
            self.assertTrue(verification["pass"])
            self.assertTrue(output.exists())
            self.assertIsNotNone(report)
            self.assertTrue(report.exists())
            produced = load_workbook(output, read_only=True)
            self.assertEqual([produced.active[f"K{row}"].value for row in range(5, 9)],
                             ["Nam", "Nữ", "Nam", "Nữ"])
            self.assertEqual([produced.active[f"H{row}"].value for row in range(5, 9)],
                             ["NGUYỄN VĂN A", "TRẦN THỊ B", "NGUYỄN VĂN A", "TRẦN THỊ B"])
            produced.close()
            report_wb = load_workbook(report, read_only=True)
            self.assertEqual(report_wb.sheetnames, ["Tong_quan", "Theo_ho", "Canh_bao", "GCN", "Mapping"])
            skipped = [row for row in report_wb["Canh_bao"].iter_rows(min_row=2, values_only=True)
                       if row[1] == "SUMMARY_ROW_SKIPPED"]
            self.assertEqual([(row[0], row[2], row[6]) for row in skipped], [("THÔNG TIN", 4, "TỔNG DT")])
            summary = dict(report_wb["Tong_quan"].iter_rows(min_row=2, max_col=2, values_only=True))
            self.assertEqual(summary["Số dòng tổng hợp đã bỏ qua"], 1)
            report_wb.close()


if __name__ == "__main__":
    unittest.main()
