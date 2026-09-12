from __future__ import annotations

import json
import logging
import shutil
from collections import defaultdict
from collections.abc import Callable, Iterable
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import column_index_from_string, get_column_letter

from .documents import DocumentScanner, best_document_candidates
from .messages import describe_issues, recommended_actions
from .models import (
    DocumentCompleteness,
    DocumentRecord,
    DocumentType,
    ExcelConfig,
    ParcelRecord,
    RunConfig,
    SignatureReadiness,
    SignatureStatus,
    UploadRow,
    ValidationMode,
    ValidationResult,
    WorkflowStatus,
)
from .normalization import (
    normalize_document_key,
    normalize_identifier,
    normalize_name,
    normalize_role,
    raw_text,
    sanitize_windows_name,
)

Progress = Callable[[int, str], None]


def sheet_names(path: str | Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def headers(config: ExcelConfig) -> list[tuple[str, str]]:
    workbook = load_workbook(config.path, read_only=False, data_only=False)
    try:
        ws = workbook[config.sheet]
        values: list[tuple[str, str]] = []
        second_row = config.header_row_2
        if config.header_row > ws.max_row:
            raise ValueError("Dòng tiêu đề tầng 1 nằm ngoài vùng dữ liệu của sheet.")
        if second_row is not None and second_row <= config.header_row:
            raise ValueError("Dòng tiêu đề tầng 2 phải nằm sau tầng 1.")
        if second_row is not None and second_row > ws.max_row:
            raise ValueError("Dòng tiêu đề tầng 2 nằm ngoài vùng dữ liệu của sheet.")
        top_merges = [
            cell_range
            for cell_range in ws.merged_cells.ranges
            if cell_range.min_row <= config.header_row <= cell_range.max_row
        ]
        for index in range(1, ws.max_column + 1):
            top = raw_text(ws.cell(config.header_row, index).value)
            if not top:
                merged = next(
                    (cell_range for cell_range in top_merges if cell_range.min_col <= index <= cell_range.max_col),
                    None,
                )
                if merged is not None:
                    top = raw_text(ws.cell(merged.min_row, merged.min_col).value)
            if second_row is None:
                label = top
            else:
                bottom = raw_text(ws.cell(second_row, index).value)
                label = " — ".join(part for part in (top, bottom) if part)
            values.append((get_column_letter(index), label))
        return values
    finally:
        workbook.close()


def _row_value(row: dict[str, Any], column: str) -> Any:
    return row.get(column.upper()) if column else None


def _read_rows(config: ExcelConfig) -> list[dict[str, Any]]:
    values_book = load_workbook(config.path, read_only=True, data_only=True)
    formula_book = load_workbook(config.path, read_only=True, data_only=False)
    try:
        values_ws = values_book[config.sheet]
        formula_ws = formula_book[config.sheet]
        rows: list[dict[str, Any]] = []
        header_end = max(config.header_row, config.header_row_2 or config.header_row)
        value_rows = values_ws.iter_rows(
            min_row=header_end + 1,
            max_row=values_ws.max_row,
            max_col=values_ws.max_column,
            values_only=True,
        )
        formula_rows = formula_ws.iter_rows(
            min_row=header_end + 1,
            max_row=formula_ws.max_row,
            max_col=formula_ws.max_column,
            values_only=True,
        )
        for row_number, (values, formulas) in enumerate(
            zip(value_rows, formula_rows, strict=True),
            header_end + 1,
        ):
            row: dict[str, Any] = {"_row": row_number}
            nonblank = False
            for column, (value, formula) in enumerate(zip(values, formulas, strict=True), 1):
                if value is None and isinstance(formula, str) and formula.startswith("="):
                    value = formula
                if value not in (None, ""):
                    nonblank = True
                row[get_column_letter(column)] = value
            if nonblank:
                rows.append(row)
        return rows
    finally:
        values_book.close()
        formula_book.close()


def _required(config: RunConfig) -> None:
    required_files = [(
            "Excel VBDLIS",
            config.upload,
            (("person", "Tên người"), ("role", "Vai trò"), ("sheet", "Số tờ"), ("parcel", "Số thửa")),
        )]
    if config.validation_mode == ValidationMode.FULL_SOURCE_COMPARE:
        if config.source is None:
            raise ValueError("Chế độ so sánh đầy đủ cần chọn Excel nguồn.")
        required_files.append((
            "Excel nguồn", config.source,
            (("owner", "Tên chủ hộ"), ("sheet", "Số tờ"), ("parcel", "Số thửa")),
        ))
    for label, excel, fields in required_files:
        if not excel.path.is_file():
            raise ValueError(f"Không tìm thấy {label}. Hãy bấm Chọn và chọn lại đúng file: {excel.path}")
        missing = [field_name for key, field_name in fields if not excel.columns.get(key)]
        if missing:
            raise ValueError(
                f"{label} chưa chọn đủ các cột bắt buộc: {', '.join(missing)}. "
                "Hãy chọn cột tương ứng trong phần Mapping cột."
            )
        if excel.header_row_2 is not None and excel.header_row_2 <= excel.header_row:
            raise ValueError(f"{label}: dòng tiêu đề tầng 2 phải nằm sau tầng 1.")
    document_mapping = config.upload.columns
    if not document_mapping.get("combined") and not (
        document_mapping.get("tbxn") and document_mapping.get("ddk")
    ):
        raise ValueError(
            "Chưa chọn cột chứa tên tài liệu. Hãy chọn cả cột TBXN và cột DDK, "
            "hoặc chọn một Cột tài liệu dùng chung."
        )
    try:
        start = column_index_from_string(config.export_start.upper())
        end = column_index_from_string(config.export_end.upper())
    except ValueError as error:
        raise ValueError("Phạm vi export phải là tên cột Excel hợp lệ, ví dụ A đến AZ.") from error
    if start > end:
        raise ValueError("Cột bắt đầu export phải đứng trước hoặc bằng cột kết thúc.")


def _append_issue(record: ParcelRecord, issue: str) -> None:
    if issue not in record.issues:
        record.issues.append(issue)


def _signature_issue(record: ParcelRecord, document: DocumentRecord, prefix: str) -> None:
    signature = document.signature
    if signature is None:
        _append_issue(record, "SIGNATURE_CHECK_ERROR")
    elif signature.file_access_error:
        _append_issue(record, "FILE_ACCESS_ERROR")
        _append_issue(record, "SIGNATURE_CHECK_ERROR")
    elif signature.invalid_pdf:
        _append_issue(record, "INVALID_PDF")
        _append_issue(record, "SIGNATURE_CHECK_ERROR")
    elif signature.status == SignatureStatus.UNSIGNED:
        _append_issue(record, f"{prefix}_UNSIGNED")
    elif signature.status == SignatureStatus.SIGNATURE_ERROR:
        _append_issue(record, "SIGNATURE_CHECK_ERROR")
    elif signature.status != SignatureStatus.SIGNED_VALID:
        _append_issue(record, f"{prefix}_SIGNATURE_INVALID")


def _references_document(record: ParcelRecord, document: DocumentRecord, kind: DocumentType) -> bool:
    expected = normalize_document_key(document.path.name)
    if not expected:
        return False
    for row in record.upload_rows:
        value = row.combined_reference or (row.tbxn_reference if kind == DocumentType.TBXN else row.ddk_reference)
        if expected in normalize_document_key(value):
            return True
    return False


def evaluate_record(
    record: ParcelRecord,
    mode: ValidationMode = ValidationMode.FULL_SOURCE_COMPARE,
) -> ParcelRecord:
    record.issues = [issue for issue in record.issues if issue != "VALID"]
    if mode == ValidationMode.FULL_SOURCE_COMPARE and not record.source_rows:
        _append_issue(record, "SOURCE_PARCEL_NOT_FOUND")
        _append_issue(record, "VBDLIS_ONLY")
    if mode == ValidationMode.FULL_SOURCE_COMPARE and not record.upload_rows:
        _append_issue(record, "VBDLIS_PARCEL_NOT_FOUND")
        _append_issue(record, "SOURCE_ONLY")
    if record.upload_rows and not record.owner_normalized and not record.household_id:
        _append_issue(record, "HOUSEHOLD_NOT_FOUND")
    if not record.tbxn_documents and not record.ddk_documents:
        _append_issue(record, "MISSING_BOTH")
        record.document_completeness = DocumentCompleteness.MISSING_BOTH
    elif not record.tbxn_documents:
        _append_issue(record, "MISSING_TBXN")
        record.document_completeness = DocumentCompleteness.MISSING_TBXN
    elif not record.ddk_documents:
        _append_issue(record, "MISSING_DDK")
        record.document_completeness = DocumentCompleteness.MISSING_DDK
    else:
        record.document_completeness = DocumentCompleteness.COMPLETE
    if len(record.tbxn_documents) > 1:
        _append_issue(record, "DUPLICATE_TBXN")
        _append_issue(record, "REVIEW_REQUIRED")
        record.document_completeness = DocumentCompleteness.DUPLICATE_DOCUMENT
    if len(record.ddk_documents) > 1:
        _append_issue(record, "DUPLICATE_DDK")
        _append_issue(record, "REVIEW_REQUIRED")
        record.document_completeness = DocumentCompleteness.DUPLICATE_DOCUMENT
    if len(record.tbxn_documents) == 1:
        _signature_issue(record, record.tbxn_documents[0], "TBXN")
        if record.upload_rows and not _references_document(record, record.tbxn_documents[0], DocumentType.TBXN):
            _append_issue(record, "DATA_MISMATCH")
    if len(record.ddk_documents) == 1:
        _signature_issue(record, record.ddk_documents[0], "DDK")
        if record.upload_rows and not _references_document(record, record.ddk_documents[0], DocumentType.DDK):
            _append_issue(record, "DATA_MISMATCH")
    if {"TBXN_UNSIGNED", "DDK_UNSIGNED"} <= set(record.issues):
        _append_issue(record, "BOTH_UNSIGNED")
    invalid_kinds = {issue for issue in record.issues if issue.endswith("SIGNATURE_INVALID")}
    if {"TBXN_SIGNATURE_INVALID", "DDK_SIGNATURE_INVALID"} <= invalid_kinds:
        _append_issue(record, "BOTH_SIGNATURE_INVALID")

    if "INVALID_PDF" in record.issues:
        record.document_completeness = DocumentCompleteness.INVALID_DOCUMENT
    elif "AMBIGUOUS_MATCH" in record.issues and record.document_completeness == DocumentCompleteness.COMPLETE:
        record.document_completeness = DocumentCompleteness.AMBIGUOUS_DOCUMENT

    if "BOTH_UNSIGNED" in record.issues:
        record.signature_readiness = SignatureReadiness.BOTH_UNSIGNED
    elif "TBXN_UNSIGNED" in record.issues:
        record.signature_readiness = SignatureReadiness.TBXN_UNSIGNED
    elif "DDK_UNSIGNED" in record.issues:
        record.signature_readiness = SignatureReadiness.DDK_UNSIGNED
    elif "BOTH_SIGNATURE_INVALID" in record.issues:
        record.signature_readiness = SignatureReadiness.BOTH_SIGNATURE_INVALID
    elif "TBXN_SIGNATURE_INVALID" in record.issues:
        record.signature_readiness = SignatureReadiness.TBXN_SIGNATURE_INVALID
    elif "DDK_SIGNATURE_INVALID" in record.issues:
        record.signature_readiness = SignatureReadiness.DDK_SIGNATURE_INVALID
    elif "SIGNATURE_CHECK_ERROR" in record.issues:
        record.signature_readiness = SignatureReadiness.SIGNATURE_CHECK_ERROR
    elif record.document_completeness == DocumentCompleteness.COMPLETE:
        record.signature_readiness = SignatureReadiness.SIGNED_READY
    else:
        record.signature_readiness = SignatureReadiness.NOT_CHECKED

    review_issues = {
        "REVIEW_REQUIRED", "AMBIGUOUS_MATCH", "DUPLICATE_TBXN", "DUPLICATE_DDK",
        "UNKNOWN_ROLE", "HOUSEHOLD_WITHOUT_HEAD", "MULTIPLE_HOUSEHOLD_HEADS",
        "MISSING_NAME", "MISSING_ROLE", "MISSING_SHEET", "MISSING_PARCEL",
    }
    data_issues = {
        "DATA_MISMATCH", "PARCEL_MISMATCH", "AX_DATA_MISMATCH", "AX_PARSE_ERROR",
        "SOURCE_PARCEL_NOT_FOUND", "VBDLIS_PARCEL_NOT_FOUND",
    }
    if review_issues.intersection(record.issues):
        record.workflow_status = WorkflowStatus.REVIEW_REQUIRED
    elif record.document_completeness in {
        DocumentCompleteness.MISSING_TBXN,
        DocumentCompleteness.MISSING_DDK,
        DocumentCompleteness.MISSING_BOTH,
    }:
        record.workflow_status = WorkflowStatus.MISSING_DOCUMENT
    elif record.signature_readiness in {
        SignatureReadiness.TBXN_SIGNATURE_INVALID,
        SignatureReadiness.DDK_SIGNATURE_INVALID,
        SignatureReadiness.BOTH_SIGNATURE_INVALID,
        SignatureReadiness.SIGNATURE_CHECK_ERROR,
    } or "INVALID_PDF" in record.issues:
        record.workflow_status = WorkflowStatus.SIGNATURE_INVALID
    elif data_issues.intersection(record.issues):
        record.workflow_status = WorkflowStatus.DATA_ERROR
    elif record.signature_readiness in {
        SignatureReadiness.TBXN_UNSIGNED,
        SignatureReadiness.DDK_UNSIGNED,
        SignatureReadiness.BOTH_UNSIGNED,
    }:
        record.workflow_status = WorkflowStatus.NEED_SIGNATURE
    else:
        record.workflow_status = WorkflowStatus.READY_TO_UPLOAD
    if not record.issues:
        record.issues.append("VALID")
    return record


class ValidationService:
    def __init__(self, scanner: DocumentScanner | None = None) -> None:
        self.scanner = scanner or DocumentScanner()

    def preview(self, config: RunConfig) -> dict[str, int]:
        _required(config)
        upload_rows = _read_rows(config.upload)
        head_rows = sum(
            normalize_role(_row_value(row, config.upload.columns.get("role", ""))) == "HEAD"
            for row in upload_rows
        )
        if config.validation_mode == ValidationMode.VBDLIS_ONLY:
            records = self._build_vbdlis_only_records(config.upload, upload_rows)
            source_rows: list[dict[str, Any]] = []
        else:
            assert config.source is not None
            source_rows = _read_rows(config.source)
            records, _ = self._build_source_records(config.source, source_rows)
        documents = self.scanner.scan(
            config.tbxn_folder, config.ddk_folder, config.shared_document_folder,
            validate_signatures=False,
        )
        return {
            "vbdlis_rows": len(upload_rows),
            "source_rows": len(source_rows),
            "source_households": len({record.owner_normalized for record in records if record.owner_normalized}),
            "source_parcels": len(records),
            "tbxn_pdfs": sum(document.document_type == DocumentType.TBXN for document in documents),
            "ddk_pdfs": sum(document.document_type == DocumentType.DDK for document in documents),
            "unknown_pdfs": sum(document.document_type == DocumentType.UNKNOWN for document in documents),
            "head_rows": head_rows,
        }

    def run(self, config: RunConfig, progress: Progress | None = None) -> ValidationResult:
        notify = progress or (lambda _percent, _message: None)
        _required(config)
        notify(3, "1/10 — Đang đọc Excel")
        upload_rows = _read_rows(config.upload)
        source_rows = [] if config.validation_mode == ValidationMode.VBDLIS_ONLY else _read_rows(config.source)  # type: ignore[arg-type]
        notify(12, "2/10 — Đang chuẩn hóa dữ liệu")
        if config.validation_mode == ValidationMode.VBDLIS_ONLY:
            records = self._build_vbdlis_only_records(config.upload, upload_rows)
        else:
            assert config.source is not None
            records, indexes = self._build_source_records(config.source, source_rows)
            notify(20, "3/10 — Đang xây dựng hộ gia đình")
            self._attach_upload_rows(config.upload, upload_rows, records, indexes)
        notify(30, "4/10 — Đang lập chỉ mục thửa")
        notify(36, "5–7/10 — Đang quét PDF và xác minh chữ ký số")
        documents = self.scanner.scan(
            config.tbxn_folder,
            config.ddk_folder,
            config.shared_document_folder,
            progress=lambda done, total: notify(36 + (34 * done // max(total, 1)), f"7/10 — Xác minh chữ ký {done:,}/{total:,}"),
        )
        notify(72, "8/10 — Đang ghép TBXN/DDK theo từng thửa")
        self._attach_documents(records, documents)
        notify(80, "9/10 — Đang kiểm tra dữ liệu VBDLIS")
        for record in records:
            self._validate_ax(record)
            evaluate_record(record, config.validation_mode)
        result = ValidationResult(
            records,
            len(upload_rows),
            len(source_rows),
            documents=documents,
            validation_mode=config.validation_mode,
            scan_stats=self.scanner.last_stats,
        )
        notify(88, "10/10 — Đang xuất kết quả")
        self._export(config, result)
        notify(100, "Hoàn thành kiểm tra")
        return result

    @staticmethod
    def _build_vbdlis_only_records(config: ExcelConfig, rows: list[dict[str, Any]]) -> list[ParcelRecord]:
        records: list[ParcelRecord] = []
        by_key: dict[tuple[str, str, str], ParcelRecord] = {}
        current_owner = current_id = ""
        head_names: dict[str, set[str]] = defaultdict(set)
        for row in rows:
            person_raw = raw_text(_row_value(row, config.columns["person"]))
            person = normalize_name(person_raw)
            role_raw = raw_text(_row_value(row, config.columns["role"]))
            role = normalize_role(role_raw)
            sheet_raw = raw_text(_row_value(row, config.columns["sheet"]))
            parcel_raw = raw_text(_row_value(row, config.columns["parcel"]))
            sheet = normalize_identifier(sheet_raw, "sheet")
            parcel = normalize_identifier(parcel_raw, "parcel")
            household_id = normalize_identifier(_row_value(row, config.columns.get("household_id", "")))
            if household_id:
                current_id = household_id
            if role == "HEAD" and person:
                current_owner = person_raw
                if not household_id:
                    current_id = ""
                head_names[household_id or person].add(person)
            owner = normalize_name(current_owner)
            household_key = current_id or owner
            # Keep malformed rows individually traceable instead of merging
            # unrelated rows that happen to share empty keys.
            row_fallback = f"ROW:{int(row['_row'])}"
            key = (household_key or row_fallback, sheet or row_fallback, parcel or row_fallback)
            record = by_key.get(key)
            if record is None:
                record = ParcelRecord(current_owner, owner, sheet_raw, sheet, parcel_raw, parcel, current_id)
                record.matching_rule = "VBDLIS_HOUSEHOLD+SHEET+PARCEL"
                record.match_evidence = f"Hộ={current_owner}; tờ={sheet_raw}; thửa={parcel_raw}"
                records.append(record); by_key[key] = record
            upload = UploadRow(
                int(row["_row"]), person_raw, person, role, sheet_raw, sheet, parcel_raw, parcel,
                current_id,
                raw_text(_row_value(row, config.columns.get("tbxn", ""))),
                raw_text(_row_value(row, config.columns.get("ddk", ""))),
                raw_text(_row_value(row, config.columns.get("combined", ""))),
                row,
            )
            record.upload_rows.append(upload)
            if not person:
                _append_issue(record, "MISSING_NAME")
            if not role_raw:
                _append_issue(record, "MISSING_ROLE")
                _append_issue(record, "REVIEW_REQUIRED")
            if not sheet:
                _append_issue(record, "MISSING_SHEET")
            if not parcel:
                _append_issue(record, "MISSING_PARCEL")
            if role == "UNKNOWN":
                _append_issue(record, "UNKNOWN_ROLE")
                _append_issue(record, "REVIEW_REQUIRED")
            if not household_key:
                _append_issue(record, "HOUSEHOLD_WITHOUT_HEAD")
                _append_issue(record, "REVIEW_REQUIRED")
        for record in records:
            household_key = record.household_id or record.owner_normalized
            if household_key and not head_names.get(household_key):
                _append_issue(record, "HOUSEHOLD_WITHOUT_HEAD")
                _append_issue(record, "REVIEW_REQUIRED")
            elif len(head_names.get(household_key, set())) > 1:
                _append_issue(record, "MULTIPLE_HOUSEHOLD_HEADS")
                _append_issue(record, "REVIEW_REQUIRED")
        return records

    @staticmethod
    def _validate_ax(record: ParcelRecord) -> None:
        values = [
            row.combined_reference or " | ".join(part for part in (row.tbxn_reference, row.ddk_reference) if part)
            for row in record.upload_rows
        ]
        record.ax_raw = " | ".join(dict.fromkeys(value for value in values if value))
        if not record.ax_raw:
            record.ax_status = "EMPTY"
            _append_issue(record, "AX_EMPTY")
            return
        folded = record.ax_raw.casefold()
        mentions_tbxn, mentions_ddk = "tbxn" in folded, "ddk" in folded
        if not mentions_tbxn and not mentions_ddk:
            record.ax_status = "PARSE_ERROR"
            _append_issue(record, "AX_PARSE_ERROR")
            _append_issue(record, "REVIEW_REQUIRED")
            return
        mismatch = False
        if record.tbxn_documents and not any(_references_document(record, doc, DocumentType.TBXN) for doc in record.tbxn_documents):
            _append_issue(record, "AX_MISSING_TBXN_REFERENCE"); mismatch = True
        if record.ddk_documents and not any(_references_document(record, doc, DocumentType.DDK) for doc in record.ddk_documents):
            _append_issue(record, "AX_MISSING_DDK_REFERENCE"); mismatch = True
        if mentions_tbxn and not record.tbxn_documents:
            _append_issue(record, "AX_REFERENCES_MISSING_TBXN"); mismatch = True
        if mentions_ddk and not record.ddk_documents:
            _append_issue(record, "AX_REFERENCES_MISSING_DDK"); mismatch = True
        if mismatch:
            _append_issue(record, "AX_DATA_MISMATCH")
            record.ax_status = "MISMATCH"
        else:
            record.ax_status = "MATCH"

    @staticmethod
    def _build_source_records(
        config: ExcelConfig, rows: list[dict[str, Any]]
    ) -> tuple[list[ParcelRecord], dict[str, Any]]:
        records: list[ParcelRecord] = []
        by_owner: dict[tuple[str, str, str], ParcelRecord] = {}
        by_id: dict[tuple[str, str, str], ParcelRecord] = {}
        by_parcel: dict[tuple[str, str], list[ParcelRecord]] = defaultdict(list)
        current_owner = current_id = ""
        for row in rows:
            owner_value = _row_value(row, config.columns.get("owner", ""))
            household_id_value = _row_value(row, config.columns.get("household_id", ""))
            if raw_text(owner_value):
                current_owner = raw_text(owner_value)
                if not raw_text(household_id_value):
                    current_id = ""
            if raw_text(household_id_value):
                current_id = normalize_identifier(household_id_value)
            sheet_raw = raw_text(_row_value(row, config.columns["sheet"]))
            parcel_raw = raw_text(_row_value(row, config.columns["parcel"]))
            sheet = normalize_identifier(sheet_raw, "sheet")
            parcel = normalize_identifier(parcel_raw, "parcel")
            # Source workbooks commonly contain subtotal rows such as "Tổng DT".
            # They are not parcels and must not become synthetic validation errors.
            if not sheet and not parcel:
                continue
            owner = normalize_name(current_owner)
            if not sheet or not parcel or not owner:
                record = ParcelRecord(current_owner, owner, sheet_raw, sheet, parcel_raw, parcel, current_id)
                record.source_rows.append(int(row["_row"]))
                _append_issue(record, "DATA_MISMATCH")
                records.append(record)
                continue
            key = (owner, sheet, parcel)
            record = by_owner.get(key)
            if record is None:
                record = ParcelRecord(current_owner, owner, sheet_raw, sheet, parcel_raw, parcel, current_id)
                records.append(record)
                by_owner[key] = record
                by_parcel[(sheet, parcel)].append(record)
                if current_id:
                    by_id[(current_id, sheet, parcel)] = record
            elif int(row["_row"]) not in record.source_rows:
                _append_issue(record, "DATA_MISMATCH")
            record.source_rows.append(int(row["_row"]))
        return records, {"owner": by_owner, "id": by_id, "parcel": by_parcel}

    @staticmethod
    def _attach_upload_rows(
        config: ExcelConfig,
        rows: list[dict[str, Any]],
        records: list[ParcelRecord],
        indexes: dict[str, Any],
    ) -> None:
        current_owner = current_id = ""
        for row in rows:
            person_raw = raw_text(_row_value(row, config.columns["person"]))
            role_raw = raw_text(_row_value(row, config.columns["role"]))
            role = normalize_role(role_raw)
            sheet_raw = raw_text(_row_value(row, config.columns["sheet"]))
            parcel_raw = raw_text(_row_value(row, config.columns["parcel"]))
            sheet = normalize_identifier(sheet_raw, "sheet")
            parcel = normalize_identifier(parcel_raw, "parcel")
            household_id = normalize_identifier(_row_value(row, config.columns.get("household_id", "")))
            if household_id:
                current_id = household_id
            if role == "HEAD" and person_raw:
                current_owner = person_raw
                if not household_id:
                    current_id = ""
            owner = normalize_name(current_owner)
            candidates = indexes["parcel"].get((sheet, parcel), [])
            record = None
            rule = ""
            if current_id:
                record = indexes["id"].get((current_id, sheet, parcel))
                rule = "HOUSEHOLD_ID+SHEET+PARCEL" if record else ""
            if record is None and owner:
                record = indexes["owner"].get((owner, sheet, parcel))
                rule = "OWNER+SHEET+PARCEL" if record else ""
            if record is None and len(candidates) == 1:
                record = candidates[0]
                if owner != record.owner_normalized:
                    _append_issue(record, "DATA_MISMATCH")
                    _append_issue(record, "REVIEW_REQUIRED")
                    record.suggested_match = record.owner_raw
                    record.confidence = 0.5
                    rule = "UNIQUE_SOURCE_SHEET+PARCEL_REVIEW"
                else:
                    rule = "UNIQUE_SOURCE_SHEET+PARCEL+HOUSEHOLD_RELATION"

            upload = UploadRow(
                row_number=int(row["_row"]), person_raw=person_raw,
                person_normalized=normalize_name(person_raw), role=role,
                sheet_raw=sheet_raw, sheet_normalized=sheet,
                parcel_raw=parcel_raw, parcel_normalized=parcel,
                household_id=current_id,
                tbxn_reference=raw_text(_row_value(row, config.columns.get("tbxn", ""))),
                ddk_reference=raw_text(_row_value(row, config.columns.get("ddk", ""))),
                combined_reference=raw_text(_row_value(row, config.columns.get("combined", ""))),
                raw=row,
            )
            if record is None:
                record = ParcelRecord(current_owner or person_raw, owner, sheet_raw, sheet, parcel_raw, parcel, current_id)
                if len(candidates) > 1:
                    _append_issue(record, "AMBIGUOUS_MATCH")
                    _append_issue(record, "REVIEW_REQUIRED")
                    record.suggested_match = " | ".join(f"{item.owner_raw}: Tờ {item.sheet_raw}, Thửa {item.parcel_raw}" for item in candidates[:3])
                    record.confidence = 0.0
                elif len(candidates) == 1:
                    _append_issue(record, "DATA_MISMATCH")
                    record.suggested_match = candidates[0].owner_raw
                    record.confidence = 0.5
                elif owner:
                    same_household = [item for item in records if item.owner_normalized == owner and item.source_rows]
                    if same_household:
                        _append_issue(record, "PARCEL_MISMATCH")
                        record.suggested_match = " | ".join(
                            f"Tờ {item.sheet_raw}, Thửa {item.parcel_raw}" for item in same_household[:3]
                        )
                        record.confidence = 0.0
                records.append(record)
            if not person_raw:
                _append_issue(record, "MISSING_NAME")
                _append_issue(record, "REVIEW_REQUIRED")
            if not role_raw:
                _append_issue(record, "MISSING_ROLE")
                _append_issue(record, "REVIEW_REQUIRED")
            elif role == "UNKNOWN":
                _append_issue(record, "UNKNOWN_ROLE")
                _append_issue(record, "REVIEW_REQUIRED")
            if not sheet:
                _append_issue(record, "MISSING_SHEET")
                _append_issue(record, "REVIEW_REQUIRED")
            if not parcel:
                _append_issue(record, "MISSING_PARCEL")
                _append_issue(record, "REVIEW_REQUIRED")
            record.upload_rows.append(upload)
            record.matching_rule = record.matching_rule or rule

    @staticmethod
    def _attach_documents(records: list[ParcelRecord], documents: list[DocumentRecord]) -> None:
        index: dict[tuple[str, str, DocumentType], list[DocumentRecord]] = defaultdict(list)
        records_by_parcel: dict[tuple[str, str], list[ParcelRecord]] = defaultdict(list)
        for record in records:
            if record.sheet_normalized and record.parcel_normalized and (record.source_rows or record.upload_rows):
                records_by_parcel[(record.sheet_normalized, record.parcel_normalized)].append(record)
        for document in documents:
            if document.sheet_normalized and document.parcel_normalized and document.document_type != DocumentType.UNKNOWN:
                index[document.index_key].append(document)
        for record in records:
            if not record.sheet_normalized or not record.parcel_normalized:
                continue
            for kind, attribute in ((DocumentType.TBXN, "tbxn_documents"), (DocumentType.DDK, "ddk_documents")):
                matches = index.get((record.sheet_normalized, record.parcel_normalized, kind), [])
                owner_matches = [doc for doc in matches if not doc.owner_normalized or doc.owner_normalized == record.owner_normalized]
                selected = owner_matches or ([] if any(doc.owner_normalized for doc in matches) else matches)
                setattr(record, attribute, selected)
                if (
                    selected
                    and len(records_by_parcel[(record.sheet_normalized, record.parcel_normalized)]) > 1
                    and all(not document.owner_normalized for document in selected)
                ):
                    _append_issue(record, "AMBIGUOUS_MATCH")
                    _append_issue(record, "REVIEW_REQUIRED")
                    record.suggested_match = "Tờ/thửa xuất hiện ở nhiều hộ; filename không chứa chủ hộ."
                    record.confidence = 0.0
                if not selected:
                    candidates = best_document_candidates(documents, record.sheet_normalized, record.parcel_normalized, kind)
                    if candidates:
                        # A nearby filename is only a suggestion. It must not
                        # replace the definite result that the exact file is absent.
                        record.suggested_match = " | ".join(str(item.path) for item in candidates)
        matched_to: dict[Path, list[ParcelRecord]] = defaultdict(list)
        for record in records:
            for document in record.tbxn_documents + record.ddk_documents:
                matched_to[document.path.resolve()].append(record)
        for document in documents:
            matches = matched_to.get(document.path.resolve(), [])
            document.matched_parcels = [
                f"{record.owner_raw} | Tờ {record.sheet_raw} | Thửa {record.parcel_raw}"
                for record in matches
            ]
            if len(matches) == 1:
                document.match_status = "MATCHED"
                document.match_method = "EXACT_SHEET+PARCEL+TYPE"
            elif len(matches) > 1:
                document.match_status = "AMBIGUOUS"
                document.match_method = "AMBIGUOUS_SHEET+PARCEL"
            else:
                document.match_status = "ORPHAN"
                document.match_method = "NONE"

    @staticmethod
    def _new_output_folder(root: Path) -> Path:
        root.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().astimezone().strftime("%Y-%m-%d_%H%M%S")
        candidate = root / f"VBDLIS_VALIDATION_{timestamp}"
        index = 1
        while candidate.exists():
            candidate = root / f"VBDLIS_VALIDATION_{timestamp}_{index:02d}"
            index += 1
        candidate.mkdir()
        return candidate

    def _export(self, config: RunConfig, result: ValidationResult) -> None:
        output = self._new_output_folder(config.output_root)
        log_file = output / "validation.log"
        logger = logging.getLogger(f"VBDLISValidation.{id(result)}")
        logger.setLevel(logging.INFO)
        handler = logging.FileHandler(log_file, encoding="utf-8")
        logger.addHandler(handler)
        try:
            logger.info(
                "Bắt đầu kiểm tra | mode=%s | upload=%s | source=%s",
                config.validation_mode.value, config.upload.path,
                config.source.path if config.source else "NOT_AVAILABLE",
            )
            logger.info(
                "Sheets | upload=%s | source=%s", config.upload.sheet,
                config.source.sheet if config.source else "NOT_AVAILABLE",
            )
            logger.info(
                "Mappings | upload=%s | source=%s", config.upload.columns,
                config.source.columns if config.source else "NOT_AVAILABLE",
            )
            ready_records = [r for r in result.records if r.workflow_status == WorkflowStatus.READY_TO_UPLOAD]
            need_signature_records = [r for r in result.records if r.workflow_status == WorkflowStatus.NEED_SIGNATURE]
            invalid_records = [r for r in result.records if r.workflow_status not in {WorkflowStatus.READY_TO_UPLOAD, WorkflowStatus.NEED_SIGNATURE}]
            ready_rows = sorted(row.row_number for record in ready_records for row in record.upload_rows)
            need_signature_rows = sorted(row.row_number for record in need_signature_records for row in record.upload_rows)
            invalid_rows = sorted(row.row_number for record in invalid_records for row in record.upload_rows)
            document_complete_rows = sorted(
                row.row_number
                for record in result.records
                if record.document_completeness == DocumentCompleteness.COMPLETE
                for row in record.upload_rows
            )
            document_complete_records = [
                record
                for record in result.records
                if record.document_completeness == DocumentCompleteness.COMPLETE
            ]
            row_records = {row.row_number: record for record in result.records for row in record.upload_rows}
            ready_excel = output / "VBDLIS_READY_TO_UPLOAD.xlsx"
            need_signature_excel = output / "VBDLIS_NEED_SIGNATURE.xlsx"
            invalid_excel = output / "VBDLIS_MISSING_OR_INVALID.xlsx"
            document_complete_excel = output / "VBDLIS_DOCUMENT_COMPLETE.xlsx"
            self._write_row_export(config, ready_rows, row_records, ready_records, ready_excel, diagnostics=False)
            self._write_row_export(
                config, need_signature_rows, row_records, need_signature_records,
                need_signature_excel, diagnostics=True,
            )
            self._write_row_export(
                config, invalid_rows, row_records, invalid_records,
                invalid_excel, diagnostics=True,
            )
            self._write_row_export(
                config, document_complete_rows, row_records, document_complete_records,
                document_complete_excel, diagnostics=False,
            )
            summary_excel = output / "Validation_Summary.xlsx"
            self._write_summary(result, summary_excel)
            document_actions = output / "DOCUMENT_ACTIONS"
            self._copy_action_documents(result.records, result.documents, document_actions, logger)
            result.output_folder = output
            result.valid_excel = ready_excel if config.separate_unsigned else document_complete_excel
            result.ready_excel = ready_excel
            result.need_signature_excel = need_signature_excel
            result.invalid_excel = invalid_excel
            result.document_complete_excel = document_complete_excel
            result.summary_excel = summary_excel
            result.document_actions = document_actions
            result.log_file = log_file
            classified_rows = len(ready_rows) + len(need_signature_rows) + len(invalid_rows)
            if classified_rows != result.total_upload_rows:
                raise RuntimeError(
                    f"Kiểm tra bảo toàn dòng thất bại: {classified_rows}/{result.total_upload_rows} dòng được phân loại."
                )
            accounted_documents = sum(
                document.match_status in {"MATCHED", "AMBIGUOUS", "ORPHAN"}
                for document in result.documents
            )
            if accounted_documents != len(result.documents):
                raise RuntimeError(
                    "Kiểm tra bảo toàn tài liệu thất bại: "
                    f"{accounted_documents}/{len(result.documents)} file được phân loại."
                )
            logger.info("Hoàn thành | stats=%s | output=%s", json.dumps(result.stats, ensure_ascii=False), output)
        finally:
            handler.close()
            logger.removeHandler(handler)

    @staticmethod
    def _copy_cell(source, target) -> None:
        target.value = source.value
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)

    @staticmethod
    def _safe_excel(value: Any) -> Any:
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    def _write_row_export(
        self,
        config: RunConfig,
        row_numbers: list[int],
        row_records: dict[int, ParcelRecord],
        detail_records: list[ParcelRecord],
        path: Path,
        diagnostics: bool,
    ) -> None:
        source_book = load_workbook(config.upload.path, data_only=False)
        target_book = Workbook()
        target_ws = target_book.active
        target_ws.title = config.upload.sheet[:31]
        source_ws = source_book[config.upload.sheet]
        start = column_index_from_string(config.export_start.upper())
        end = column_index_from_string(config.export_end.upper())
        diagnostic_headers = [
            "Validation_Mode", "Document_Completeness", "Signature_Readiness", "Workflow_Status",
            "Validation_Issues", "Recommended_Action", "Technical_Issue_Codes",
            "Household_Key", "Household_Name", "Map_Sheet", "Parcel_Number", "Member_Count",
            "TBXN_Status", "TBXN_File", "TBXN_Signature_Status", "TBXN_Trust_Verified",
            "TBXN_Signature_Count", "TBXN_Certificate_Info",
            "DDK_Status", "DDK_File", "DDK_Signature_Status", "DDK_Trust_Verified",
            "DDK_Signature_Count", "DDK_Certificate_Info",
            "AX_Raw", "AX_vs_Physical_Status", "Source_Completeness_Status",
            "Suggested_Match", "Match_Confidence", "Matching_Rule", "Match_Evidence",
            "Original_Row_Number",
        ]
        try:
            for source_column in range(start, end + 1):
                target_column = source_column - start + 1
                letter = get_column_letter(source_column)
                target_ws.column_dimensions[get_column_letter(target_column)].width = source_ws.column_dimensions[letter].width
            target_row = 1
            header_end = max(
                config.upload.header_row,
                config.upload.header_row_2 or config.upload.header_row,
            )
            for source_row in range(1, header_end + 1):
                for source_column in range(start, end + 1):
                    self._copy_cell(source_ws.cell(source_row, source_column), target_ws.cell(target_row, source_column - start + 1))
                target_ws.row_dimensions[target_row].height = source_ws.row_dimensions[source_row].height
                target_row += 1
            for merged in source_ws.merged_cells.ranges:
                if (
                    merged.min_row >= 1
                    and merged.max_row <= header_end
                    and merged.min_col >= start
                    and merged.max_col <= end
                ):
                    target_ws.merge_cells(
                        start_row=merged.min_row,
                        end_row=merged.max_row,
                        start_column=merged.min_col - start + 1,
                        end_column=merged.max_col - start + 1,
                    )
            if diagnostics:
                for offset, value in enumerate(diagnostic_headers, end - start + 2):
                    cell = target_ws.cell(header_end, offset, value)
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                diagnostic_start = end - start + 2
                for offset, width in enumerate((16, 72, 72, 42, 28, 14, 16), diagnostic_start):
                    target_ws.column_dimensions[get_column_letter(offset)].width = width
            for source_row in row_numbers:
                for source_column in range(start, end + 1):
                    self._copy_cell(source_ws.cell(source_row, source_column), target_ws.cell(target_row, source_column - start + 1))
                target_ws.row_dimensions[target_row].height = source_ws.row_dimensions[source_row].height
                if diagnostics:
                    record = row_records[source_row]
                    tbxn = record.tbxn_documents[0] if len(record.tbxn_documents) == 1 else None
                    ddk = record.ddk_documents[0] if len(record.ddk_documents) == 1 else None
                    values = [
                        config.validation_mode.value, record.document_completeness.value,
                        record.signature_readiness.value, record.workflow_status.value,
                        describe_issues(record.issues), recommended_actions(record.issues),
                        " | ".join(record.issues), record.household_id or record.owner_normalized,
                        record.owner_raw, record.sheet_raw, record.parcel_raw, len(record.upload_rows),
                        "FOUND" if record.tbxn_documents else "MISSING", str(tbxn.path) if tbxn else "",
                        tbxn.signature.status.value if tbxn and tbxn.signature else "", tbxn.signature.trust_verified if tbxn and tbxn.signature else "",
                        tbxn.signature.signature_count if tbxn and tbxn.signature else "",
                        json.dumps(tbxn.signature.certificates, ensure_ascii=False) if tbxn and tbxn.signature else "",
                        "FOUND" if record.ddk_documents else "MISSING", str(ddk.path) if ddk else "",
                        ddk.signature.status.value if ddk and ddk.signature else "", ddk.signature.trust_verified if ddk and ddk.signature else "",
                        ddk.signature.signature_count if ddk and ddk.signature else "",
                        json.dumps(ddk.signature.certificates, ensure_ascii=False) if ddk and ddk.signature else "",
                        record.ax_raw, record.ax_status,
                        "NOT_RUN_NOT_AVAILABLE" if config.validation_mode == ValidationMode.VBDLIS_ONLY else ("MATCHED" if record.source_rows else "NOT_FOUND"),
                        record.suggested_match, record.confidence, record.matching_rule,
                        record.match_evidence, source_row,
                    ]
                    for offset, value in enumerate(values, end - start + 2):
                        cell = target_ws.cell(target_row, offset, self._safe_excel(value))
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                target_row += 1
            if diagnostics:
                details = target_book.create_sheet("PARCEL_DETAILS")
                details.append([
                    "Household", "Map_Sheet", "Parcel_Number", "Source_Rows", "VBDLIS_Rows",
                    "Lý do dễ hiểu", "Cách xử lý đề nghị", "Technical_Issue_Codes", "TBXN", "DDK",
                    "Suggested_Match", "Confidence", "Matching_Rule",
                ])
                seen: set[int] = set()
                # Include source-only parcels as well; they cannot appear in the row-level sheet.
                for record in detail_records:
                    if id(record) in seen:
                        continue
                    seen.add(id(record))
                    details.append([
                        self._safe_excel(record.owner_raw), record.sheet_raw, record.parcel_raw,
                        ", ".join(map(str, record.source_rows)), ", ".join(str(row.row_number) for row in record.upload_rows),
                        describe_issues(record.issues), recommended_actions(record.issues),
                        " | ".join(record.issues), " | ".join(str(doc.path) for doc in record.tbxn_documents),
                        " | ".join(str(doc.path) for doc in record.ddk_documents), self._safe_excel(record.suggested_match),
                        record.confidence, record.matching_rule,
                    ])
            self._save_verified(
                target_book,
                path,
                target_ws.title,
                expected_rows=header_end + len(row_numbers),
            )
        finally:
            source_book.close()
            target_book.close()

    def _write_summary(self, result: ValidationResult, path: Path) -> None:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "SUMMARY"
        summary.append(["Chỉ số", "Giá trị"])
        labels = {
            "total_vbdlis_rows": "Total Excel Rows", "total_households": "Total Households",
            "total_parcels": "Total Parcels", "tbxn_pdfs_found": "TBXN PDFs Found",
            "ddk_pdfs_found": "DDK PDFs Found", "document_complete_parcels": "Document Complete Parcels",
            "ready_to_upload": "Ready To Upload", "need_signature": "Need Signature",
            "missing_document": "Missing Document", "signature_invalid": "Signature Invalid",
            "review_required": "Review Required",
            "missing_tbxn": "Missing TBXN", "missing_ddk": "Missing DDK", "missing_both": "Missing both",
            "tbxn_unsigned": "Unsigned TBXN", "ddk_unsigned": "Unsigned DDK",
            "tbxn_signature_invalid": "Invalid TBXN signature", "ddk_signature_invalid": "Invalid DDK signature",
            "ambiguous_match": "Ambiguous matches", "data_mismatch": "Source/data mismatch",
            "duplicate_tbxn": "Duplicate TBXN", "duplicate_ddk": "Duplicate DDK",
            "both_unsigned": "Both documents unsigned", "source_only": "Source-only parcels",
            "vbdlis_only": "VBDLIS-only parcels", "parcel_mismatch": "Parcel mismatches",
            "household_not_found": "Households not found",
            "ax_match": "AX Match", "ax_mismatch": "AX Mismatch", "ax_blank": "AX Blank",
            "ax_parse_error": "AX Parse Error", "orphan_tbxn": "Orphan TBXN",
            "orphan_ddk": "Orphan DDK", "invalid_pdfs": "Invalid PDFs",
            "valid_percentage": "READY percentage", "error_percentage": "NOT READY percentage",
        }
        stats = result.stats | result.scan_stats
        summary.append(["Validation Mode", result.validation_mode.value])
        summary.append([
            "Source completeness validation",
            "NOT RUN / NOT AVAILABLE" if result.validation_mode == ValidationMode.VBDLIS_ONLY else "COMPLETED",
        ])
        for key, label in labels.items():
            summary.append([label, stats.get(key, 0)])
        details = workbook.create_sheet("PARCEL_RESULTS")
        details.append([
            "Household", "Sheet", "Parcel", "MemberCount", "OriginalRows",
            "TBXN_Status", "TBXN_Path", "TBXN_Signature", "DDK_Status", "DDK_Path", "DDK_Signature",
            "Document_Completeness", "Signature_Readiness", "Workflow_Status", "AX_Raw", "AX_Status",
            "Match_Method", "Match_Confidence", "Issues", "Lý do dễ hiểu", "Cách xử lý đề nghị",
        ])
        for record in result.records:
            tbxn = record.tbxn_documents[0] if len(record.tbxn_documents) == 1 else None
            ddk = record.ddk_documents[0] if len(record.ddk_documents) == 1 else None
            details.append([
                self._safe_excel(record.owner_raw), record.sheet_raw, record.parcel_raw, len(record.upload_rows),
                ", ".join(str(row.row_number) for row in record.upload_rows),
                "FOUND" if record.tbxn_documents else "MISSING", str(tbxn.path) if tbxn else "",
                tbxn.signature.status.value if tbxn and tbxn.signature else "",
                "FOUND" if record.ddk_documents else "MISSING", str(ddk.path) if ddk else "",
                ddk.signature.status.value if ddk and ddk.signature else "",
                record.document_completeness.value, record.signature_readiness.value,
                record.workflow_status.value, record.ax_raw, record.ax_status,
                record.matching_rule, record.confidence, " | ".join(record.issues),
                describe_issues(record.issues), recommended_actions(record.issues),
            ])

        households = workbook.create_sheet("HOUSEHOLD_SUMMARY")
        households.append([
            "Household", "MemberCount", "ParcelCount", "ReadyParcelCount", "NeedSignatureParcelCount",
            "MissingParcelCount", "ReviewParcelCount", "MissingTBXN", "MissingDDK", "UnsignedDocuments",
        ])
        household_records: dict[str, list[ParcelRecord]] = defaultdict(list)
        for record in result.records:
            household_records[record.household_id or record.owner_normalized].append(record)
        for key, records in sorted(household_records.items()):
            people = {row.person_normalized for record in records for row in record.upload_rows if row.person_normalized}
            households.append([
                records[0].owner_raw or key, len(people), len(records),
                sum(r.workflow_status == WorkflowStatus.READY_TO_UPLOAD for r in records),
                sum(r.workflow_status == WorkflowStatus.NEED_SIGNATURE for r in records),
                sum(r.workflow_status == WorkflowStatus.MISSING_DOCUMENT for r in records),
                sum(r.workflow_status == WorkflowStatus.REVIEW_REQUIRED for r in records),
                sum("MISSING_TBXN" in r.issues for r in records), sum("MISSING_DDK" in r.issues for r in records),
                sum("UNSIGNED" in issue for r in records for issue in r.issues if issue != "BOTH_UNSIGNED"),
            ])

        inventory = workbook.create_sheet("DOCUMENT_INVENTORY")
        inventory.append([
            "Type", "Filename", "RelativePath", "FullPath", "ParentFolder", "OwnerParsed", "SheetParsed",
            "ParcelParsed", "ParseStatus", "SignatureStatus", "SignatureCount", "CryptographicIntegrity",
            "CertificateTrust", "ModifiedAfterSigning", "MatchedParcel", "MatchMethod", "MatchStatus",
        ])
        for document in result.documents:
            signature = document.signature
            inventory.append([
                document.document_type.value, document.path.name, document.relative_path, str(document.path),
                document.parent_folder, document.owner_raw, document.sheet_raw, document.parcel_raw,
                document.parse_status, signature.status.value if signature else "NOT_CHECKED",
                signature.signature_count if signature else 0,
                signature.cryptographic_integrity if signature else False,
                signature.trust_verified if signature else False,
                signature.document_modified_after_signing if signature else None,
                " | ".join(document.matched_parcels), document.match_method, document.match_status,
            ])

        actions = workbook.create_sheet("SIGNATURE_ACTIONS")
        actions.append(["Household", "Sheet", "Parcel", "Document_Type", "Filename", "Full_Path", "Signature_Status", "Action"])
        missing = workbook.create_sheet("MISSING_DOCUMENTS")
        missing.append(["Household", "Sheet", "Parcel", "Document_Completeness", "Issues", "Action"])
        for record in result.records:
            if record.workflow_status == WorkflowStatus.NEED_SIGNATURE:
                for document in record.tbxn_documents + record.ddk_documents:
                    if document.signature and document.signature.status == SignatureStatus.UNSIGNED:
                        actions.append([
                            record.owner_raw, record.sheet_raw, record.parcel_raw, document.document_type.value,
                            document.path.name, str(document.path), document.signature.status.value,
                            f"KÝ {document.document_type.value}",
                        ])
            elif record.workflow_status == WorkflowStatus.SIGNATURE_INVALID:
                for document in record.tbxn_documents + record.ddk_documents:
                    if document.signature and document.signature.status != SignatureStatus.SIGNED_VALID:
                        actions.append([
                            record.owner_raw, record.sheet_raw, record.parcel_raw, document.document_type.value,
                            document.path.name, str(document.path), document.signature.status.value,
                            "KIỂM TRA LẠI CHỮ KÝ",
                        ])
            if record.workflow_status == WorkflowStatus.MISSING_DOCUMENT:
                missing.append([
                    record.owner_raw, record.sheet_raw, record.parcel_raw,
                    record.document_completeness.value, " | ".join(record.issues),
                    recommended_actions(record.issues),
                ])
        self._save_verified(workbook, path, "SUMMARY")
        workbook.close()

    @staticmethod
    def _save_verified(
        workbook: Workbook,
        path: Path,
        required_sheet: str,
        expected_rows: int | None = None,
    ) -> None:
        temporary = path.with_name(f".{path.stem}.tmp.xlsx")
        workbook.save(temporary)
        try:
            verified = load_workbook(temporary, read_only=True, data_only=False)
            try:
                if required_sheet not in verified.sheetnames:
                    raise RuntimeError(f"File kết quả thiếu sheet {required_sheet}.")
                if expected_rows is not None and verified[required_sheet].max_row != expected_rows:
                    raise RuntimeError(
                        f"File kết quả có {verified[required_sheet].max_row} dòng, "
                        f"không khớp {expected_rows} dòng dự kiến."
                    )
            finally:
                verified.close()
            temporary.replace(path)
        finally:
            if temporary.exists():
                temporary.unlink()

    @staticmethod
    def _copy_action_documents(
        records: Iterable[ParcelRecord],
        documents: Iterable[DocumentRecord],
        root: Path,
        logger: logging.Logger | None = None,
    ) -> None:
        categories = (
            "NEED_SIGNATURE/TBXN_UNSIGNED", "NEED_SIGNATURE/DDK_UNSIGNED", "NEED_SIGNATURE/BOTH_UNSIGNED",
            "MISSING_DOCUMENTS/MISSING_TBXN", "MISSING_DOCUMENTS/MISSING_DDK",
            "MISSING_DOCUMENTS/MISSING_BOTH", "INVALID_SIGNATURE", "DUPLICATE_DOCUMENT",
            "AMBIGUOUS", "ORPHAN_DOCUMENTS", "OTHER",
        )
        for category in categories: (root / category).mkdir(parents=True, exist_ok=True)
        copied: set[tuple[Path, Path]] = set()
        for record in records:
            if record.workflow_status == WorkflowStatus.READY_TO_UPLOAD: continue
            if record.workflow_status == WorkflowStatus.NEED_SIGNATURE:
                if record.signature_readiness == SignatureReadiness.BOTH_UNSIGNED:
                    category = "NEED_SIGNATURE/BOTH_UNSIGNED"
                elif record.signature_readiness == SignatureReadiness.TBXN_UNSIGNED:
                    category = "NEED_SIGNATURE/TBXN_UNSIGNED"
                else: category = "NEED_SIGNATURE/DDK_UNSIGNED"
                selected_documents = [
                    d for d in record.tbxn_documents + record.ddk_documents
                    if d.signature and d.signature.status == SignatureStatus.UNSIGNED
                ]
            elif "MISSING_DDK" in record.issues:
                category, selected_documents = "MISSING_DOCUMENTS/MISSING_DDK", record.tbxn_documents
            elif "MISSING_TBXN" in record.issues:
                category, selected_documents = "MISSING_DOCUMENTS/MISSING_TBXN", record.ddk_documents
            elif "MISSING_BOTH" in record.issues:
                category, selected_documents = "MISSING_DOCUMENTS/MISSING_BOTH", []
            elif any("UNSIGNED" in issue or "SIGNATURE_INVALID" in issue or issue == "INVALID_PDF" for issue in record.issues):
                category, selected_documents = "INVALID_SIGNATURE", record.tbxn_documents + record.ddk_documents
            elif any(issue in record.issues for issue in ("DUPLICATE_TBXN", "DUPLICATE_DDK")):
                category, selected_documents = "DUPLICATE_DOCUMENT", record.tbxn_documents + record.ddk_documents
            elif any(issue in record.issues for issue in ("AMBIGUOUS_MATCH", "DUPLICATE_TBXN", "DUPLICATE_DDK", "REVIEW_REQUIRED")):
                category, selected_documents = "AMBIGUOUS", record.tbxn_documents + record.ddk_documents
            else:
                category, selected_documents = "OTHER", record.tbxn_documents + record.ddk_documents
            parcel_folder = sanitize_windows_name(
                f"{record.owner_raw}__TO_{record.sheet_raw}__THUA_{record.parcel_raw}"
            )
            destination = root / category / parcel_folder
            for document in selected_documents:
                try:
                    destination.mkdir(parents=True, exist_ok=True)
                    target = destination / sanitize_windows_name(document.path.name, "document.pdf")
                    safe_stem = sanitize_windows_name(document.path.stem, "document")
                    counter = 1
                    while target.exists() and target.read_bytes() != document.path.read_bytes():
                        target = destination / f"{safe_stem}__{counter:03d}{document.path.suffix}"
                        counter += 1
                    key = (document.path.resolve(), target)
                    if key not in copied and not target.exists():
                        shutil.copy2(document.path, target)
                        copied.add(key)
                except OSError as error:
                    if logger is not None:
                        logger.warning("Không thể copy tài liệu lỗi %s: %s", document.path, error)
        orphan_root = root / "ORPHAN_DOCUMENTS"
        for document in documents:
            if document.match_status != "ORPHAN": continue
            destination = orphan_root / document.document_type.value
            destination.mkdir(parents=True, exist_ok=True)
            target = destination / sanitize_windows_name(document.path.name, "document.pdf")
            safe_stem = sanitize_windows_name(document.path.stem, "document")
            counter = 1
            try:
                while target.exists() and target.read_bytes() != document.path.read_bytes():
                    target = destination / f"{safe_stem}__{counter:03d}{document.path.suffix}"
                    counter += 1
                if not target.exists():
                    shutil.copy2(document.path, target)
            except OSError as error:
                if logger is not None: logger.warning("Không thể copy tài liệu orphan %s: %s", document.path, error)
