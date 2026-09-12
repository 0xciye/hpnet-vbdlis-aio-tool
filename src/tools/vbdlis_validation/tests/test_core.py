from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from cryptography import x509
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import rsa
from cryptography.hazmat.primitives.serialization import pkcs12
from cryptography.x509.oid import NameOID
from openpyxl import Workbook, load_workbook
from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
from pyhanko.sign import signers
from pypdf import PdfWriter

from tools.vbdlis_validation.documents import (
    DocumentScanner,
    SignatureValidator,
    parse_document,
)
from tools.vbdlis_validation.messages import (
    describe_issues,
    friendly_failure,
    recommended_actions,
)
from tools.vbdlis_validation.models import (
    DocumentCompleteness,
    DocumentRecord,
    DocumentType,
    ExcelConfig,
    ParcelRecord,
    RunConfig,
    SignatureReadiness,
    SignatureResult,
    SignatureStatus,
    UploadRow,
    ValidationMode,
    WorkflowStatus,
)
from tools.vbdlis_validation.normalization import (
    normalize_identifier,
    normalize_name,
    normalize_role,
)
from tools.vbdlis_validation.service import (
    ValidationService,
    _read_rows,
    evaluate_record,
    headers,
)


def _signature(status: SignatureStatus) -> SignatureResult:
    return SignatureResult(
        status, signature_count=0 if status == SignatureStatus.UNSIGNED else 1,
        cryptographic_integrity=status == SignatureStatus.SIGNED_VALID,
    )


def _document(kind: DocumentType, parcel: str, status: SignatureStatus) -> DocumentRecord:
    name = f"CHUACOGIAY_10930_12_{parcel}-{kind.value}.signed.pdf"
    return DocumentRecord(
        Path(name), kind, sheet_raw="12", sheet_normalized="12", parcel_raw=parcel,
        parcel_normalized=parcel, parse_confidence=1.0, signature=_signature(status),
    )


def _row(number: int, parcel: str, person: str = "Nguyễn Văn A", role: str = "HEAD") -> UploadRow:
    return UploadRow(
        number, person, normalize_name(person), role, "12", "12", parcel, parcel,
        tbxn_reference=f"CHUACOGIAY_10930_12_{parcel}-TBXN.pdf",
        ddk_reference=f"CHUACOGIAY_10930_12_{parcel}-DDK.pdf",
    )


def _record(parcel: str = "100", members: int = 1) -> ParcelRecord:
    record = ParcelRecord("Nguyễn Văn A", normalize_name("Nguyễn Văn A"), "12", "12", parcel, parcel)
    record.source_rows = [2]
    record.upload_rows = [_row(2 + index, parcel, f"Người {index}", "HEAD" if index == 0 else "MEMBER") for index in range(members)]
    record.tbxn_documents = [_document(DocumentType.TBXN, parcel, SignatureStatus.SIGNED_VALID)]
    record.ddk_documents = [_document(DocumentType.DDK, parcel, SignatureStatus.SIGNED_VALID)]
    return record


def _create_signed_pdf(folder: Path, signature_count: int = 1) -> Path:
    unsigned = folder / "unsigned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_metadata({"/Title": "SyntheticToken"})
    with unsigned.open("wb") as stream:
        writer.write(stream)

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    subject = issuer = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "VBDLIS Test Signer")])
    now = datetime.now(timezone.utc)
    certificate = (
        x509.CertificateBuilder().subject_name(subject).issuer_name(issuer).public_key(key.public_key())
        .serial_number(x509.random_serial_number()).not_valid_before(now - timedelta(days=1))
        .not_valid_after(now + timedelta(days=30)).add_extension(x509.BasicConstraints(ca=True, path_length=None), critical=True)
        .sign(key, hashes.SHA256())
    )
    pfx = folder / "signer.p12"
    pfx.write_bytes(pkcs12.serialize_key_and_certificates(
        b"signer", key, certificate, None, serialization.BestAvailableEncryption(b"secret")
    ))
    signer = signers.SimpleSigner.load_pkcs12(str(pfx), passphrase=b"secret")
    assert signer is not None
    current = unsigned
    for index in range(1, signature_count + 1):
        signed = folder / ("signed.pdf" if index == 1 else f"signed-{index}.pdf")
        with current.open("rb") as source, signed.open("wb") as target:
            signers.sign_pdf(
                IncrementalPdfFileWriter(source),
                signature_meta=signers.PdfSignatureMetadata(field_name=f"Signature{index}"),
                signer=signer,
                output=target,
            )
        current = signed
    return current


def test_case_1_valid_parcel():
    record = evaluate_record(_record())
    assert record.valid


def test_case_2_three_members_two_parcels_six_rows_not_duplicates():
    records = [evaluate_record(_record("100", 3)), evaluate_record(_record("101", 3))]
    assert all(record.valid for record in records)
    assert sum(len(record.upload_rows) for record in records) == 6
    assert not any("DUPLICATE" in issue for record in records for issue in record.issues)


@pytest.mark.parametrize(
    ("missing", "issue"),
    [("ddk", "MISSING_DDK"), ("tbxn", "MISSING_TBXN"), ("both", "MISSING_BOTH")],
)
def test_cases_3_to_5_missing_documents(missing, issue):
    record = _record()
    if missing in {"ddk", "both"}: record.ddk_documents = []
    if missing in {"tbxn", "both"}: record.tbxn_documents = []
    assert issue in evaluate_record(record).issues


@pytest.mark.parametrize(
    ("kind", "issue"),
    [(DocumentType.TBXN, "TBXN_UNSIGNED"), (DocumentType.DDK, "DDK_UNSIGNED")],
)
def test_cases_6_and_7_unsigned(kind, issue):
    record = _record()
    target = record.tbxn_documents if kind == DocumentType.TBXN else record.ddk_documents
    target[0].signature = _signature(SignatureStatus.UNSIGNED)
    assert issue in evaluate_record(record).issues


def test_both_unsigned_keeps_individual_issues_and_composite_status():
    record = _record()
    record.tbxn_documents[0].signature = _signature(SignatureStatus.UNSIGNED)
    record.ddk_documents[0].signature = _signature(SignatureStatus.UNSIGNED)
    issues = evaluate_record(record).issues
    assert {"TBXN_UNSIGNED", "DDK_UNSIGNED", "BOTH_UNSIGNED"} <= set(issues)
    assert describe_issues(issues).casefold().count("cả file tbxn và file ddk") == 1
    assert "Thay cả hai file" in recommended_actions(issues)


def test_failure_messages_give_nontechnical_next_step():
    message = friendly_failure("[Errno 13] Permission denied")
    assert "đóng file Excel/PDF" in message
    assert "thử lại" in message


def test_case_8_cryptographically_invalid_signature(tmp_path):
    signed = _create_signed_pdf(tmp_path)
    original = signed.read_bytes()
    assert b"SyntheticToken" in original
    signed.write_bytes(original.replace(b"SyntheticToken", b"SynthetixToken", 1))
    result = SignatureValidator().validate(signed)
    assert result.status == SignatureStatus.SIGNED_INVALID
    assert not result.cryptographic_integrity


def test_multiple_signatures_are_all_checked_without_false_invalid(tmp_path):
    result = SignatureValidator().validate(_create_signed_pdf(tmp_path, signature_count=3))
    assert result.status == SignatureStatus.SIGNED_VALID
    assert result.signature_count == 3
    assert result.cryptographic_integrity


def test_cases_9_and_10_source_only_and_vbdlis_only():
    source_only = _record(); source_only.upload_rows = []
    vbdlis_only = _record(); vbdlis_only.source_rows = []
    assert "SOURCE_ONLY" in evaluate_record(source_only).issues
    assert "VBDLIS_ONLY" in evaluate_record(vbdlis_only).issues


def test_case_11_parcel_stats_once_but_all_five_rows_exportable():
    record = _record(members=5); record.ddk_documents = []
    evaluate_record(record)
    assert record.issues.count("MISSING_DDK") == 1
    assert len(record.upload_rows) == 5


def test_case_12_normalization_variants():
    assert normalize_name("  NGUYỄN__Văn-A. ") == normalize_name("nguyen van a")
    assert normalize_identifier("Tờ số 12", "sheet") == normalize_identifier(12.0, "sheet") == "12"
    assert normalize_identifier("CN", "parcel") == ""


def test_invalid_letter_parcel_is_reported_as_source_data_error():
    record = ParcelRecord("Hộ A", "ho a", "92", "92", "CN", "")
    record.source_rows.append(288)

    evaluate_record(record)

    assert record.workflow_status == WorkflowStatus.DATA_ERROR
    assert "INVALID_PARCEL_IDENTIFIER" in record.issues
    assert "MISSING_BOTH" not in record.issues
    assert "số nguyên dương" in describe_issues(record.issues).casefold()
    assert "vbdlis" not in describe_issues(record.issues).casefold()


def test_case_13_ambiguous_household_is_not_auto_matched(tmp_path):
    source_config = ExcelConfig(tmp_path / "source.xlsx", "Data", 1, {"owner": "A", "sheet": "B", "parcel": "C"})
    source_rows = [
        {"_row": 2, "A": "Nguyễn Văn A", "B": 12, "C": 100},
        {"_row": 3, "A": "Nguyễn Văn B", "B": 12, "C": 100},
    ]
    records, indexes = ValidationService._build_source_records(source_config, source_rows)
    upload_config = ExcelConfig(tmp_path / "upload.xlsx", "Data", 1, {"person": "A", "role": "B", "sheet": "C", "parcel": "D", "tbxn": "E", "ddk": "F"})
    ValidationService._attach_upload_rows(upload_config, [{"_row": 2, "A": "Thành viên", "B": "Thành viên", "C": 12, "D": 100, "E": "x", "F": "y"}], records, indexes)
    ambiguous = next(record for record in records if record.upload_rows)
    assert {"AMBIGUOUS_MATCH", "REVIEW_REQUIRED"} <= set(ambiguous.issues)


def test_source_subtotals_are_ignored_and_unique_owner_mismatch_stays_one_parcel(tmp_path):
    source_config = ExcelConfig(
        tmp_path / "source.xlsx", "Data", 1,
        {"owner": "B", "sheet": "G", "parcel": "H"},
    )
    source_rows = [
        {"_row": 2, "B": "Lê Thị Dịu", "G": 92, "H": 330},
        {"_row": 3, "B": "Tổng DT", "G": None, "H": None},
    ]
    records, indexes = ValidationService._build_source_records(source_config, source_rows)
    assert len(records) == 1

    upload_config = ExcelConfig(
        tmp_path / "upload.xlsx", "Data", 1,
        {"person": "B", "role": "C", "sheet": "G", "parcel": "H", "combined": "I"},
    )
    upload_rows = [
        {"_row": 2, "B": "Lê Thị Êm", "C": "Chủ hộ", "G": 92, "H": 330, "I": "TBXN; DDK"},
        {"_row": 3, "B": "Người nhà", "C": "Thành viên", "G": 92, "H": 330, "I": "TBXN; DDK"},
    ]
    ValidationService._attach_upload_rows(upload_config, upload_rows, records, indexes)
    assert len(records) == 1
    assert len(records[0].upload_rows) == 2
    assert records[0].matching_rule == "UNIQUE_SOURCE_SHEET+PARCEL_REVIEW"
    assert {"DATA_MISMATCH", "REVIEW_REQUIRED"} <= set(records[0].issues)


def test_case_14_duplicate_tbxn_requires_review():
    record = _record(); record.tbxn_documents.append(_document(DocumentType.TBXN, "100", SignatureStatus.SIGNED_VALID))
    assert {"DUPLICATE_TBXN", "REVIEW_REQUIRED"} <= set(evaluate_record(record).issues)


def test_same_sheet_parcel_in_two_households_needs_owner_evidence():
    first = _record(); second = _record(); second.owner_raw = "Nguyễn Văn B"; second.owner_normalized = normalize_name(second.owner_raw)
    documents = first.tbxn_documents + first.ddk_documents
    first.tbxn_documents = []; first.ddk_documents = []; second.tbxn_documents = []; second.ddk_documents = []
    ValidationService._attach_documents([first, second], documents)
    assert all({"AMBIGUOUS_MATCH", "REVIEW_REQUIRED"} <= set(record.issues) for record in (first, second))


def test_case_15_corrupt_pdf_does_not_crash(tmp_path):
    path = tmp_path / "corrupt.pdf"; path.write_bytes(b"not a pdf")
    result = SignatureValidator().validate(path)
    assert result.status == SignatureStatus.SIGNATURE_ERROR and result.invalid_pdf


def test_project_filename_convention_is_reused():
    record = parse_document(Path("CHUACOGIAY_10930_12_100-TBXN.ldsigned.pdf"))
    assert record.document_type == DocumentType.TBXN
    assert (record.sheet_normalized, record.parcel_normalized, record.parse_confidence) == ("12", "100", 1.0)


def _save_workbook(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook(); ws = workbook.active; ws.title = "Data"; ws.append(headers)
    for row in rows: ws.append(row)
    workbook.save(path); workbook.close()


def test_two_level_headers_are_merged_and_data_starts_after_second_level(tmp_path):
    path = tmp_path / "two-level.xlsx"
    workbook = Workbook(); ws = workbook.active; ws.title = "Data"
    ws.merge_cells("A1:B1"); ws["A1"] = "Hộ gia đình"
    ws.merge_cells("C1:D1"); ws["C1"] = "Thửa đất"
    ws.append(["Tên chủ hộ", "Mã hộ", "Số tờ", "Số thửa"])
    ws.append(["Nguyễn Văn A", "H001", 12, 100])
    workbook.save(path); workbook.close()
    config = ExcelConfig(path, "Data", 1, {}, 2)

    assert headers(config) == [
        ("A", "Hộ gia đình — Tên chủ hộ"),
        ("B", "Hộ gia đình — Mã hộ"),
        ("C", "Thửa đất — Số tờ"),
        ("D", "Thửa đất — Số thửa"),
    ]
    assert _read_rows(config) == [{"_row": 3, "A": "Nguyễn Văn A", "B": "H001", "C": 12, "D": 100}]

    output = tmp_path / "two-level-output.xlsx"
    upload_record = _record(); upload_record.upload_rows = [_row(3, "100")]
    run_config = RunConfig(config, config, tmp_path, tmp_path, tmp_path, "A", "D")
    ValidationService()._write_row_export(
        run_config,
        [3],
        {3: upload_record},
        [upload_record],
        output,
        diagnostics=False,
    )
    exported = load_workbook(output)
    try:
        exported_ws = exported["Data"]
        assert {str(cell_range) for cell_range in exported_ws.merged_cells.ranges} == {"A1:B1", "C1:D1"}
        assert exported_ws["A3"].value == "Nguyễn Văn A"
    finally:
        exported.close()


def test_end_to_end_exports_all_member_rows_and_preserves_inputs(tmp_path):
    source = tmp_path / "source.xlsx"; upload = tmp_path / "upload.xlsx"
    _save_workbook(source, ["Chủ hộ", "Số tờ", "Số thửa"], [["Nguyễn Văn A", 12, 100], ["Nguyễn Văn A", 12, 101]])
    upload_rows = []
    for parcel in (100, 101):
        tbxn = f"CHUACOGIAY_10930_12_{parcel}-TBXN.pdf"
        ddk = f"CHUACOGIAY_10930_12_{parcel}-DDK.pdf"
        upload_rows.extend([
            ["Nguyễn Văn A", "Chủ hộ", 12, parcel, tbxn, ddk],
            ["Nguyễn Thị B", "Thành viên hộ gia đình", 12, parcel, tbxn, ddk],
            ["Nguyễn Văn C", "Thành viên hộ gia đình", 12, parcel, tbxn, ddk],
        ])
    _save_workbook(upload, ["Tên", "Vai trò", "Số tờ", "Số thửa", "TBXN", "DDK"], upload_rows)
    source_before, upload_before = source.read_bytes(), upload.read_bytes()
    documents = tmp_path / "documents"; documents.mkdir()
    signed = _create_signed_pdf(tmp_path)
    for parcel in (100, 101):
        for kind in ("TBXN", "DDK"):
            (documents / f"CHUACOGIAY_10930_12_{parcel}-{kind}.signed.pdf").write_bytes(signed.read_bytes())
    config = RunConfig(
        ExcelConfig(upload, "Data", 1, {"person": "A", "role": "B", "sheet": "C", "parcel": "D", "tbxn": "E", "ddk": "F"}),
        ExcelConfig(source, "Data", 1, {"owner": "A", "sheet": "B", "parcel": "C"}),
        documents, documents, tmp_path / "outputs", "A", "F", True,
    )
    result = ValidationService().run(config)
    assert result.stats["total_valid_parcels"] == 2
    assert result.stats["total_vbdlis_rows"] == 6
    valid_book = load_workbook(result.valid_excel, data_only=False)
    try:
        assert valid_book["Data"].max_row == 7
    finally:
        valid_book.close()
    assert source.read_bytes() == source_before and upload.read_bytes() == upload_before


@pytest.mark.parametrize(
    ("tbxn", "ddk", "completeness", "readiness", "workflow", "issue"),
    [
        (SignatureStatus.SIGNED_VALID, SignatureStatus.SIGNED_VALID, DocumentCompleteness.COMPLETE, SignatureReadiness.SIGNED_READY, WorkflowStatus.READY_TO_UPLOAD, "VALID"),
        (SignatureStatus.UNSIGNED, SignatureStatus.SIGNED_VALID, DocumentCompleteness.COMPLETE, SignatureReadiness.TBXN_UNSIGNED, WorkflowStatus.NEED_SIGNATURE, "TBXN_UNSIGNED"),
        (SignatureStatus.SIGNED_VALID, SignatureStatus.UNSIGNED, DocumentCompleteness.COMPLETE, SignatureReadiness.DDK_UNSIGNED, WorkflowStatus.NEED_SIGNATURE, "DDK_UNSIGNED"),
        (SignatureStatus.UNSIGNED, SignatureStatus.UNSIGNED, DocumentCompleteness.COMPLETE, SignatureReadiness.BOTH_UNSIGNED, WorkflowStatus.NEED_SIGNATURE, "BOTH_UNSIGNED"),
        (None, SignatureStatus.SIGNED_VALID, DocumentCompleteness.MISSING_TBXN, SignatureReadiness.NOT_CHECKED, WorkflowStatus.MISSING_DOCUMENT, "MISSING_TBXN"),
        (SignatureStatus.SIGNED_VALID, None, DocumentCompleteness.MISSING_DDK, SignatureReadiness.NOT_CHECKED, WorkflowStatus.MISSING_DOCUMENT, "MISSING_DDK"),
        (None, None, DocumentCompleteness.MISSING_BOTH, SignatureReadiness.NOT_CHECKED, WorkflowStatus.MISSING_DOCUMENT, "MISSING_BOTH"),
        (None, SignatureStatus.UNSIGNED, DocumentCompleteness.MISSING_TBXN, SignatureReadiness.DDK_UNSIGNED, WorkflowStatus.MISSING_DOCUMENT, "DDK_UNSIGNED"),
        (SignatureStatus.UNSIGNED, None, DocumentCompleteness.MISSING_DDK, SignatureReadiness.TBXN_UNSIGNED, WorkflowStatus.MISSING_DOCUMENT, "TBXN_UNSIGNED"),
        (SignatureStatus.SIGNED_INVALID, SignatureStatus.SIGNED_VALID, DocumentCompleteness.COMPLETE, SignatureReadiness.TBXN_SIGNATURE_INVALID, WorkflowStatus.SIGNATURE_INVALID, "TBXN_SIGNATURE_INVALID"),
        (SignatureStatus.SIGNED_VALID, SignatureStatus.SIGNED_INVALID, DocumentCompleteness.COMPLETE, SignatureReadiness.DDK_SIGNATURE_INVALID, WorkflowStatus.SIGNATURE_INVALID, "DDK_SIGNATURE_INVALID"),
    ],
)
def test_required_classification_matrix(tbxn, ddk, completeness, readiness, workflow, issue):
    record = _record()
    record.tbxn_documents = [] if tbxn is None else [_document(DocumentType.TBXN, "100", tbxn)]
    record.ddk_documents = [] if ddk is None else [_document(DocumentType.DDK, "100", ddk)]
    evaluate_record(record)
    assert record.document_completeness == completeness
    assert record.signature_readiness == readiness
    assert record.workflow_status == workflow
    assert issue in record.issues


@pytest.mark.parametrize("kind", [DocumentType.TBXN, DocumentType.DDK])
def test_invalid_pdf_is_not_treated_as_missing(kind):
    record = _record()
    target = record.tbxn_documents if kind == DocumentType.TBXN else record.ddk_documents
    target[0].signature = SignatureResult(SignatureStatus.SIGNATURE_ERROR, invalid_pdf=True)
    evaluate_record(record)
    assert record.document_completeness == DocumentCompleteness.INVALID_DOCUMENT
    assert record.signature_readiness == SignatureReadiness.SIGNATURE_CHECK_ERROR
    assert record.workflow_status == WorkflowStatus.SIGNATURE_INVALID
    assert "INVALID_PDF" in record.issues


@pytest.mark.parametrize("kind", [DocumentType.TBXN, DocumentType.DDK])
def test_duplicate_document_requires_review_for_each_type(kind):
    record = _record()
    target = record.tbxn_documents if kind == DocumentType.TBXN else record.ddk_documents
    target.append(_document(kind, "100", SignatureStatus.SIGNED_VALID))
    evaluate_record(record)
    assert record.document_completeness == DocumentCompleteness.DUPLICATE_DOCUMENT
    assert record.workflow_status == WorkflowStatus.REVIEW_REQUIRED


@pytest.mark.parametrize("kind", [DocumentType.TBXN, DocumentType.DDK])
def test_ambiguous_document_beats_signature_state(kind):
    record = _record()
    record.issues.extend(["AMBIGUOUS_MATCH", "REVIEW_REQUIRED"])
    target = record.tbxn_documents if kind == DocumentType.TBXN else record.ddk_documents
    target[0].signature = _signature(SignatureStatus.UNSIGNED)
    evaluate_record(record)
    assert record.document_completeness == DocumentCompleteness.AMBIGUOUS_DOCUMENT
    assert record.workflow_status == WorkflowStatus.REVIEW_REQUIRED


@pytest.mark.parametrize("kind", [DocumentType.TBXN, DocumentType.DDK])
def test_orphan_document_is_accounted_for(kind):
    record = _record("100")
    orphan = _document(kind, "999", SignatureStatus.SIGNED_VALID)
    ValidationService._attach_documents([record], [orphan])
    assert orphan.match_status == "ORPHAN"
    assert not orphan.matched_parcels


def test_nearby_filename_is_only_a_suggestion_not_an_ambiguous_match():
    record = _record("100")
    record.tbxn_documents = []
    record.ddk_documents = []
    nearby = _document(DocumentType.TBXN, "999", SignatureStatus.SIGNED_VALID)
    ValidationService._attach_documents([record], [nearby])
    evaluate_record(record, ValidationMode.VBDLIS_ONLY)
    assert record.workflow_status == WorkflowStatus.MISSING_DOCUMENT
    assert "AMBIGUOUS_MATCH" not in record.issues
    assert record.suggested_match
    assert record.confidence == 1.0


def test_ax_matrix_empty_missing_mismatch_and_parse_error():
    empty = _record(); empty.upload_rows[0].tbxn_reference = ""; empty.upload_rows[0].ddk_reference = ""
    ValidationService._validate_ax(empty)
    assert empty.ax_status == "EMPTY" and "AX_EMPTY" in empty.issues

    missing = _record(); missing.ddk_documents = []
    ValidationService._validate_ax(missing)
    assert missing.ax_status == "MISMATCH" and "AX_REFERENCES_MISSING_DDK" in missing.issues

    mismatch = _record(); mismatch.upload_rows[0].tbxn_reference = "wrong-TBXN.pdf"
    ValidationService._validate_ax(mismatch)
    assert mismatch.ax_status == "MISMATCH" and "AX_MISSING_TBXN_REFERENCE" in mismatch.issues

    unparseable = _record(); unparseable.upload_rows[0].tbxn_reference = "không rõ"; unparseable.upload_rows[0].ddk_reference = ""
    ValidationService._validate_ax(unparseable)
    assert unparseable.ax_status == "PARSE_ERROR" and "AX_PARSE_ERROR" in unparseable.issues


def test_role_variants_and_blank_key_fields_require_review(tmp_path):
    assert normalize_role(" CHỦ   HỘ ") == "HEAD"
    assert normalize_role("TV hộ") == "MEMBER"
    assert normalize_role("một giá trị khác") == "UNKNOWN"
    config = ExcelConfig(tmp_path / "v.xlsx", "Data", 1, {
        "person": "A", "role": "B", "sheet": "C", "parcel": "D", "combined": "E",
    })
    records = ValidationService._build_vbdlis_only_records(config, [
        {"_row": 2, "A": "", "B": "", "C": "", "D": "", "E": ""},
        {"_row": 3, "A": "Nguyễn Văn A", "B": "khác", "C": 12, "D": 100, "E": "TBXN; DDK"},
    ])
    issues = {issue for record in records for issue in record.issues}
    assert {"MISSING_NAME", "MISSING_ROLE", "MISSING_SHEET", "MISSING_PARCEL", "UNKNOWN_ROLE", "REVIEW_REQUIRED"} <= issues
    assert len(records) == 2


def test_household_parcel_collision_and_member_grouping(tmp_path):
    config = ExcelConfig(tmp_path / "v.xlsx", "Data", 1, {
        "person": "A", "role": "B", "sheet": "C", "parcel": "D", "combined": "E",
    })
    rows = [
        {"_row": 2, "A": "Nguyễn Văn A", "B": "Chủ hộ", "C": 10, "D": 50, "E": "TBXN; DDK"},
        {"_row": 3, "A": "Thành viên A", "B": "TV hộ", "C": 10, "D": 50, "E": "TBXN; DDK"},
        {"_row": 4, "A": "Nguyễn Văn A", "B": "Chủ hộ", "C": 11, "D": 50, "E": "TBXN; DDK"},
        {"_row": 5, "A": "Nguyễn Văn B", "B": "Chủ hộ", "C": 10, "D": 50, "E": "TBXN; DDK"},
    ]
    records = ValidationService._build_vbdlis_only_records(config, rows)
    assert len(records) == 3
    assert sorted(len(record.upload_rows) for record in records) == [1, 1, 2]
    documents = [_document(DocumentType.TBXN, "50", SignatureStatus.SIGNED_VALID)]
    documents[0].sheet_raw = documents[0].sheet_normalized = "10"
    ValidationService._attach_documents(records, documents)
    same_parcel = [record for record in records if record.sheet_normalized == "10"]
    assert all("AMBIGUOUS_MATCH" in record.issues for record in same_parcel)


def test_multiple_heads_and_household_without_head_require_review(tmp_path):
    config = ExcelConfig(tmp_path / "v.xlsx", "Data", 1, {
        "person": "A", "role": "B", "household_id": "C", "sheet": "D",
        "parcel": "E", "combined": "F",
    })
    records = ValidationService._build_vbdlis_only_records(config, [
        {"_row": 2, "A": "Chủ hộ A", "B": "Chủ hộ", "C": "H01", "D": 12, "E": 100, "F": "TBXN; DDK"},
        {"_row": 3, "A": "Chủ hộ B", "B": "Chủ hộ", "C": "H01", "D": 12, "E": 101, "F": "TBXN; DDK"},
        {"_row": 4, "A": "Thành viên lẻ", "B": "Thành viên", "C": "", "D": 12, "E": 102, "F": "TBXN; DDK"},
    ])
    first_household = [record for record in records if record.household_id == "h01"]
    assert first_household
    assert all({"MULTIPLE_HOUSEHOLD_HEADS", "REVIEW_REQUIRED"} <= set(record.issues) for record in first_household)

    no_id_config = ExcelConfig(tmp_path / "v.xlsx", "Data", 1, {
        "person": "A", "role": "B", "sheet": "D", "parcel": "E", "combined": "F",
    })
    no_head = ValidationService._build_vbdlis_only_records(no_id_config, [
        {"_row": 4, "A": "Thành viên lẻ", "B": "Thành viên", "D": 12, "E": 102, "F": "TBXN; DDK"},
    ])[0]
    assert {"HOUSEHOLD_WITHOUT_HEAD", "REVIEW_REQUIRED"} <= set(no_head.issues)


def test_recursive_scanner_paths_extensions_and_non_pdf(tmp_path):
    tbxn = tmp_path / "thư mục có dấu" / "a" / "b"; tbxn.mkdir(parents=True)
    ddk = tmp_path / "folder with spaces"; ddk.mkdir()
    (tbxn / "Tài liệu giống_12_100-TBXN.PDF").write_bytes(b"pdf")
    duplicate = tbxn.parent / "Tài liệu giống_12_100-TBXN.PDF"; duplicate.write_bytes(b"other")
    (ddk / "Đơn đăng ký_12_100-DDK.Pdf").write_bytes(b"pdf")
    (ddk / "bo-qua.txt").write_text("not pdf", encoding="utf-8")
    scanner = DocumentScanner()
    documents = scanner.scan(tbxn.parents[1], ddk, False, validate_signatures=False)
    assert len(documents) == 3
    assert len({document.path.resolve() for document in documents}) == 3
    assert scanner.last_stats["tbxn_maximum_depth"] >= 2
    assert all(document.path.suffix.casefold() == ".pdf" for document in documents)


def test_file_access_error_is_a_safe_signature_failure():
    record = _record()
    record.tbxn_documents[0].signature = SignatureResult(
        SignatureStatus.SIGNATURE_ERROR,
        file_access_error=True,
    )
    evaluate_record(record)
    assert {"FILE_ACCESS_ERROR", "SIGNATURE_CHECK_ERROR"} <= set(record.issues)
    assert record.workflow_status == WorkflowStatus.SIGNATURE_INVALID


def test_formula_and_numeric_identifier_cells_are_preserved(tmp_path):
    path = tmp_path / "formula.xlsx"
    workbook = Workbook(); ws = workbook.active; ws.title = "Data"
    ws.append(["Số tờ", "Công thức"]); ws.append([12.0, "=1+1"])
    workbook.save(path); workbook.close()
    rows = _read_rows(ExcelConfig(path, "Data", 1, {}))
    assert normalize_identifier(rows[0]["A"], "sheet") == "12"
    assert rows[0]["B"] == "=1+1"


def test_signed_fixture_is_intact_even_when_trust_root_is_unknown(tmp_path):
    result = SignatureValidator().validate(_create_signed_pdf(tmp_path))
    assert result.status == SignatureStatus.SIGNED_VALID
    assert result.cryptographic_integrity
    assert not result.trust_verified


def test_unsigned_five_member_integration_default_and_separate_outputs(tmp_path):
    upload = tmp_path / "upload.xlsx"
    tbxn_name = "CHUACOGIAY_10930_12_100-TBXN.pdf"
    ddk_name = "CHUACOGIAY_10930_12_100-DDK.pdf"
    rows = [
        ["Nguyễn Văn A" if index == 0 else f"Thành viên {index}", "Chủ hộ" if index == 0 else "Thành viên hộ gia đình", 12, 100, f"{tbxn_name}, {ddk_name}"]
        for index in range(5)
    ]
    _save_workbook(upload, ["Tên", "Vai trò", "Số tờ", "Số thửa", "Tài liệu"], rows)
    before = upload.read_bytes()
    documents = tmp_path / "documents"; documents.mkdir()
    writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
    with (documents / tbxn_name).open("wb") as stream: writer.write(stream)
    (documents / ddk_name).write_bytes(_create_signed_pdf(tmp_path).read_bytes())
    excel = ExcelConfig(upload, "Data", 1, {
        "person": "A", "role": "B", "sheet": "C", "parcel": "D", "combined": "E",
    })

    results = []
    for separate in (False, True):
        config = RunConfig(
            excel, None, documents, documents, tmp_path / "outputs", "A", "E", True,
            ValidationMode.VBDLIS_ONLY, separate,
        )
        results.append(ValidationService().run(config))
    for result in results:
        assert result.stats["total_parcels"] == 1
        assert result.stats["need_signature"] == 1
        assert result.stats["missing_document"] == 0
        assert result.records[0].document_completeness == DocumentCompleteness.COMPLETE
        assert result.records[0].workflow_status == WorkflowStatus.NEED_SIGNATURE
        need_book = load_workbook(result.need_signature_excel, read_only=True)
        try:
            assert need_book["Data"].max_row == 6
            assert need_book["PARCEL_DETAILS"].max_row == 2
        finally: need_book.close()
        copied = list((result.document_actions / "NEED_SIGNATURE" / "TBXN_UNSIGNED").rglob("*.pdf"))
        assert len(copied) == 1
    assert results[0].valid_excel == results[0].document_complete_excel
    assert results[1].valid_excel == results[1].ready_excel
    assert upload.read_bytes() == before
