from pathlib import Path
from openpyxl import Workbook
from tools.excel_safety import atomic_save, style_report


def export_report(result, path: str | Path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = workbook.create_sheet("Summary")
    summary.append(["Chỉ số", "Số lượng"])
    for key, value in result.counts().items(): summary.append([key, value])
    names = workbook.create_sheet("Name Changes")
    names.append(["Row", "Original Name", "Normalized Name", "Changed?", "Status", "Warning"])
    birth_headers = ["Row", "Original Value", "Normalized Value", "Input Format", "Output Format", "Status", "Warning"]
    births = workbook.create_sheet("Birthdate Changes")
    births.append(birth_headers)
    warnings = workbook.create_sheet("Warnings")
    warnings.append(["Row", "Field", "Original", "Status", "Warning"])
    invalid = workbook.create_sheet("Invalid Dates"); invalid.append(birth_headers)
    ambiguous = workbook.create_sheet("Ambiguous Dates"); ambiguous.append(birth_headers)
    for change in result.changes:
        if change.field == "name":
            names.append([change.row, change.original, change.display_value, change.changed, change.status, change.warning])
        else:
            row = [change.row, change.original, change.display_value, result.config.input_mode,
                   result.config.output_format, change.status, change.warning]
            births.append(row)
            if change.status == "INVALID": invalid.append(row)
            if change.status == "AMBIGUOUS": ambiguous.append(row)
        if change.warning: warnings.append([change.row, change.field, change.original, change.status, change.warning])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    style_report(workbook)
    try:
        return atomic_save(workbook, path)
    finally:
        workbook.close()
