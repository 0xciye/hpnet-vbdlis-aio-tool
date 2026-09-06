from __future__ import annotations

from pathlib import Path

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string

from tools.excel_safety import atomic_save, create_backup, open_workbook
from .date_normalizer import FORMATS, normalize_birthdate
from .models import CellChange, NormalizeConfig, NormalizeScanResult
from .name_normalizer import normalize_person_name


EXCEL_FORMATS = {"DD/MM/YYYY": "dd/mm/yyyy", "DD-MM-YYYY": "dd-mm-yyyy", "YYYY-MM-DD": "yyyy-mm-dd"}


def _inspect_name(cell, row, column):
    if isinstance(cell, MergedCell):
        return CellChange(row, "name", column, None, status="MERGED_CELL", warning="Không ghi vào ô gộp không phải ô neo.")
    if cell.data_type == "f":
        return CellChange(row, "name", column, cell.value, status="FORMULA_CELL", warning="Công thức được bảo vệ.")
    result = normalize_person_name(cell.value)
    status = "EMPTY" if result.normalized is None else ("CHANGE" if result.changed else "UNCHANGED")
    return CellChange(row, "name", column, cell.value, result.normalized, result.normalized or "", status,
                      "; ".join(result.warnings), result.changed, result.changed)


def _inspect_date(cell, row, column, config, epoch):
    if isinstance(cell, MergedCell):
        return CellChange(row, "birthdate", column, None, status="MERGED_CELL", warning="Không ghi vào ô gộp không phải ô neo.")
    if cell.data_type == "f":
        return CellChange(row, "birthdate", column, cell.value, status="FORMULA_CELL", warning="Công thức được bảo vệ.")
    result = normalize_birthdate(cell.value, config.input_mode, config.output_format,
                                 number_format=cell.number_format, epoch=epoch)
    if result.status == "VALID" and config.store_as_date and hasattr(cell.value, "year"):
        result.changed = cell.number_format.lower() != EXCEL_FORMATS[config.output_format]
    status = "CHANGE" if result.status == "VALID" and result.changed else ("UNCHANGED" if result.status == "VALID" else result.status)
    normalized = result.parsed_value if config.store_as_date else result.display_value
    return CellChange(row, "birthdate", column, cell.value, normalized, result.display_value, status,
                      result.warning, status == "CHANGE", status == "CHANGE")


def scan(config: NormalizeConfig) -> NormalizeScanResult:
    if not config.normalize_names and not config.normalize_birthdates:
        raise ValueError("Hãy chọn ít nhất một trường cần chuẩn hóa.")
    if config.end_row is not None and config.end_row < config.start_row:
        raise ValueError("Dòng kết thúc phải lớn hơn hoặc bằng dòng bắt đầu.")
    selected_columns = []
    if config.normalize_names: selected_columns.append(config.name_column)
    if config.normalize_birthdates: selected_columns.append(config.birthdate_column)
    for column in selected_columns: column_index_from_string(column)
    if len(selected_columns) == 2 and selected_columns[0] == selected_columns[1]:
        raise ValueError("Cột Họ tên và Ngày sinh phải khác nhau khi bật cả hai.")
    workbook = open_workbook(config.source, read_only=False, data_only=False)
    try:
        if config.sheet_name not in workbook.sheetnames:
            raise ValueError("Sheet đã chọn không tồn tại.")
        sheet = workbook[config.sheet_name]
        end = min(config.end_row or sheet.max_row, sheet.max_row)
        changes = []
        for row in range(config.start_row, end + 1):
            if config.normalize_names:
                changes.append(_inspect_name(sheet[f"{config.name_column}{row}"], row, config.name_column))
            if config.normalize_birthdates:
                changes.append(_inspect_date(sheet[f"{config.birthdate_column}{row}"], row, config.birthdate_column, config, workbook.epoch))
        return NormalizeScanResult(config, changes, max(0, end - config.start_row + 1))
    finally:
        workbook.close()


def apply_changes(result: NormalizeScanResult, output: str | Path, selected: set[tuple[int, str]] | None = None):
    planned = {(change.row, change.field) for change in result.changes if change.status == "CHANGE"}
    selected = planned if selected is None else set(selected)
    if not selected <= planned:
        raise ValueError("Có ô không thuộc kế hoạch thay đổi đã xác nhận.")
    if result.config.source.suffix.lower() == ".xlsm" and Path(output).suffix.lower() != ".xlsm":
        raise ValueError("Nguồn .xlsm phải lưu đầu ra .xlsm để bảo toàn macro.")
    backup = create_backup(result.config.source, "backup_before_normalize")
    workbook = open_workbook(result.config.source, read_only=False, data_only=False)
    try:
        sheet = workbook[result.config.sheet_name]
        by_key = {(change.row, change.field): change for change in result.changes}
        for key in selected:
            change = by_key[key]
            cell = sheet[f"{change.column}{change.row}"]
            if isinstance(cell, MergedCell) or cell.data_type == "f":
                raise ValueError(f"Ô {change.column}{change.row} hiện không thể ghi an toàn.")
            cell.value = change.normalized
            if change.field == "birthdate" and result.config.store_as_date:
                cell.number_format = EXCEL_FORMATS[result.config.output_format]
        output = atomic_save(workbook, output)
    finally:
        workbook.close()
    return output, backup
