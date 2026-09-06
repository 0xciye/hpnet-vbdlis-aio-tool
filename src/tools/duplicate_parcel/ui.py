from __future__ import annotations

import logging
import os
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtWidgets import QHeaderView
from PySide6.QtWidgets import (QCheckBox, QComboBox, QFileDialog, QFormLayout, QHBoxLayout, QLabel,
    QLineEdit, QMainWindow, QMessageBox, QPushButton, QSpinBox, QTableWidget, QTableWidgetItem,
    QVBoxLayout, QWidget)
from openpyxl.utils import get_column_letter, column_index_from_string

from launcher_ui.theme import STYLE, palette
from tools.excel_safety import default_output, open_workbook, timestamp
from tools.qt_worker import Worker
from .models import ScanConfig
from .report import export_report
from .service import apply_cleanup, scan


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Kiểm tra & Làm sạch thửa trùng")
        self.resize(1180, 760); self.setMinimumSize(900, 620)
        self.setStyleSheet(STYLE); self.setPalette(palette())
        self.result = None; self.visible_records = []; self.worker = None
        self._build()

    def _build(self):
        root = QWidget(); box = QVBoxLayout(root); box.setContentsMargins(24, 20, 24, 20); box.setSpacing(12)
        title = QLabel("KIỂM TRA & LÀM SẠCH THỬA TRÙNG"); title.setObjectName("pageTitle"); box.addWidget(title)
        box.addWidget(QLabel("Tách hộ bằng dòng Tổng DT; mọi dòng cùng Số tờ + Số thửa bị trùng, kể cả khác hộ, sẽ được chọn clear sau khi xem trước và xác nhận."))
        form = QFormLayout(); self.source = QLineEdit(); browse = QPushButton("Chọn…"); browse.clicked.connect(self.browse)
        row = QHBoxLayout(); row.addWidget(self.source, 1); row.addWidget(browse); form.addRow("File Excel", row)
        self.sheet = QComboBox(); form.addRow("Sheet", self.sheet)
        self.sheet.currentTextChanged.connect(self._refresh_column_choices)
        self.columns = {}
        for label, key, value in (("Cột Họ và tên", "name", "B"), ("Cột Số tờ", "sheet", "G"),
                                  ("Cột Số thửa", "parcel", "H"), ("Cột Diện tích", "area", "I"),
                                  ("Cột Xứ đồng", "location", "J"), ("Cột bắt đầu clear", "clear_start", "G"),
                                  ("Cột kết thúc clear", "clear_end", "X")):
            edit = QComboBox(); edit.setProperty("default_column", value)
            edit.addItem("— Chọn file Excel trước —", None)
            edit.setMinimumWidth(240); edit.setMaximumWidth(320)
            self.columns[key] = edit; form.addRow(label, edit)
        self.start = QSpinBox(); self.start.setRange(1, 1_048_576); self.start.setValue(4); form.addRow("Dòng bắt đầu", self.start)
        box.addLayout(form)
        actions = QHBoxLayout(); self.scan_button = QPushButton("QUÉT & XEM TRƯỚC"); self.scan_button.clicked.connect(self.scan_file)
        self.report_button = QPushButton("Xuất báo cáo"); self.report_button.clicked.connect(self.export); self.report_button.setEnabled(False)
        self.apply_button = QPushButton("Áp dụng làm sạch"); self.apply_button.clicked.connect(self.apply); self.apply_button.setEnabled(False)
        actions.addWidget(self.scan_button); actions.addStretch(); actions.addWidget(self.report_button); actions.addWidget(self.apply_button); box.addLayout(actions)
        self.summary = QLabel("Chưa quét dữ liệu."); box.addWidget(self.summary)
        self.table = QTableWidget(0, 11); self.table.setHorizontalHeaderLabels(["Chọn", "Hộ", "Chủ hộ", "Dòng", "Số tờ", "Số thửa", "Diện tích", "Xứ đồng", "Trạng thái", "Hành động", "Trùng với"])
        header = self.table.horizontalHeader()
        header.setStretchLastSection(True)
        header.setSectionResizeMode(QHeaderView.ResizeToContents)
        header.setMinimumSectionSize(72)
        box.addWidget(self.table, 1)
        self.setCentralWidget(root); self.statusBar().showMessage("Quét chỉ đọc dữ liệu và không sửa file.")

    def browse(self):
        path, _ = QFileDialog.getOpenFileName(self, "Chọn file Excel", "", "Excel (*.xlsx *.xlsm)")
        if not path: return
        self.source.setText(path)
        try:
            wb = open_workbook(path, read_only=True); self.sheet.clear(); self.sheet.addItems(wb.sheetnames); wb.close()
        except Exception as error: self._error(error)

    def _refresh_column_choices(self):
        if not self.source.text() or not self.sheet.currentText():
            return
        try:
            wb = open_workbook(self.source.text(), read_only=True, data_only=True)
            ws = wb[self.sheet.currentText()]
            max_column = max(26, ws.max_column)
            headers = {}
            for column in range(1, max_column + 1):
                letter = get_column_letter(column)
                for row in range(1, min(ws.max_row, 5) + 1):
                    value = ws.cell(row, column).value
                    if value is not None and str(value).strip():
                        headers[letter] = " ".join(str(value).split())
                        break
            for combo in self.columns.values():
                selected = combo.currentData() or combo.property("default_column")
                combo.blockSignals(True); combo.clear()
                for column in range(1, max_column + 1):
                    letter = get_column_letter(column)
                    label = f"{letter} — {headers[letter]}" if letter in headers else letter
                    combo.addItem(label, letter)
                index = combo.findData(selected)
                combo.setCurrentIndex(index if index >= 0 else 0)
                combo.blockSignals(False)
            wb.close()
        except Exception as error:
            self._error(error)

    def config(self):
        values = {key: str(value.currentData() or value.currentText()).strip().upper() for key, value in self.columns.items()}
        if not self.sheet.currentText(): raise ValueError("Hãy chọn file và sheet Excel.")
        return ScanConfig(Path(self.source.text()), self.sheet.currentText(), values["name"], values["sheet"], values["parcel"],
                          values["area"], values["location"], self.start.value(), values["clear_start"], values["clear_end"])

    def _run(self, function, done):
        if self.worker: return
        self.setEnabled(False); self.statusBar().showMessage("Đang xử lý…")
        self.worker = Worker(function, self); self.worker.succeeded.connect(done); self.worker.failed.connect(self._error)
        self.worker.finished.connect(self._finished); self.worker.start()

    def _finished(self):
        worker = self.worker; self.worker = None; self.setEnabled(True); self.statusBar().showMessage("Sẵn sàng.")
        if worker: worker.deleteLater()

    def scan_file(self):
        try: config = self.config()
        except Exception as error: self._error(error); return
        self._run(lambda: scan(config), self.show_result)

    def show_result(self, result):
        self.result = result; counts = result.counts()
        conflicts = sum(counts.get(status, 0) for status in ("SAME_PARCEL_DIFFERENT_AREA",
                                                               "SAME_PARCEL_DIFFERENT_LOCATION",
                                                               "SAME_PARCEL_DIFFERENT_AREA_AND_LOCATION",
                                                               "INCOMPLETE_DATA"))
        self.summary.setText(f"Đã quét {result.rows_scanned} dòng · {counts.get('HOUSEHOLDS', 0)} hộ · "
                             f"{counts.get('EXACT_DUPLICATE', 0)} dòng sẽ clear · "
                             f"{conflicts} dòng xung đột/thiếu dữ liệu cần xem lại")
        exact_groups = {}
        for record in result.records:
            if record.status == "EXACT_DUPLICATE":
                exact_groups.setdefault((record.sheet, record.parcel), []).append(record)
        self.visible_records = [record for record in result.records if record.status == "EXACT_DUPLICATE"]
        by_row = {record.row: record for record in self.visible_records}
        self.table.setRowCount(len(self.visible_records))
        for index, record in enumerate(self.visible_records):
            check = QTableWidgetItem(); check.setFlags(Qt.ItemIsEnabled | Qt.ItemIsUserCheckable)
            check.setCheckState(Qt.Checked if record.selected else Qt.Unchecked)
            self.table.setItem(index, 0, check)
            related = [by_row[row] for row in record.related_rows if row in by_row]
            related_text = "; ".join(
                f"Hộ {other.household} — {other.name or '(không có tên)'} — dòng {other.row}"
                for other in related
            )
            values = (record.household, record.name or "(không có tên)", record.row,
                      record.sheet, record.parcel, record.area or "", record.location,
                      record.status, record.action, related_text)
            for column, value in enumerate(values, 1):
                self.table.setItem(index, column, QTableWidgetItem(str(value)))
        self.report_button.setEnabled(True); self.apply_button.setEnabled(bool(result.clear_rows))
        self._log("scan=" + repr(result.counts()) + "\n")

    def selected_rows(self):
        return {self.visible_records[row].row for row in range(self.table.rowCount()) if self.table.item(row, 0).checkState() == Qt.Checked}

    def export(self):
        if not self.result: return
        path, _ = QFileDialog.getSaveFileName(self, "Lưu báo cáo", str(Path(self.source.text()).with_name("duplicate_parcel_report.xlsx")), "Excel (*.xlsx)")
        if path: self._run(lambda: export_report(self.result, path), lambda value: QMessageBox.information(self, "Đã xuất báo cáo", str(value)))

    def apply(self):
        rows = self.selected_rows()
        if not rows: self._error("Chưa chọn dòng duplicate nào để clear."); return
        message = f"Sắp clear dữ liệu trong {len(rows)} dòng duplicate.\nBackup sẽ được tạo trước khi chỉnh sửa."
        if QMessageBox.question(self, "Xác nhận làm sạch", message, QMessageBox.Yes | QMessageBox.No, QMessageBox.No) != QMessageBox.Yes: return
        output, _ = QFileDialog.getSaveFileName(self, "Lưu file đã làm sạch", str(default_output(self.source.text(), "cleaned")), "Excel (*.xlsx *.xlsm)")
        if not output: return
        self._run(lambda: apply_cleanup(self.result, output, rows), self._applied)

    def _applied(self, value):
        output, backup = value; self._log(f"output={output}\nbackup={backup}\n")
        QMessageBox.information(self, "Đã làm sạch", f"File mới: {output}\nBackup: {backup}")

    def _log(self, text):
        folder = Path(os.getenv("APPDATA", str(Path.home()))) / "HPNET_VBDLIS_Tools" / "logs"; folder.mkdir(parents=True, exist_ok=True)
        (folder / f"parcel_cleanup_{timestamp().replace('-', '')}.log").write_text(text, encoding="utf-8")

    def _error(self, error):
        self._log(f"ERROR={error}\n")
        QMessageBox.warning(self, "Không thể xử lý file Excel", f"{error}\n\nHãy kiểm tra file có đang mở, bị khóa hoặc bạn không có quyền truy cập.")
