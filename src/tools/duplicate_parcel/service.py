from __future__ import annotations

import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl.utils import column_index_from_string

from tools.excel_safety import atomic_save, create_backup, open_workbook
from .models import ParcelRecord, ScanConfig, ScanResult


def normalized_text(value) -> str:
    return " ".join(str(value).split()).casefold() if value is not None else ""


def is_total_dt(value) -> bool:
    return bool(re.fullmatch(r"tổng\s*dt", " ".join(str(value or "").split()), re.IGNORECASE))


def normalized_identifier(value) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        return ""
    text = str(value).strip()
    try:
        number = Decimal(text)
        if number == number.to_integral_value():
            return str(number.quantize(Decimal("1")))
        return format(number.normalize(), "f")
    except InvalidOperation:
        return text.casefold()


def normalized_area(value) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return format(Decimal(str(value).strip()).normalize(), "f")
    except InvalidOperation:
        return normalized_text(value)


def _classify_group(records: list[ParcelRecord]) -> None:
    if len(records) < 2:
        return
    signatures: dict[tuple[str | None, str], list[ParcelRecord]] = defaultdict(list)
    for record in records:
        signatures[(record.area, record.location)].append(record)
    for same in signatures.values():
        if len(same) < 2:
            continue
        for duplicate in same:
            duplicate.status = "EXACT_DUPLICATE"
            duplicate.action = "Clear"
            duplicate.related_rows = [other.row for other in same if other.row != duplicate.row]
            duplicate.selected = True
    representatives = [same[0] for same in signatures.values()]
    for record in representatives:
        if record.status == "EXACT_DUPLICATE":
            continue
        others = [other for other in representatives if other is not record]
        if not others:
            continue
        if record.area is None or any(other.area is None for other in others):
            status = "INCOMPLETE_DATA"
        else:
            different_area = any(other.area != record.area for other in others)
            different_location = any(other.location != record.location for other in others)
            if different_area and different_location:
                status = "SAME_PARCEL_DIFFERENT_AREA_AND_LOCATION"
            elif different_area:
                status = "SAME_PARCEL_DIFFERENT_AREA"
            elif different_location:
                status = "SAME_PARCEL_DIFFERENT_LOCATION"
            else:
                continue
        record.status = status
        record.action = "Xem lại"
        record.related_rows = [other.row for other in others]


def scan(config: ScanConfig) -> ScanResult:
    for column in (config.name_column, config.sheet_column, config.parcel_column, config.area_column,
                   config.location_column, config.clear_start_column, config.clear_end_column):
        column_index_from_string(column)
    if column_index_from_string(config.clear_start_column) > column_index_from_string(config.clear_end_column):
        raise ValueError("Cột bắt đầu clear phải đứng trước cột kết thúc.")
    workbook = open_workbook(config.source, read_only=False, data_only=False)
    try:
        if config.sheet_name not in workbook.sheetnames:
            raise ValueError("Sheet đã chọn không tồn tại.")
        sheet = workbook[config.sheet_name]
        household = 1
        summary_rows: list[int] = []
        records: list[ParcelRecord] = []
        groups: dict[tuple[int, str, str], list[ParcelRecord]] = defaultdict(list)
        for row in range(config.start_row, sheet.max_row + 1):
            if is_total_dt(sheet[f"{config.name_column}{row}"].value):
                summary_rows.append(row)
                household += 1
                continue
            sheet_number = normalized_identifier(sheet[f"{config.sheet_column}{row}"].value)
            parcel_number = normalized_identifier(sheet[f"{config.parcel_column}{row}"].value)
            if not sheet_number or not parcel_number:
                continue
            record = ParcelRecord(household, row, sheet_number, parcel_number,
                                  normalized_area(sheet[f"{config.area_column}{row}"].value),
                                  normalized_text(sheet[f"{config.location_column}{row}"].value))
            records.append(record)
            groups[(household, sheet_number, parcel_number)].append(record)
        for group in groups.values():
            _classify_group(group)

        global_groups: dict[tuple[str, str], list[ParcelRecord]] = defaultdict(list)
        for record in records:
            global_groups[(record.sheet, record.parcel)].append(record)
        for group in global_groups.values():
            if len(group) < 2:
                continue
            rows = [other.row for other in group]
            for record in group:
                record.status = "EXACT_DUPLICATE"
                record.action = "Clear"
                record.related_rows = [row for row in rows if row != record.row]
                record.selected = True
        return ScanResult(config, records, summary_rows, max(0, sheet.max_row - config.start_row + 1))
    finally:
        workbook.close()


def apply_cleanup(result: ScanResult, output: str | Path, selected_rows: set[int] | None = None):
    selected_rows = set(result.clear_rows if selected_rows is None else selected_rows)
    allowed = set(result.clear_rows)
    if not selected_rows <= allowed:
        raise ValueError("Có dòng không thuộc kế hoạch duplicate đã xác nhận.")
    if result.config.source.suffix.lower() == ".xlsm" and Path(output).suffix.lower() != ".xlsm":
        raise ValueError("Nguồn .xlsm phải lưu đầu ra .xlsm để bảo toàn macro.")
    backup = create_backup(result.config.source, "backup_before_cleanup")
    workbook = open_workbook(result.config.source, read_only=False, data_only=False)
    try:
        sheet = workbook[result.config.sheet_name]
        start = column_index_from_string(result.config.clear_start_column)
        end = column_index_from_string(result.config.clear_end_column)
        if start > end:
            raise ValueError("Cột bắt đầu clear phải đứng trước cột kết thúc.")
        for row in selected_rows:
            if row in result.summary_rows or is_total_dt(sheet[f"{result.config.name_column}{row}"].value):
                raise ValueError(f"Dòng Tổng DT {row} được bảo vệ, không thể clear.")
            for column in range(start, end + 1):
                sheet.cell(row, column).value = None
        output = atomic_save(workbook, output)
    finally:
        workbook.close()
    return output, backup
