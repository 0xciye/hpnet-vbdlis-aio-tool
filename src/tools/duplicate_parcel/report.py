from pathlib import Path
from openpyxl import Workbook
from tools.excel_safety import atomic_save, style_report


HEADERS = ["Hộ", "Dòng", "Số tờ", "Số thửa", "Diện tích", "Xứ đồng", "Trạng thái", "Hành động", "Dòng liên quan"]


def export_report(result, path: str | Path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = workbook.create_sheet("Summary")
    summary.append(["Chỉ số", "Số lượng"])
    for key, value in result.counts().items():
        summary.append([key, value])
    sections = {
        "Exact Duplicates": {"EXACT_DUPLICATE"},
        "Parcel Conflicts": {"SAME_PARCEL_DIFFERENT_AREA", "SAME_PARCEL_DIFFERENT_LOCATION", "SAME_PARCEL_DIFFERENT_AREA_AND_LOCATION"},
        "Cross Household": {"CROSS_HOUSEHOLD_DUPLICATE"},
        "Incomplete Data": {"INCOMPLETE_DATA"},
    }
    for title, statuses in sections.items():
        sheet = workbook.create_sheet(title)
        sheet.append(HEADERS)
        for record in result.records:
            if record.status in statuses:
                sheet.append([record.household, record.row, record.sheet, record.parcel, record.area, record.location,
                              record.status, record.action, ", ".join(map(str, record.related_rows))])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    style_report(workbook)
    try:
        return atomic_save(workbook, path)
    finally:
        workbook.close()
