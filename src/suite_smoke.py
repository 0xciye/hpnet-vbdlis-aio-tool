"""Offline release checks; only creates synthetic files in a temporary directory."""
from datetime import datetime, timedelta, timezone
from pathlib import Path
from tempfile import TemporaryDirectory


def run(hub):
    from openpyxl import Workbook, load_workbook

    from launcher_ui.help_page import guide_sources
    from tools.hpnet_file_generator.core.file_generator import FileGenerator
    from tools.hpnet_file_generator.models.data_models import (
        ActionStatus,
        GenerationAction,
        Parcel,
        SourceFile,
    )
    from tools.signed_pdf_cleaner.core.models import ProcessStatus
    from tools.signed_pdf_cleaner.core.processor import FileProcessor
    from tools.signed_pdf_cleaner.core.scanner import FileScanner
    from tools.vbdlis_excel_builder.core.service import BuilderService
    from tools.vbdlis_excel_builder.models import MappingProfile
    from tools.vbdlis_excel_builder.utils.paths import resource_path
    from tools.vbdlis_validation.documents import SignatureValidator
    from tools.vbdlis_validation.messages import describe_issues, recommended_actions
    from tools.vbdlis_validation.models import ExcelConfig
    from tools.vbdlis_validation.service import headers

    # Presentation-only checks also run inside the packaged Windows executable.
    hub.open_help()
    view = hub.launcher_view
    assert view.pages.currentIndex() == 1
    assert view.help_page.contents.count() == view.help_page.documents.count() == 3
    for index, (_, path) in enumerate(guide_sources()):
        view.help_page.contents.setCurrentRow(index)
        displayed = " ".join(view.help_page.browser.toPlainText().split())
        assert path.is_file(), str(path)
        for line in path.read_text(encoding="utf-8-sig").splitlines():
            line = line.strip()
            if line and not set(line) <= {"=", "-"}:
                assert " ".join(line.split()) in displayed, line
    view.help_page.contents.setCurrentRow(0)
    view.help_page.search.setText("giấy tờ nhân thân")
    view.help_page.find_text()
    assert "Đã tìm thấy" in view.help_page.search_status.text()
    view.reset_search()
    assert len(view.visible_tool_keys) == 10

    hub.launch_excel_builder()
    hub.launch_vbdlis_validation()
    hub.launch_auto_rename()
    hub.launch_pdf_cleaner()
    hub.launch_notice_builder()
    hub.launch_duplicate_parcel()
    hub.launch_data_normalizer()
    assert set(hub.tool_windows) == {"excel", "validation", "rename", "cleaner", "notice", "duplicate_parcel", "data_normalizer"}
    assert all(not window.windowIcon().isNull() for window in hub.tool_windows.values())
    assert hub.tool_windows["validation"].windowIcon().cacheKey() != hub.tool_windows["excel"].windowIcon().cacheKey()
    from PySide6.QtWidgets import QListView, QStyle
    notice_window = hub.tool_windows["notice"]
    assert notice_window.steps.count() == notice_window.stack.count() == 8
    for combo in notice_window.mapping.values():
        assert isinstance(combo.view(), QListView)
        assert combo.style().styleHint(QStyle.SH_ComboBox_Popup, None, combo) == 0
    from tools.notice_builder.smoke import run as notice_smoke
    notice = notice_smoke(hub.tool_windows["notice"])
    with TemporaryDirectory(prefix="suite-release-smoke-") as temporary:
        root = Path(temporary)
        source = root / "synthetic.xlsx"
        wb = Workbook()
        wb.active.title = "Du lieu"
        wb.active.append(["STT", "Họ tên", "CCCD", "Số tờ", "Số thửa", "Diện tích"])
        wb.active.append([1, "Nguyễn Văn Kiểm Thử", "031080001234", 10, 300, 100])
        wb.active.append([None, "TỔNG DT", None, None, None, 100])
        wb.save(source)
        wb.close()
        original = source.read_bytes()
        profile = MappingProfile(commune_code="10930", address="Địa chỉ kiểm thử",
            source_mapping={"household_stt": "A", "person_name": "B", "cccd": "C",
                            "sheet_number": "D", "parcel_number": "E", "area": "F"},
            field_rules={"MUC_29": {"mode": "fixed", "default": "Không xác định", "required": True}})
        service = BuilderService(resource_path("resources/BieuMauThuThapThongTinGiayChungNhan.xlsx"),
                                 resource_path("config/template_schema.json"), root / "logs")
        result = service.process(source, "Du lieu", 1, profile)
        assert result.can_export, repr(result.issues)
        assert len(result.rows) == 1
        assert result.stats["summary_rows_skipped"] == 1
        output, report, verification = service.export(result, profile, root, "output.xlsx")
        assert verification["pass"] and report.is_file()
        assert len(result.diagnostic_files.paths) == 2
        saved = load_workbook(output, data_only=True)
        try:
            sheet = saved.worksheets[0]
            assert sheet["AD5"].value == "Không xác định", sheet["AD5"].value
            assert sheet["K5"].value == "Nam", sheet["K5"].value
            assert sheet["I5"].value == "031080001234"
        finally:
            saved.close()
        assert source.read_bytes() == original

        two_level = root / "two-level.xlsx"
        two_level_book = Workbook(); two_level_sheet = two_level_book.active
        two_level_sheet.title = "Data"; two_level_sheet.merge_cells("A1:B1")
        two_level_sheet["A1"] = "Hộ gia đình"; two_level_sheet.append(["Tên chủ hộ", "Mã hộ"])
        two_level_sheet.append(["Nguyễn Văn A", "H001"]); two_level_book.save(two_level)
        two_level_book.close()
        assert headers(ExcelConfig(two_level, "Data", 1, {}, 2))[1][1] == "Hộ gia đình — Mã hộ"
        assert "Chưa tìm thấy file Đơn đăng ký" in describe_issues(["MISSING_DDK"])
        assert "Bổ sung đúng file DDK" in recommended_actions(["MISSING_DDK"])

        pdf = root / "original.pdf"
        pdf.write_bytes(b"%PDF-offline-synthetic-test")
        action = GenerationAction(SourceFile(pdf, pdf.name, "test", ".pdf"),
                                  Parcel("10", "300"), "DDK", "copy.pdf", root / "copy.pdf")
        assert FileGenerator().execute_action(action).status == ActionStatus.SUCCESS
        assert action.target_path.read_bytes() == pdf.read_bytes()

        clean_dir = root / "cleaner"
        clean_dir.mkdir()
        signed = clean_dir / "sample.signed.pdf"
        signed.write_bytes(pdf.read_bytes())
        plans = FileScanner().scan_directory(str(clean_dir))
        assert len(plans) == 1
        assert FileProcessor(use_recycle_bin=True).process_plan(plans[0]).status == ProcessStatus.SKIPPED
        assert signed.read_bytes() == pdf.read_bytes()
        assert not (clean_dir / "sample.pdf").exists()

        # Exercise the real CMS/ByteRange verifier inside the frozen package.
        from cryptography import x509
        from cryptography.hazmat.primitives import hashes, serialization
        from cryptography.hazmat.primitives.asymmetric import rsa
        from cryptography.hazmat.primitives.serialization import pkcs12
        from cryptography.x509.oid import NameOID
        from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
        from pyhanko.sign import signers
        from pypdf import PdfWriter
        unsigned_pdf = root / "signature-input.pdf"
        pdf_writer = PdfWriter(); pdf_writer.add_blank_page(width=200, height=200)
        with unsigned_pdf.open("wb") as stream: pdf_writer.write(stream)
        key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        subject = issuer = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "Frozen Smoke")])
        now = datetime.now(timezone.utc)
        certificate = (x509.CertificateBuilder().subject_name(subject).issuer_name(issuer)
                       .public_key(key.public_key()).serial_number(x509.random_serial_number())
                       .not_valid_before(now - timedelta(days=1)).not_valid_after(now + timedelta(days=1))
                       .add_extension(x509.BasicConstraints(ca=True, path_length=None), critical=True)
                       .sign(key, hashes.SHA256()))
        pfx = root / "smoke.p12"
        pfx.write_bytes(pkcs12.serialize_key_and_certificates(
            b"smoke", key, certificate, None, serialization.BestAvailableEncryption(b"smoke")
        ))
        signer = signers.SimpleSigner.load_pkcs12(str(pfx), passphrase=b"smoke")
        signed_pdf = root / "signature-signed.pdf"
        with unsigned_pdf.open("rb") as source, signed_pdf.open("wb") as target:
            signers.sign_pdf(IncrementalPdfFileWriter(source), signers.PdfSignatureMetadata(field_name="Sig"), signer=signer, output=target)
        signature_result = SignatureValidator().validate(signed_pdf)
        assert signature_result.cryptographic_integrity, signature_result.message

    return {"status": "PASS", "embedded_help_complete": True, "independent_help_pages": 3, "launcher_tools": 10,
            "python_windows": 7, "python_window_icons": True, "notice_builder": notice, "notice_styled_popups": True, "excel_export": True,
            "summary_row_excluded": True, "fixed_item_29": True, "cccd_gender": True,
            "diagnostic_reports": True, "source_unchanged": True,
            "auto_rename_copy": True, "cleaner_fake_signature_blocked": True,
            "vbdlis_signature_validation": True, "validation_two_level_headers": True,
            "validation_friendly_guidance": True, "validation_unique_icon": True,
            "hpnet_live_operations": False}
