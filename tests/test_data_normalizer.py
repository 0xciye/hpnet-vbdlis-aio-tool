from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from tools.data_normalizer.date_normalizer import normalize_birthdate
from tools.data_normalizer.models import NormalizeConfig
from tools.data_normalizer.name_normalizer import normalize_person_name
from tools.data_normalizer.report import export_report
from tools.data_normalizer.service import apply_changes, scan


def test_vietnamese_names_whitespace_unicode_and_punctuation():
    examples = {
        " NGUYỄN   VĂN  A ": "Nguyễn Văn A",
        "trần thị b": "Trần Thị B",
        "PHẠM VĂN C": "Phạm Văn C",
        "\t ĐỖ THỊ ÁNH \n": "Đỗ Thị Ánh",
        "JEAN-PIERRE D'ARC": "Jean-Pierre D'Arc",
        "HÀ-AN": "Hà-An",
    }
    for original, expected in examples.items(): assert normalize_person_name(original).normalized == expected
    assert normalize_person_name(None).normalized is None


def test_controlled_date_parsing_and_leap_years():
    assert normalize_birthdate("1-2-1990", "DMY").display_value == "01/02/1990"
    assert normalize_birthdate("1990-02-01", "YMD").parsed_value == date(1990, 2, 1)
    assert normalize_birthdate("03/04/1990", "AUTO").status == "AMBIGUOUS"
    assert normalize_birthdate("25/03/1990", "AUTO").parsed_value == date(1990, 3, 25)
    assert normalize_birthdate("31/02/1990", "DMY").status == "INVALID"
    assert normalize_birthdate("29/02/2024", "DMY").status == "VALID"
    assert normalize_birthdate("29/02/2023", "DMY").status == "INVALID"
    assert normalize_birthdate("1990", "DMY").status == "PARTIAL"
    assert normalize_birthdate("02/1990", "DMY").status == "PARTIAL"
    assert normalize_birthdate("01/02/90", "DMY").status == "PARTIAL"
    assert normalize_birthdate(45000, "DMY", number_format="General").status == "PARTIAL"
    assert normalize_birthdate(datetime(2026, 8, 30), "DMY").display_value == "30/08/2026"


def make_source(path):
    wb = Workbook(); ws = wb.active; ws.title = "People"
    ws.append(["Họ tên", "Ngày sinh"]); ws.append([]); ws.append([])
    ws.append([" NGUYỄN   VĂN  A ", "25-2-1990"])
    ws.append(["Đỗ Thị Ánh", "03/04/1990"])
    ws.append(["Trần Thị B", "31/02/1990"])
    ws.append(["=UPPER(\"abc\")", "=DATE(1990,2,1)"])
    ws.append([None, "1990"])
    ws["A7"].data_type = "f"
    wb.save(path); wb.close()


def test_scan_formula_ambiguous_invalid_partial_apply_and_report(tmp_path):
    source = tmp_path / "people.xlsx"; make_source(source); original = source.read_bytes()
    config = NormalizeConfig(source, "People", 4, None, True, "A", True, "B", "AUTO", "DD/MM/YYYY", True)
    result = scan(config); keyed = {(item.row, item.field): item for item in result.changes}
    assert keyed[(4, "name")].display_value == "Nguyễn Văn A"
    assert keyed[(4, "birthdate")].status == "CHANGE"
    assert keyed[(5, "birthdate")].status == "AMBIGUOUS"
    assert keyed[(6, "birthdate")].status == "INVALID"
    assert keyed[(7, "name")].status == keyed[(7, "birthdate")].status == "FORMULA_CELL"
    assert keyed[(8, "birthdate")].status == "PARTIAL"
    output, backup = apply_changes(result, tmp_path / "normalized.xlsx")
    assert source.read_bytes() == original and backup.read_bytes() == original
    wb = load_workbook(output, data_only=False); ws = wb["People"]
    assert ws["A4"].value == "Nguyễn Văn A" and ws["B4"].value == datetime(1990, 2, 25)
    assert ws["B5"].value == "03/04/1990" and ws["B6"].value == "31/02/1990"
    assert ws["A7"].value.startswith("=") and ws["B7"].value.startswith("=")
    wb.close()
    report = export_report(result, tmp_path / "normalization_report.xlsx")
    wb = load_workbook(report, read_only=False); assert {"Summary", "Name Changes", "Birthdate Changes", "Warnings", "Invalid Dates", "Ambiguous Dates"} <= set(wb.sheetnames)
    assert wb["Birthdate Changes"].column_dimensions["G"].width >= 30
    assert all(cell.data_type != "f" for sheet in wb.worksheets for row in sheet.iter_rows() for cell in row)
    wb.close()


def test_name_and_birthdate_modes_are_independent(tmp_path):
    source = tmp_path / "people.xlsx"; make_source(source)
    name_only = scan(NormalizeConfig(source, "People", 4, 4, True, "A", False, "B"))
    date_only = scan(NormalizeConfig(source, "People", 4, 4, False, "A", True, "B", "DMY"))
    assert {item.field for item in name_only.changes} == {"name"}
    assert {item.field for item in date_only.changes} == {"birthdate"}
