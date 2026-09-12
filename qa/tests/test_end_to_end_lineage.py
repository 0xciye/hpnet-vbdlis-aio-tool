from __future__ import annotations

import json
import shutil

from openpyxl import Workbook

from qa.generators.pipeline_fixture import REPO, SHEET, run_golden_pipeline
from tools.hpnet_file_generator.core.excel_reader import ExcelReader
from tools.notice_builder.core import ColumnMapping, inspect_workbook
from tools.vbdlis_excel_builder.core import BuilderService
from tools.vbdlis_excel_builder.models import MappingProfile
from tools.vbdlis_validation.models import ExcelConfig, RunConfig
from tools.vbdlis_validation.service import ValidationService


def test_golden_pipeline_has_no_silent_loss(tmp_path):
    stats,lineage=run_golden_pipeline(tmp_path)
    expected=json.loads((REPO/"qa/fixtures/golden_expected.json").read_text(encoding="utf-8"))

    assert stats==expected|{"input_unchanged":True}
    assert len(lineage)==10
    assert all(item["word_generated"] and item["tbxn_found"] and item["ddk_found"] for item in lineage)
    assert all(item["ax_tbxn_reference_found"] and item["ax_ddk_reference_found"] for item in lineage)


def test_one_owner_five_thousand_parcels_survives_reader(tmp_path):
    path=tmp_path/"large.xlsx"; workbook=Workbook(); sheet=workbook.active
    sheet.append(["STT","Tên hộ","Số tờ","Số thửa"])
    sheet.append([1,"NGUYỄN VĂN A",10,1])
    for parcel in range(2,5001): sheet.append([None,None,10,parcel])
    workbook.save(path); workbook.close()

    records=ExcelReader(str(path)).read_data(sheet.title,1,{"secondary_key":"STT","ho_ten":"Tên hộ","so_to":"Số tờ","so_thua":"Số thửa"})

    assert len(records)==1 and len(records[0].parcels)==5000
    assert not records[0].data_issues


def test_one_owner_twenty_parcels_survives_word_and_ddk_readers(tmp_path):
    path=tmp_path/"twenty.xlsx"; workbook=Workbook(); sheet=workbook.active; sheet.title=SHEET
    sheet.append(["STT","Tên hộ","Ngày sinh","CCCD","Địa chỉ","Nơi ở","Số tờ","Số thửa","Diện tích","Xứ đồng"])
    for parcel in range(1,21):
        sheet.append([1 if parcel==1 else None,"NGUYỄN VĂN A" if parcel==1 else None,1980 if parcel==1 else None,
                      "031080001234" if parcel==1 else None,None,None,10,parcel,100,"Đồng thử"])
    workbook.save(path); workbook.close()

    inspection=inspect_workbook(path,SHEET,1,1,ColumnMapping("B","G","H","I","J","D","A","C"),require_identity=True)
    records=ExcelReader(str(path)).read_data(SHEET,1,{"secondary_key":"STT","ho_ten":"Tên hộ","so_to":"Số tờ","so_thua":"Số thửa"})

    assert len(inspection.valid_records)==20
    assert len(records)==1 and len(records[0].parcels)==20


def test_five_members_ten_parcels_produce_fifty_vbdlis_rows(tmp_path):
    path=tmp_path/"members.xlsx"; workbook=Workbook(); sheet=workbook.active; sheet.title=SHEET
    sheet.append(["STT","Tên hộ","CCCD","Số tờ","Số thửa","Diện tích"])
    people=[(f"NGƯỜI KIỂM THỬ {letter}",f"03108{index:07d}") for index,letter in enumerate("ABCDE",1)]
    for parcel in range(100,110):
        for index,(name,cccd) in enumerate(people):
            sheet.append([1 if parcel==100 and index==0 else None,name,cccd,10,parcel,100])
    workbook.save(path); workbook.close()
    root=REPO/"src/tools/vbdlis_excel_builder"
    service=BuilderService(root/"resources/BieuMauThuThapThongTinGiayChungNhan.xlsx",root/"config/template_schema.json",tmp_path/"logs")
    profile=MappingProfile(commune_code="10930",address="Xã kiểm thử",source_mapping={
        "household_stt":"A","person_name":"B","cccd":"C","sheet_number":"D","parcel_number":"E","area":"F"
    })

    result=service.process(path,SHEET,1,profile)

    assert result.stats["people"]==5 and result.stats["parcels"]==10
    assert result.stats["output_rows"]==50
    assert not any(issue.code=="PERSON_PARCEL_PRODUCT_MISMATCH" for issue in result.issues)


def test_mixed_documents_are_explicitly_classified(tmp_path):
    run_golden_pipeline(tmp_path)
    tbxn=tmp_path/"tbxn/actual"; ddk=tmp_path/"ddk/actual"
    tbxn_files=sorted(tbxn.rglob("*.pdf")); ddk_files=sorted(ddk.rglob("*.pdf"))
    tbxn_files[0].unlink(); ddk_files[1].unlink()
    duplicate=tbxn/"trùng thư mục"/tbxn_files[2].name; duplicate.parent.mkdir(parents=True); shutil.copy2(tbxn_files[2],duplicate)
    upload=tmp_path/"vbdlis_prepare/actual/VBDLIS_GOLDEN.xlsx"
    source=tmp_path/"word_generation/input/golden_source.xlsx"
    result=ValidationService().run(RunConfig(
        ExcelConfig(upload,"Sheet1",4,{"person":"H","role":"N","sheet":"T","parcel":"U","combined":"AX"}),
        ExcelConfig(source,SHEET,1,{"owner":"B","sheet":"G","parcel":"H"}),
        tbxn,ddk,tmp_path/"mixed_validation","A","AZ",False,
    ))

    assert result.stats["missing_tbxn"]==1
    assert result.stats["missing_ddk"]==1
    assert result.stats["duplicate_tbxn"]==1
    assert all(record.workflow_status.value in {"NEED_SIGNATURE","MISSING_DOCUMENT","REVIEW_REQUIRED"} for record in result.records)
