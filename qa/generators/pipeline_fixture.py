from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import json
import re

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter

from tools.hpnet_file_generator.core.action_planner import ActionPlanner
from tools.hpnet_file_generator.core.excel_reader import ExcelReader
from tools.hpnet_file_generator.core.file_generator import FileGenerator
from tools.hpnet_file_generator.core.person_matcher import PersonMatcher
from tools.hpnet_file_generator.core.source_scanner import SourceScanner
from tools.hpnet_file_generator.models.data_models import ActionStatus, ProfileConfig
from tools.notice_builder.core import BatchConfig, ColumnMapping, NoticeService, inspect_workbook
from tools.notice_builder.paths import default_template_fields, default_template_path, resource
from tools.vbdlis_excel_builder.core import BuilderService
from tools.vbdlis_excel_builder.models import MappingProfile
from tools.vbdlis_validation.models import ExcelConfig, RunConfig
from tools.vbdlis_validation.service import ValidationService


REPO = Path(__file__).resolve().parents[2]
SHEET = "Dữ liệu nguồn"
MAPPING = ColumnMapping(
    household_index="A", owner="B", birth_date="C", identity="D",
    sheet="G", parcel="H", area="I", location="J",
)


def _stage(root: Path, name: str) -> dict[str, Path]:
    paths = {part: root / name / part for part in ("input", "expected", "actual", "logs")}
    for path in paths.values():
        path.mkdir(parents=True, exist_ok=True)
    return paths


def make_golden_source(path: Path, households: int = 5, parcels_per_household: int = 2) -> list[dict]:
    workbook = Workbook(); sheet = workbook.active; sheet.title = SHEET
    sheet.append(["STT", "Tên hộ", "Ngày sinh", "CCCD", "Địa chỉ", "Nơi ở", "Số tờ", "Số thửa", "Diện tích", "Xứ đồng"])
    lineage=[]
    owner_names=("NGUYỄN VĂN AN","TRẦN THỊ BÌNH","LÊ VĂN CHÍ","PHẠM THỊ DUNG","ĐỖ VĂN EM")
    for household in range(1, households + 1):
        owner=owner_names[household-1]
        for parcel_index in range(parcels_per_household):
            sheet_number=10 + household
            parcel_number=household * 100 + parcel_index
            row=sheet.max_row + 1
            sheet.append([
                household if parcel_index == 0 else None,
                owner if parcel_index == 0 else None,
                1980 + household if parcel_index == 0 else None,
                f"03108{household:07d}" if parcel_index == 0 else None,
                "Thôn kiểm thử" if parcel_index == 0 else None,
                "Xã kiểm thử" if parcel_index == 0 else None,
                float(sheet_number) if parcel_index else str(sheet_number),
                str(parcel_number) if parcel_index else float(parcel_number),
                100 + parcel_index,
                "Đồng kiểm thử",
            ])
            lineage.append({"test_case":f"CASE{len(lineage)+1:03d}","household":owner,"source_row":row,
                            "sheet":str(sheet_number),"parcel":str(parcel_number)})
    sheet.row_dimensions[lineage[-1]["source_row"]].hidden = True
    workbook.save(path); workbook.close()
    return lineage


def _write_pdf(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    writer=PdfWriter(); writer.add_blank_page(width=200,height=200)
    with path.open("wb") as stream: writer.write(stream)
    assert len(PdfReader(path).pages)==1


def _docx_text(path: Path) -> str:
    with ZipFile(path) as archive:
        xml=archive.read("word/document.xml").decode("utf-8")
    return " ".join(re.sub(r"<[^>]+>", " ", xml).split())


def run_golden_pipeline(root: Path) -> tuple[dict, list[dict]]:
    stages={name:_stage(root,name) for name in (
        "word_generation","tbxn","ddk","vbdlis_prepare","vbdlis_validation","end_to_end"
    )}
    source=stages["word_generation"]["input"] / "golden_source.xlsx"
    lineage=make_golden_source(source)
    source_hash=sha256(source.read_bytes()).hexdigest()

    inspection=inspect_workbook(source,SHEET,1,1,MAPPING,require_identity=True)
    legal=json.loads(resource("config/legal_defaults.json").read_text(encoding="utf-8"))
    notice=NoticeService(default_template_path(),legal)
    config=BatchConfig(
        "10930","Thôn kiểm thử, xã kiểm thử","Thôn kiểm thử","Xã kiểm thử","Xã kiểm thử","Xã kiểm thử",
        12,9,2026,number_mode="list",number_list="700",continue_number=701,
        template_fields=default_template_fields(),
    )
    word_output=stages["word_generation"]["actual"]
    preview=notice.preview(inspection,config,word_output,0,stages["word_generation"]["expected"])
    result=notice.generate(inspection,config,word_output,notice.approve(preview))
    words=sorted(word_output.glob("*.docx"))
    word_text={word.stem:_docx_text(word) for word in words}
    for record in inspection.valid_records:
        key=f"CHUACOGIAY_10930_{record.sheet}_{record.parcel}-TBXN"
        text=word_text[key]
        assert record.owner in text and record.sheet in text and record.parcel in text and record.area in text

    tbxn_root=stages["tbxn"]["actual"]
    for index,word in enumerate(words):
        _write_pdf(tbxn_root / f"nhóm {index%2+1}" / word.with_suffix(".pdf").name)

    source_people=stages["ddk"]["input"] / "nguồn người"
    source_people.mkdir(parents=True,exist_ok=True)
    for household,owner in enumerate(("NGUYỄN VĂN AN","TRẦN THỊ BÌNH","LÊ VĂN CHÍ","PHẠM THỊ DUNG","ĐỖ VĂN EM"),1):
        _write_pdf(source_people / f"{household}. {owner}.pdf")
    excel_reader=ExcelReader(str(source))
    people=excel_reader.read_data(SHEET,1,{"secondary_key":"STT","ho_ten":"Tên hộ","so_to":"Số tờ","so_thua":"Số thửa"})
    sources=SourceScanner(str(source_people),[".pdf"]).scan(); PersonMatcher(people,sources).match()
    ddk_root=stages["ddk"]["actual"]; ddk_root.mkdir(parents=True,exist_ok=True)
    actions=ActionPlanner(sources,str(ddk_root),ProfileConfig(ma_dvhc="10930",suffixes=["DDK"])).build_plan()
    for action in actions: FileGenerator().execute_action(action)
    assert all(action.status==ActionStatus.SUCCESS for action in actions)

    builder_root=REPO / "src/tools/vbdlis_excel_builder"
    builder=BuilderService(builder_root/"resources/BieuMauThuThapThongTinGiayChungNhan.xlsx",
                           builder_root/"config/template_schema.json",stages["vbdlis_prepare"]["logs"])
    profile=MappingProfile(commune_code="10930",address="Thôn kiểm thử, xã kiểm thử",output_folder=str(stages["vbdlis_prepare"]["actual"]),
        output_filename="VBDLIS_GOLDEN.xlsx",source_mapping={"household_stt":"A","person_name":"B","birth_date":"C","cccd":"D",
        "sheet_number":"G","parcel_number":"H","area":"I","land_location":"J"})
    transformed=builder.process(source,SHEET,1,profile)
    upload,_,verification=builder.export(transformed,profile,stages["vbdlis_prepare"]["actual"],"VBDLIS_GOLDEN.xlsx")
    assert verification["pass"]
    book=load_workbook(upload,read_only=True,data_only=False); upload_sheet=book.active
    rows=list(range(5,upload_sheet.max_row+1))
    ax_values=[str(upload_sheet[f"AX{row}"].value or "") for row in rows]
    upload_sheet_name=upload_sheet.title; book.close()

    validation=ValidationService().run(RunConfig(
        ExcelConfig(upload,upload_sheet_name,4,{"person":"H","role":"N","sheet":"T","parcel":"U","combined":"AX"}),
        ExcelConfig(source,SHEET,1,{"owner":"B","sheet":"G","parcel":"H"}),
        tbxn_root,ddk_root,stages["vbdlis_validation"]["actual"],"A","AZ",False,
    ))
    by_key={(record.sheet_normalized,record.parcel_normalized):record for record in validation.records}
    for item in lineage:
        record=by_key[(item["sheet"],item["parcel"])]
        item.update({"source_present":True,"word_expected":True,"word_generated":True,"tbxn_expected":True,
                     "tbxn_found":bool(record.tbxn_documents),"ddk_expected":True,"ddk_found":bool(record.ddk_documents),
                     "vbdlis_parcel_found":bool(record.upload_rows),"vbdlis_rows":len(record.upload_rows),
                     "ax_tbxn_reference_found":"-TBXN.pdf" in record.ax_raw,
                     "ax_ddk_reference_found":"-DDK.pdf" in record.ax_raw,"final_status":record.workflow_status.value,
                     "failure_stage":"","failure_reason":""})
    stats={"households":5,"source_parcels":len(lineage),"word_records":len(inspection.valid_records),"word_files":result["success"],
           "tbxn_pdfs":len(list(tbxn_root.rglob("*.pdf"))),"ddk_pdfs":len(list(ddk_root.rglob("*.pdf"))),
           "vbdlis_parcels":len(validation.records),"vbdlis_rows":validation.total_upload_rows,
           "ax_tbxn_references":sum("-TBXN.pdf" in value for value in ax_values),
           "ax_ddk_references":sum("-DDK.pdf" in value for value in ax_values),
           "missing_tbxn":validation.stats["missing_tbxn"],"missing_ddk":validation.stats["missing_ddk"],
           "silently_missing":sum(not item["vbdlis_parcel_found"] for item in lineage),
           "input_unchanged":sha256(source.read_bytes()).hexdigest()==source_hash}
    report=stages["end_to_end"]["logs"] / "golden_lineage.json"
    report.write_text(json.dumps({"stats":stats,"lineage":lineage},ensure_ascii=False,indent=2),encoding="utf-8")
    return stats,lineage


if __name__ == "__main__":
    run_golden_pipeline(REPO / "qa_runtime")
